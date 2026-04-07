from __future__ import annotations

import re
import time
from pathlib import Path
from typing import Any, Dict, List, Optional, Sequence

from openpyxl import load_workbook
from openpyxl.styles import Alignment
from openpyxl.utils import get_column_letter

from tools.pdf_fetcher.core.config import (
    ACTIVE_PROJECT_ROOT,
    BASE_AVAILABLE_STATUSES,
    BENCHMARK_RECORD_IDS,
    ENABLE_BENCHMARK_MODE,
    MAX_ROWS_TO_PROCESS,
    PDF_BASE_DIR,
    REQUIRED_COLUMNS,
    SPECIALIZED_DOWNLOADED_NOW_STATUSES,
    SHEET_NAME,
    TOOL_ROOT_DIR,
    build_current_workbook_path,
    build_final_workbook_path,
    build_history_workbook_path,
    build_temp_workbook_path,
)
from tools.pdf_fetcher.core.models import DownloadResult, Record, SessionState
from tools.pdf_fetcher.core.utils import build_pdf_filename


DOWNLOADED_STATUS_VALUES = BASE_AVAILABLE_STATUSES | SPECIALIZED_DOWNLOADED_NOW_STATUSES
DEFAULT_CELL_ALIGNMENT = Alignment(horizontal="left", vertical="top", wrap_text=False, shrink_to_fit=False)
WRAPPED_CELL_ALIGNMENT = Alignment(horizontal="left", vertical="top", wrap_text=True, shrink_to_fit=False)
CENTERED_CELL_ALIGNMENT = Alignment(horizontal="center", vertical="top", wrap_text=False, shrink_to_fit=False)
WRAPPED_HEADER_ALIGNMENT = Alignment(horizontal="center", vertical="center", wrap_text=True, shrink_to_fit=False)
HEADER_ROW_HEIGHT = 22
BODY_ROW_HEIGHT = 18
HEADER_WIDTH_PADDING = 3
HEADER_MAX_WIDTH = 36
COLUMN_WIDTH_PRESETS = {
    "record_id": 10,
    "source_record_id": 12,
    "master_id": 12,
    "source": 10,
    "source_priority": 12,
    "run_id": 12,
    "source_path": 18,
    "source_row_index": 10,
    "date_imported": 18,
    "Authors": 24,
    "Publication Year": 12,
    "Title": 34,
    "Abstract": 30,
    "Author Keywords": 24,
    "Source Title": 18,
    "Volume": 10,
    "Issue": 10,
    "Begin Page": 10,
    "End Page": 10,
    "Article Number": 12,
    "E-mail Address": 24,
    "E-mail Adress": 24,
    "DOI": 18,
    "DOI Link": 24,
    "pdf_downloaded": 12,
    "pdf_download_status": 20,
    "pdf_file_name": 28,
    "pdf_source_url": 22,
    "pdf_local_path": 24,
    "pdf_http_status": 12,
    "pdf_checked_at": 20,
    "Label": 14,
    "Reason to Exclude": 24,
    "Comment": 24,
}
WRAPPED_BODY_COLUMNS = {
    "Authors",
    "Title",
    "Abstract",
    "Author Keywords",
    "Source Title",
    "E-mail Address",
    "E-mail Adress",
    "DOI",
    "DOI Link",
    "pdf_download_status",
    "pdf_file_name",
    "pdf_source_url",
    "pdf_local_path",
    "pdf_checked_at",
    "Reason to Exclude",
    "Comment",
}
CENTERED_BODY_COLUMNS = {
    "record_id",
    "source_priority",
    "source_row_index",
    "Publication Year",
    "Volume",
    "Issue",
    "Begin Page",
    "End Page",
    "Article Number",
    "pdf_downloaded",
    "pdf_http_status",
    "Label",
}


def load_workbook_and_sheet(workbook_path: Path, sheet_name: str):
    """
    Open the Excel workbook and return (workbook, worksheet).
    """
    if not workbook_path.exists():
        raise FileNotFoundError(f"Workbook not found: {workbook_path}")

    wb = load_workbook(workbook_path)

    if sheet_name not in wb.sheetnames:
        raise ValueError(f"Sheet not found: {sheet_name}")

    ws = wb[sheet_name]
    return wb, ws


def build_col_map(ws) -> Dict[str, int]:
    """
    Build a mapping from header name to Excel column index.
    """
    col_map: Dict[str, int] = {}

    for col_idx in range(1, ws.max_column + 1):
        header = ws.cell(row=1, column=col_idx).value
        if header is not None and str(header).strip():
            col_map[str(header).strip()] = col_idx

    return col_map


def validate_required_columns(
    col_map: Dict[str, int],
    sheet_name: str,
    required_columns: Optional[Sequence[str]] = None,
) -> None:
    """
    Ensure all required columns exist in the target sheet.
    """
    required = list(required_columns or REQUIRED_COLUMNS)
    missing = [c for c in required if c not in col_map]
    if missing:
        raise ValueError(f"Missing required columns in '{sheet_name}': {missing}")


def get_cell(ws, row_idx: int, col_map: Dict[str, int], header: str):
    return ws.cell(row=row_idx, column=col_map[header]).value


def get_optional_cell(ws, row_idx: int, col_map: Dict[str, int], header: str, default: Any = ""):
    """
    Return the cell value if the column exists, otherwise return default.
    """
    if header not in col_map:
        return default
    return ws.cell(row=row_idx, column=col_map[header]).value


def set_cell(ws, row_idx: int, col_map: Dict[str, int], header: str, value) -> None:
    ws.cell(row=row_idx, column=col_map[header], value=value)


def extract_row_values(
    ws,
    row_idx: int,
    col_map: Dict[str, int],
    headers: Sequence[str],
) -> Dict[str, Any]:
    """
    Extract a dictionary of values for the requested headers.
    Missing columns are returned as empty strings.
    """
    return {header: get_optional_cell(ws, row_idx, col_map, header, "") for header in headers}


def is_blank(value) -> bool:
    return value is None or str(value).strip() == ""


def normalize_doi_like(value: str) -> Optional[str]:
    """
    Lightweight DOI normalizer for workbook eligibility checks.
    """
    if is_blank(value):
        return None

    doi = str(value).strip()
    doi = re.sub(r"^https?://(dx\.)?doi\.org/", "", doi, flags=re.IGNORECASE).strip()

    if re.match(r"^10\.\d{4,9}/\S+$", doi):
        return doi

    return None


def row_has_usable_doi(ws, row_idx: int, col_map: Dict[str, int]) -> bool:
    """
    Return True when the row already has a usable DOI or DOI link.
    """
    doi_raw = str(get_optional_cell(ws, row_idx, col_map, "DOI", "") or "").strip()
    doi_link_raw = str(get_optional_cell(ws, row_idx, col_map, "DOI Link", "") or "").strip()

    if doi_link_raw:
        return True

    return normalize_doi_like(doi_raw) is not None


def row_is_already_downloaded(ws, row_idx: int, col_map: Dict[str, int]) -> bool:
    """
    Return True when the row already has a downloaded PDF according to workbook state.
    """
    pdf_downloaded = get_optional_cell(ws, row_idx, col_map, "pdf_downloaded", "")
    pdf_status = str(get_optional_cell(ws, row_idx, col_map, "pdf_download_status", "") or "").strip().lower()

    if str(pdf_downloaded).strip() == "1":
        return True

    return pdf_status in DOWNLOADED_STATUS_VALUES


def parse_pdf_downloaded_flag(value: Any) -> int:
    try:
        return 1 if int(str(value or "0").strip()) == 1 else 0
    except Exception:
        return 0


def status_indicates_downloaded(status: Any) -> bool:
    normalized = str(status or "").strip().lower()
    return normalized in DOWNLOADED_STATUS_VALUES


def resolve_existing_pdf_path_for_row(
    ws,
    row_idx: int,
    col_map: Dict[str, int],
) -> Optional[Path]:
    candidate_paths: list[Path] = []

    pdf_local_path = str(get_optional_cell(ws, row_idx, col_map, "pdf_local_path", "") or "").strip()
    if pdf_local_path:
        local_path = Path(pdf_local_path)
        if local_path.is_absolute():
            candidate_paths.append(local_path)
        else:
            candidate_paths.append(ACTIVE_PROJECT_ROOT / local_path)

    pdf_file_name = str(get_optional_cell(ws, row_idx, col_map, "pdf_file_name", "") or "").strip()
    if pdf_file_name:
        candidate_paths.append(PDF_BASE_DIR / pdf_file_name)

    record_id = str(get_optional_cell(ws, row_idx, col_map, "record_id", "") or "").strip()
    title = str(get_optional_cell(ws, row_idx, col_map, "Title", "") or "").strip()
    if record_id and title:
        candidate_paths.append(PDF_BASE_DIR / build_pdf_filename(record_id, title))

    seen: set[str] = set()
    for candidate in candidate_paths:
        normalized = str(candidate.resolve(strict=False)).lower()
        if normalized in seen:
            continue
        seen.add(normalized)

        try:
            if candidate.exists() and candidate.is_file() and candidate.stat().st_size > 0:
                return candidate
        except Exception:
            continue

    return None


def reconcile_pdf_state_for_row(
    ws,
    row_idx: int,
    col_map: Dict[str, int],
) -> bool:
    record_id = str(get_optional_cell(ws, row_idx, col_map, "record_id", "") or "").strip()
    if not record_id:
        return False

    current_downloaded = parse_pdf_downloaded_flag(get_optional_cell(ws, row_idx, col_map, "pdf_downloaded", 0))
    current_status = str(get_optional_cell(ws, row_idx, col_map, "pdf_download_status", "") or "").strip()
    existing_pdf_path = resolve_existing_pdf_path_for_row(ws, row_idx, col_map)
    changed = False

    if existing_pdf_path is not None:
        desired_file_name = existing_pdf_path.name
        try:
            desired_local_path = str(existing_pdf_path.relative_to(ACTIVE_PROJECT_ROOT))
        except ValueError:
            desired_local_path = str(existing_pdf_path)

        if current_downloaded != 1:
            set_cell(ws, row_idx, col_map, "pdf_downloaded", 1)
            changed = True

        if not status_indicates_downloaded(current_status):
            set_cell(ws, row_idx, col_map, "pdf_download_status", "downloaded")
            changed = True

        if str(get_optional_cell(ws, row_idx, col_map, "pdf_file_name", "") or "").strip() != desired_file_name:
            set_cell(ws, row_idx, col_map, "pdf_file_name", desired_file_name)
            changed = True

        if str(get_optional_cell(ws, row_idx, col_map, "pdf_local_path", "") or "").strip() != desired_local_path:
            set_cell(ws, row_idx, col_map, "pdf_local_path", desired_local_path)
            changed = True

        return changed

    if current_downloaded != 0:
        set_cell(ws, row_idx, col_map, "pdf_downloaded", 0)
        changed = True

    if status_indicates_downloaded(current_status):
        set_cell(ws, row_idx, col_map, "pdf_download_status", "file_missing")
        changed = True

    return changed


def reconcile_pdf_state_with_filesystem(ws, col_map: Dict[str, int]) -> dict:
    if "record_id" not in col_map or "pdf_downloaded" not in col_map or "pdf_download_status" not in col_map:
        return {"rows_checked": 0, "rows_changed": 0, "downloaded_in_folder": 0}

    rows_checked = 0
    rows_changed = 0
    downloaded_in_folder = 0

    for row_idx in range(2, ws.max_row + 1):
        record_id = str(get_optional_cell(ws, row_idx, col_map, "record_id", "") or "").strip()
        if not record_id:
            continue

        rows_checked += 1
        existing_pdf_path = resolve_existing_pdf_path_for_row(ws, row_idx, col_map)
        if existing_pdf_path is not None:
            downloaded_in_folder += 1

        if reconcile_pdf_state_for_row(ws, row_idx, col_map):
            rows_changed += 1

    return {
        "rows_checked": rows_checked,
        "rows_changed": rows_changed,
        "downloaded_in_folder": downloaded_in_folder,
    }


def build_record(ws, row_idx: int, col_map: Dict[str, int]) -> Optional[Record]:
    """
    Build a Record object from a given worksheet row.
    Return None when the row is not usable.
    """
    record_id = str(get_cell(ws, row_idx, col_map, "record_id") or "").strip()
    if not record_id:
        return None

    if ENABLE_BENCHMARK_MODE and BENCHMARK_RECORD_IDS and record_id not in BENCHMARK_RECORD_IDS:
        return None

    title = str(get_cell(ws, row_idx, col_map, "Title") or "").strip()
    doi_raw = str(get_cell(ws, row_idx, col_map, "DOI") or "").strip()
    doi_link_raw = str(get_cell(ws, row_idx, col_map, "DOI Link") or "").strip()

    file_name = build_pdf_filename(record_id, title)
    pdf_path = PDF_BASE_DIR / file_name
    try:
        relative_path = str(pdf_path.relative_to(ACTIVE_PROJECT_ROOT))
    except ValueError:
        relative_path = str(Path(TOOL_ROOT_DIR.name) / "PDFs" / file_name)

    return Record(
        row_idx=row_idx,
        record_id=record_id,
        title=title,
        doi_raw=doi_raw,
        doi_link_raw=doi_link_raw,
        file_name=file_name,
        pdf_path=pdf_path,
        relative_path=relative_path,
    )


def collect_records(ws, col_map: Dict[str, int], max_rows_to_process: Optional[int]) -> List[Record]:
    """
    Read all usable rows from the worksheet and return Record objects.
    """
    records: List[Record] = []
    logical_seen = 0

    for row_idx in range(2, ws.max_row + 1):
        record = build_record(ws, row_idx, col_map)
        if record is None:
            continue

        records.append(record)
        logical_seen += 1

        if max_rows_to_process is not None and logical_seen >= max_rows_to_process:
            break

    return records


def collect_phase0_candidate_rows(
    ws,
    col_map: Dict[str, int],
    max_rows_to_process: Optional[int],
) -> List[int]:
    """
    Return worksheet row indices eligible for DOI enrichment.

    Rules:
    - must have record_id
    - must have Title
    - must NOT already have a usable DOI
    - must NOT already have a downloaded PDF
    """
    candidate_rows: List[int] = []
    logical_seen = 0

    for row_idx in range(2, ws.max_row + 1):
        record_id = str(get_optional_cell(ws, row_idx, col_map, "record_id", "") or "").strip()
        if not record_id:
            continue

        if ENABLE_BENCHMARK_MODE and BENCHMARK_RECORD_IDS and record_id not in BENCHMARK_RECORD_IDS:
            continue

        title = str(get_optional_cell(ws, row_idx, col_map, "Title", "") or "").strip()
        if not title:
            continue

        if row_has_usable_doi(ws, row_idx, col_map):
            continue

        if row_is_already_downloaded(ws, row_idx, col_map):
            continue

        candidate_rows.append(row_idx)
        logical_seen += 1

        if max_rows_to_process is not None and logical_seen >= max_rows_to_process:
            break

    return candidate_rows


def write_result(ws, row_idx: int, col_map: Dict[str, int], result: DownloadResult) -> None:
    """
    Write a DownloadResult back into the worksheet.
    """
    set_cell(ws, row_idx, col_map, "pdf_downloaded", result.pdf_downloaded)
    set_cell(ws, row_idx, col_map, "pdf_download_status", result.pdf_download_status)
    set_cell(ws, row_idx, col_map, "pdf_file_name", result.pdf_file_name)
    set_cell(ws, row_idx, col_map, "pdf_source_url", result.pdf_source_url)
    set_cell(ws, row_idx, col_map, "pdf_local_path", result.pdf_local_path)
    set_cell(
        ws,
        row_idx,
        col_map,
        "pdf_http_status",
        result.pdf_http_status if result.pdf_http_status is not None else "",
    )
    set_cell(ws, row_idx, col_map, "pdf_checked_at", result.pdf_checked_at)


def write_doi_enrichment_result(
    ws,
    row_idx: int,
    col_map: Dict[str, int],
    doi: str,
    doi_link: str,
) -> None:
    """
    Write DOI enrichment results back into the main worksheet.

    Important:
    - only DOI and DOI Link are written
    - no auxiliary DOI columns are created in the main workbook
    """
    set_cell(ws, row_idx, col_map, "DOI", doi)
    set_cell(ws, row_idx, col_map, "DOI Link", doi_link)


def build_download_result_from_row(
    ws,
    row_idx: int,
    col_map: Dict[str, int],
    record: Record,
) -> DownloadResult:
    """
    Rebuild a DownloadResult from workbook values already written by phase 1/2.
    """
    pdf_downloaded_raw = get_optional_cell(ws, row_idx, col_map, "pdf_downloaded", 0)
    pdf_downloaded = parse_pdf_downloaded_flag(pdf_downloaded_raw)

    pdf_http_status_raw = get_optional_cell(ws, row_idx, col_map, "pdf_http_status", "")
    pdf_http_status = None
    try:
        if str(pdf_http_status_raw or "").strip():
            pdf_http_status = int(str(pdf_http_status_raw).strip())
    except Exception:
        pdf_http_status = None

    status = str(get_optional_cell(ws, row_idx, col_map, "pdf_download_status", "") or "").strip()
    source_url = str(get_optional_cell(ws, row_idx, col_map, "pdf_source_url", "") or "").strip()
    local_path = str(get_optional_cell(ws, row_idx, col_map, "pdf_local_path", record.relative_path) or record.relative_path).strip()
    checked_at = str(get_optional_cell(ws, row_idx, col_map, "pdf_checked_at", "") or "").strip()

    # If the workbook is blank but the physical PDF exists, use the file system as
    # the source of truth so phase 2 can still be reconstructed after reopening.
    if record.pdf_path.exists() and record.pdf_path.stat().st_size > 0:
        if pdf_downloaded != 1:
            pdf_downloaded = 1
        if not status:
            status = "duplicate_pdf"
        if not local_path:
            local_path = record.relative_path
    elif not status:
        # Workbook rows with no saved phase-1 status but no physical PDF should
        # remain eligible for specialized retry reconstruction.
        status = "not_downloaded"

    return DownloadResult(
        pdf_downloaded=pdf_downloaded,
        pdf_download_status=status,
        pdf_file_name=str(get_optional_cell(ws, row_idx, col_map, "pdf_file_name", record.file_name) or record.file_name).strip(),
        pdf_source_url=source_url,
        pdf_local_path=local_path,
        pdf_http_status=pdf_http_status,
        pdf_checked_at=checked_at,
    )


def reconstruct_phase1_results_from_workbook(
    session_state: SessionState,
) -> tuple[Dict[int, DownloadResult], dict]:
    """
    Rebuild phase 1 outputs from the current workbook so phase 2 can start directly
    after reopening the app.
    """
    if not session_state.is_workbook_loaded:
        raise ValueError("SessionState workbook is not loaded.")

    phase1_results_by_row: Dict[int, DownloadResult] = {}
    pdfs_available_global = 0

    for record in session_state.records:
        result = build_download_result_from_row(
            session_state.worksheet,
            record.row_idx,
            session_state.col_map,
            record,
        )
        phase1_results_by_row[record.row_idx] = result

        if result.pdf_download_status in BASE_AVAILABLE_STATUSES or result.pdf_download_status in SPECIALIZED_DOWNLOADED_NOW_STATUSES:
            pdfs_available_global += 1

    phase1_summary = {
        "processed": len(session_state.records),
        "total": len(session_state.records),
        "workers": 0,
        "start_time": time.time(),
        "last_record_id": "",
        "last_status": "reconstructed_from_workbook",
        "pdfs_available_global": pdfs_available_global,
        "pdfs_in_folder": sum(1 for path in PDF_BASE_DIR.glob("*.pdf") if path.is_file()) if PDF_BASE_DIR.exists() else 0,
        "downloaded_now_total": 0,
        "stopped_early": False,
    }

    return phase1_results_by_row, phase1_summary


def refresh_record_from_worksheet(ws, row_idx: int, col_map: Dict[str, int]) -> Optional[Record]:
    """
    Rebuild a single Record object from the current worksheet state.
    Useful after DOI enrichment changes.
    """
    return build_record(ws, row_idx, col_map)


def refresh_records_from_worksheet(
    ws,
    col_map: Dict[str, int],
    max_rows_to_process: Optional[int] = None,
) -> List[Record]:
    """
    Rebuild the full Record list from the current worksheet state.
    """
    return collect_records(ws, col_map, max_rows_to_process)


def initialize_session_state(
    workbook_path: Path,
    sheet_name: str = SHEET_NAME,
    max_rows_to_process: Optional[int] = None,
) -> SessionState:
    """
    Load workbook + worksheet into a SessionState and build the initial records list.
    """
    wb, ws = load_workbook_and_sheet(workbook_path, sheet_name)
    col_map = build_col_map(ws)
    validate_required_columns(col_map, sheet_name)

    effective_max_rows = max_rows_to_process if max_rows_to_process is not None else MAX_ROWS_TO_PROCESS
    records = collect_records(ws, col_map, effective_max_rows)

    return SessionState(
        workbook_path=workbook_path,
        sheet_name=sheet_name,
        project_root=ACTIVE_PROJECT_ROOT,
        tool_root=TOOL_ROOT_DIR,
        workbook=wb,
        worksheet=ws,
        col_map=col_map,
        records=records,
        current_workbook_path=workbook_path,
    )


def refresh_session_records(
    session_state: SessionState,
    max_rows_to_process: Optional[int] = None,
) -> None:
    """
    Rebuild session_state.records from the current worksheet state.
    """
    if not session_state.is_workbook_loaded:
        raise ValueError("SessionState workbook is not loaded.")

    effective_max_rows = max_rows_to_process if max_rows_to_process is not None else MAX_ROWS_TO_PROCESS
    session_state.records = refresh_records_from_worksheet(
        session_state.worksheet,
        session_state.col_map,
        effective_max_rows,
    )


def rebuild_session_column_map(session_state: SessionState) -> None:
    """
    Refresh the header map from the current worksheet.
    """
    if not session_state.is_workbook_loaded:
        raise ValueError("SessionState workbook is not loaded.")

    session_state.col_map = build_col_map(session_state.worksheet)


def apply_contained_cell_layout(ws) -> None:
    """
    Keep cell contents visually contained inside cell boundaries.
    """
    headers_by_col: Dict[int, str] = {}
    for col_idx in range(1, ws.max_column + 1):
        header = str(ws.cell(row=1, column=col_idx).value or "").strip()
        headers_by_col[col_idx] = header
        column_letter = get_column_letter(col_idx)
        width = COLUMN_WIDTH_PRESETS.get(header)
        if width is None:
            width = min(max(len(header) + 2, 10), 18) if header else 12
        if header:
            width = max(width, min(len(header) + HEADER_WIDTH_PADDING, HEADER_MAX_WIDTH))
        ws.column_dimensions[column_letter].width = width

    ws.row_dimensions[1].height = HEADER_ROW_HEIGHT

    for row_idx in range(1, ws.max_row + 1):
        if row_idx > 1:
            ws.row_dimensions[row_idx].height = BODY_ROW_HEIGHT
        for col_idx in range(1, ws.max_column + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            if row_idx == 1:
                cell.alignment = WRAPPED_HEADER_ALIGNMENT
                continue

            header = headers_by_col.get(col_idx, "")
            if header in CENTERED_BODY_COLUMNS:
                cell.alignment = CENTERED_CELL_ALIGNMENT
            elif header in WRAPPED_BODY_COLUMNS:
                cell.alignment = WRAPPED_CELL_ALIGNMENT
            else:
                cell.alignment = DEFAULT_CELL_ALIGNMENT


def safe_save_workbook(wb, output_path: Path, max_attempts: int = 2) -> bool:
    """
    Save workbook safely through a temporary file and replacement.
    """
    output_path.parent.mkdir(parents=True, exist_ok=True)

    try:
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            if sheet_name == SHEET_NAME:
                col_map = build_col_map(ws)
                reconcile_pdf_state_with_filesystem(ws, col_map)
            apply_contained_cell_layout(ws)
    except Exception:
        pass

    for attempt in range(1, max_attempts + 1):
        try:
            temp_path = output_path.with_name(output_path.stem + ".__tmp__.xlsx")
            wb.save(temp_path)

            if output_path.exists():
                try:
                    output_path.unlink()
                except PermissionError:
                    pass

            temp_path.replace(output_path)
            return True

        except KeyboardInterrupt:
            raise
        except Exception:
            if attempt == max_attempts:
                return False
            time.sleep(0.5)

    return False


def save_session_checkpoint(
    session_state: SessionState,
    label: str,
    timestamp: Optional[str] = None,
) -> Optional[Path]:
    """
    Save a temporary checkpoint workbook for recovery/debugging purposes.
    """
    if not session_state.is_workbook_loaded:
        raise ValueError("SessionState workbook is not loaded.")

    checkpoint_path = build_temp_workbook_path(label=label, timestamp=timestamp)
    ok = safe_save_workbook(session_state.workbook, checkpoint_path)
    if not ok:
        return None

    session_state.temp_workbook_path = checkpoint_path
    return checkpoint_path


def save_final_workbook(
    session_state: SessionState,
    output_path: Optional[Path] = None,
    timestamp: Optional[str] = None,
) -> Optional[Path]:
    """
    Save the final workbook for the session.
    """
    if not session_state.is_workbook_loaded:
        raise ValueError("SessionState workbook is not loaded.")

    effective_timestamp = timestamp or None
    current_path = build_current_workbook_path()
    history_path = build_history_workbook_path(timestamp=effective_timestamp)
    final_path = output_path or build_final_workbook_path(timestamp=effective_timestamp)

    save_targets = [current_path, history_path]
    if final_path not in save_targets:
        save_targets.append(final_path)

    for target_path in save_targets:
        ok = safe_save_workbook(session_state.workbook, target_path)
        if not ok:
            return None

    session_state.workbook_path = current_path
    session_state.current_workbook_path = current_path
    session_state.history_workbook_path = history_path
    session_state.final_workbook_path = final_path
    return final_path


def save_current_workbook(session_state: SessionState) -> Optional[Path]:
    """
    Save the active in-memory workbook to the project's Current workbook path.
    """
    if not session_state.is_workbook_loaded:
        raise ValueError("SessionState workbook is not loaded.")

    current_path = build_current_workbook_path()
    ok = safe_save_workbook(session_state.workbook, current_path)
    if not ok:
        return None

    session_state.workbook_path = current_path
    session_state.current_workbook_path = current_path
    return current_path


def reload_session_from_current(
    session_state: SessionState,
    max_rows_to_process: Optional[int] = None,
) -> SessionState:
    """
    Reload the session from the Current workbook so the next phase starts from
    the latest persisted workbook state.
    """
    current_path = build_current_workbook_path()
    reloaded = initialize_session_state(
        workbook_path=current_path,
        sheet_name=session_state.sheet_name or SHEET_NAME,
        max_rows_to_process=max_rows_to_process,
    )
    reloaded.timestamp = session_state.timestamp
    reloaded._dashboard_reporter = getattr(session_state, "_dashboard_reporter", None)
    return reloaded


def current_workbook_has_phase1_data(sheet_name: str = SHEET_NAME) -> bool:
    """
    Return True when the Current workbook already contains saved phase-1 style
    PDF fields, so the app can safely allow a direct start from phase 2.
    """
    current_path = build_current_workbook_path()
    if not current_path.exists():
        return False

    try:
        wb, ws = load_workbook_and_sheet(current_path, sheet_name)
        col_map = build_col_map(ws)
    except Exception:
        return False

    phase1_markers = ("pdf_downloaded", "pdf_download_status", "pdf_source_url", "pdf_checked_at")
    available_markers = [header for header in phase1_markers if header in col_map]
    if not available_markers:
        return False

    for row_idx in range(2, ws.max_row + 1):
        for header in available_markers:
            value = get_optional_cell(ws, row_idx, col_map, header, "")
            if value is None:
                continue
            if str(value).strip() != "":
                wb.close()
                return True

    wb.close()
    return False
