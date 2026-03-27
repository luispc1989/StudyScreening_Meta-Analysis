from __future__ import annotations

import re
import threading
import time
import unicodedata
from dataclasses import dataclass
from difflib import SequenceMatcher
from typing import Any, Dict, List, Optional, Sequence

import requests
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from requests.adapters import HTTPAdapter
from urllib3.util.retry import Retry

from tools.pdf_fetcher.core.config import (
    BENCHMARK_RECORD_IDS,
    DOI_ACCEPT_FIRST_VARIANT_IF_PLAUSIBLE,
    DOI_AUTO_APPLY_MIN_AUTHOR_SCORE,
    DOI_AUTO_APPLY_MIN_CONFIDENCE_SCORE,
    DOI_EARLY_STOP_MIN_CONFIDENCE_SCORE,
    DOI_EARLY_STOP_MIN_SOURCE_COUNT,
    DOI_EARLY_STOP_MIN_TITLE_SIMILARITY,
    DOI_AUTO_APPLY_MIN_SOURCE_COUNT,
    DOI_AUTO_APPLY_MIN_TITLE_SIMILARITY,
    DOI_AUTO_APPLY_MIN_TOP2_GAP,
    DOI_AUTO_APPLY_MIN_YEAR_SCORE,
    DOI_CONNECT_TIMEOUT,
    DOI_CROSSREF_MAILTO,
    DOI_GENERATE_REPORT_BY_DEFAULT,
    DOI_ENABLE_SOURCE_CASCADE,
    DOI_MAX_ROWS_TO_PROCESS,
    DOI_MAX_QUERY_VARIANTS,
    DOI_MIN_CONFIDENCE_SCORE,
    DOI_MIN_TITLE_SIMILARITY,
    DOI_NEEDS_REVIEW_MIN_CONFIDENCE_SCORE,
    DOI_NEEDS_REVIEW_MIN_TITLE_SIMILARITY,
    DOI_OVERWRITE_EXISTING,
    DOI_PRIMARY_SOURCE_ORDER,
    DOI_READ_TIMEOUT,
    DOI_REPORT_ALL_PROCESSED_SHEET_NAME,
    DOI_REPORT_ENRICHED_SHEET_NAME,
    DOI_REPORT_NEEDS_REVIEW_SHEET_NAME,
    DOI_REPORT_STATS_SHEET_NAME,
    DOI_REPORT_UNRESOLVED_SHEET_NAME,
    DOI_REQUIRED_INPUT_COLUMNS,
    DOI_REQUEST_TIMEOUT,
    DOI_SECONDARY_SOURCE_ORDER,
    DOI_SLEEP_BETWEEN_REQUESTS,
    ENABLE_BENCHMARK_MODE,
    REPORT_NAME_FASE0_DOI,
    build_report_output_path,
)
from tools.pdf_fetcher.core.excel_io import (
    get_optional_cell,
    refresh_session_records,
    row_has_usable_doi,
    row_is_already_downloaded,
    safe_save_workbook,
    validate_required_columns,
    write_doi_enrichment_result,
)
from tools.pdf_fetcher.core.models import Phase0Summary, SessionState


ReporterType = Optional[callable]
_thread_local = threading.local()
_doi_query_cache: dict[tuple[str, str], list[dict]] = {}


# =========================
# STYLE CONSTANTS
# =========================

GREEN_TITLE_FILL = "FF1F6E43"
DARK_HEADER_FILL = "FF000000"
LIGHT_BLUE_HEADER_FILL = "FFD9EAF7"
WHITE_FONT = "FFFFFFFF"

CENTER = Alignment(horizontal="center", vertical="center")
CENTER_TOP = Alignment(horizontal="center", vertical="top")


# =========================
# DATA STRUCTURES
# =========================

@dataclass
class Candidate:
    doi: str
    doi_link: str
    source: str
    matched_title: str
    matched_year: str
    matched_journal: str
    matched_authors: str
    source_count: int
    title_similarity: float
    year_score: float
    author_score: float
    confidence_score: float


@dataclass
class Phase0RowResult:
    record_id: str
    title: str
    authors: str
    publication_year: str
    original_doi: str
    original_doi_link: str
    phase0_status: str
    matched_doi: str
    matched_doi_link: str
    confidence_score: float
    title_similarity: float
    year_score: float
    author_score: float
    match_source_count: int
    match_source: str
    matched_title: str
    matched_year: str
    matched_journal: str
    second_candidate_doi: str
    second_candidate_title: str
    second_candidate_confidence: float
    top2_gap: float
    decision_reason: str


# =========================
# EVENT EMITTER
# =========================

def emit(reporter: ReporterType, event_name: str, payload: dict) -> None:
    if reporter is None:
        return
    reporter(event_name, payload)


def set_current_stop_event(stop_event) -> None:
    _thread_local.stop_event = stop_event


def clear_current_stop_event() -> None:
    if hasattr(_thread_local, "stop_event"):
        delattr(_thread_local, "stop_event")


def get_current_stop_event():
    return getattr(_thread_local, "stop_event", None)


def is_stop_requested() -> bool:
    stop_event = get_current_stop_event()
    return bool(stop_event is not None and stop_event.is_set())


def interruptible_sleep(seconds: float, step: float = 0.05) -> None:
    if seconds <= 0:
        return

    remaining = float(seconds)
    while remaining > 0:
        if is_stop_requested():
            return
        current_step = min(step, remaining)
        time.sleep(current_step)
        remaining -= current_step


# =========================
# TEXT / DOI HELPERS
# =========================

def is_blank(value) -> bool:
    return value is None or str(value).strip() == ""


def normalize_text(value: str) -> str:
    value = str(value or "").strip()
    value = unicodedata.normalize("NFKD", value)
    value = "".join(ch for ch in value if not unicodedata.combining(ch))
    value = value.lower()
    value = re.sub(r"\s+", " ", value)
    return value.strip()


def normalize_title(value: str) -> str:
    value = normalize_text(value)
    value = re.sub(r"[^\w\s]", " ", value)
    value = re.sub(r"\s+", " ", value)
    return value.strip()


def normalize_doi(value: str) -> Optional[str]:
    if is_blank(value):
        return None

    doi = str(value).strip()
    doi = re.sub(r"^https?://(dx\.)?doi\.org/", "", doi, flags=re.IGNORECASE)
    doi = doi.strip()

    if re.match(r"^10\.\d{4,9}/\S+$", doi):
        return doi

    return None


def make_doi_url(doi: str) -> str:
    return f"https://doi.org/{doi}"


def parse_year(value: Any) -> Optional[int]:
    if value is None:
        return None

    text = str(value).strip()
    if not text:
        return None

    match = re.search(r"\b(19|20)\d{2}\b", text)
    if not match:
        return None

    try:
        return int(match.group(0))
    except ValueError:
        return None


def safe_float(value: Any) -> float:
    try:
        return float(value)
    except Exception:
        return 0.0


def extract_author_surnames(authors_value: Any) -> set[str]:
    """
    Extract a conservative surname set from the author string.
    """
    text = normalize_text(str(authors_value or ""))
    if not text:
        return set()

    text = re.sub(r"\band\b", ";", text)
    text = text.replace("|", ";")
    parts = re.split(r";|\n", text)

    surnames: set[str] = set()

    for part in parts:
        cleaned = normalize_text(part)
        if not cleaned:
            continue

        cleaned = re.sub(r"[^\w\s,-]", " ", cleaned).strip()
        if not cleaned:
            continue

        tokens = [t for t in re.split(r"[\s,]+", cleaned) if t]
        if not tokens:
            continue

        surname = tokens[0] if "," in cleaned else tokens[-1]
        surname = surname.strip().lower()

        if len(surname) >= 3 and surname not in {"et", "al"}:
            surnames.add(surname)

    return surnames


def join_authors(authors: Sequence[str]) -> str:
    values = [str(a).strip() for a in authors if str(a).strip()]
    return "; ".join(values)


# =========================
# SCORING
# =========================

def compute_title_similarity(source_title: str, candidate_title: str) -> float:
    src = normalize_title(source_title)
    cand = normalize_title(candidate_title)

    if not src or not cand:
        return 0.0

    return SequenceMatcher(None, src, cand).ratio()


def compute_year_score(source_year: Optional[int], candidate_year: Optional[int]) -> float:
    if source_year is None or candidate_year is None:
        return 0.0

    delta = abs(source_year - candidate_year)
    if delta == 0:
        return 100.0
    if delta == 1:
        return 70.0
    if delta == 2:
        return 40.0
    return 0.0


def compute_author_score(source_authors: Any, candidate_authors: Sequence[str]) -> float:
    source_surnames = extract_author_surnames(source_authors)
    candidate_surnames = set()

    for author in candidate_authors:
        candidate_surnames.update(extract_author_surnames(author))

    if not source_surnames or not candidate_surnames:
        return 0.0

    overlap = len(source_surnames.intersection(candidate_surnames))
    return min(100.0, (overlap / len(source_surnames)) * 100.0)


def compute_confidence_score(
    title_similarity: float,
    year_score: float,
    author_score: float,
) -> float:
    return (title_similarity * 100.0 * 0.70) + (year_score * 0.15) + (author_score * 0.15)


def build_query_variants(title: str) -> list[str]:
    base_title = str(title or "").strip()
    if not base_title:
        return []

    variants: list[str] = []

    def _add(candidate: str) -> None:
        cleaned = re.sub(r"\s+", " ", str(candidate or "").strip())
        if cleaned and cleaned not in variants:
            variants.append(cleaned)

    _add(base_title)
    _add(re.sub(r"\s*\([^)]*\)", "", base_title).strip())
    _add(base_title.split(":", 1)[0].strip())
    _add(" ".join(base_title.split()[:12]).strip())
    _add(normalize_title(base_title))

    return variants[:DOI_MAX_QUERY_VARIANTS]


# =========================
# HTTP SESSION
# =========================

def build_retry_strategy() -> Retry:
    return Retry(
        total=2,
        connect=2,
        read=2,
        redirect=2,
        status=2,
        status_forcelist=[429, 500, 502, 503, 504],
        allowed_methods=frozenset(["GET"]),
        backoff_factor=0.4,
        raise_on_status=False,
        respect_retry_after_header=True,
    )


def get_thread_session() -> requests.Session:
    session = getattr(_thread_local, "session", None)
    if session is not None:
        return session

    session = requests.Session()
    retry_strategy = build_retry_strategy()

    adapter = HTTPAdapter(
        max_retries=retry_strategy,
        pool_connections=20,
        pool_maxsize=20,
    )
    session.mount("http://", adapter)
    session.mount("https://", adapter)

    session.headers.update(
        {
            "User-Agent": (
                "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
                "AppleWebKit/537.36 (KHTML, like Gecko) "
                "Chrome/122.0 Safari/537.36"
            ),
            "Accept": "application/json,text/plain,*/*",
        }
    )

    if DOI_CROSSREF_MAILTO:
        session.headers["mailto"] = DOI_CROSSREF_MAILTO

    _thread_local.session = session
    return session


def safe_get_json(
    session: requests.Session,
    url: str,
    params: Optional[dict] = None,
) -> Optional[dict]:
    if is_stop_requested():
        return None
    try:
        response = session.get(
            url,
            params=params,
            timeout=DOI_REQUEST_TIMEOUT,
        )
        response.raise_for_status()
        return response.json()
    except Exception:
        return None


# =========================
# SOURCE QUERIES
# =========================

def search_crossref(title: str) -> list[dict]:
    session = get_thread_session()

    params = {
        "query.bibliographic": title,
        "rows": 3,
        "select": "DOI,title,author,container-title,published-print,published-online,issued",
    }
    payload = safe_get_json(session, "https://api.crossref.org/works", params=params)
    if not payload:
        return []

    items = payload.get("message", {}).get("items", [])
    results: list[dict] = []

    for item in items:
        doi = normalize_doi(item.get("DOI"))
        if not doi:
            continue

        title_values = item.get("title") or []
        container_values = item.get("container-title") or []

        authors = []
        for author in item.get("author") or []:
            given = str(author.get("given") or "").strip()
            family = str(author.get("family") or "").strip()
            full_name = " ".join(part for part in [given, family] if part)
            if full_name:
                authors.append(full_name)

        year = None
        for key in ("published-print", "published-online", "issued"):
            date_parts = item.get(key, {}).get("date-parts") or []
            if date_parts and date_parts[0]:
                try:
                    year = int(date_parts[0][0])
                    break
                except Exception:
                    pass

        results.append(
            {
                "doi": doi,
                "source": "Crossref",
                "title": str(title_values[0]).strip() if title_values else "",
                "year": str(year or ""),
                "journal": str(container_values[0]).strip() if container_values else "",
                "authors": authors,
            }
        )

    return results


def search_openalex(title: str) -> list[dict]:
    session = get_thread_session()

    params = {
        "search": title,
        "per-page": 3,
    }
    payload = safe_get_json(session, "https://api.openalex.org/works", params=params)
    if not payload:
        return []

    results: list[dict] = []

    for item in payload.get("results") or []:
        doi_url = item.get("doi") or ""
        doi = normalize_doi(doi_url)
        if not doi:
            continue

        authors = []
        for authorship in item.get("authorships") or []:
            author_name = ((authorship or {}).get("author") or {}).get("display_name") or ""
            if author_name:
                authors.append(str(author_name).strip())

        results.append(
            {
                "doi": doi,
                "source": "OpenAlex",
                "title": str(item.get("display_name") or "").strip(),
                "year": str(item.get("publication_year") or ""),
                "journal": str(((item.get("primary_location") or {}).get("source") or {}).get("display_name") or "").strip(),
                "authors": authors,
            }
        )

    return results


def search_datacite(title: str) -> list[dict]:
    session = get_thread_session()

    params = {
        "query": title,
        "page[size]": 3,
    }
    payload = safe_get_json(session, "https://api.datacite.org/dois", params=params)
    if not payload:
        return []

    results: list[dict] = []

    for item in payload.get("data") or []:
        attributes = item.get("attributes") or {}
        doi = normalize_doi(attributes.get("doi"))
        if not doi:
            continue

        titles = attributes.get("titles") or []
        creators = attributes.get("creators") or []
        publisher = str(attributes.get("publisher") or "").strip()
        year = str(attributes.get("publicationYear") or "").strip()

        authors = []
        for creator in creators:
            name = str((creator or {}).get("name") or "").strip()
            if name:
                authors.append(name)

        results.append(
            {
                "doi": doi,
                "source": "DataCite",
                "title": str((titles[0] or {}).get("title") or "").strip() if titles else "",
                "year": year,
                "journal": publisher,
                "authors": authors,
            }
        )

    return results


def search_europe_pmc(title: str) -> list[dict]:
    session = get_thread_session()

    params = {
        "query": title,
        "format": "json",
        "pageSize": 3,
    }
    payload = safe_get_json(
        session,
        "https://www.ebi.ac.uk/europepmc/webservices/rest/search",
        params=params,
    )
    if not payload:
        return []

    results: list[dict] = []

    for item in ((payload.get("resultList") or {}).get("result") or []):
        doi = normalize_doi(item.get("doi"))
        if not doi:
            continue

        author_string = str(item.get("authorString") or "").strip()
        authors = []
        if author_string:
            authors = [a.strip() for a in author_string.split(",") if a.strip()]

        results.append(
            {
                "doi": doi,
                "source": "Europe PMC",
                "title": str(item.get("title") or "").strip(),
                "year": str(item.get("pubYear") or "").strip(),
                "journal": str(item.get("journalTitle") or "").strip(),
                "authors": authors,
            }
        )

    return results


SOURCE_SEARCH_FUNCTIONS = {
    "Crossref": search_crossref,
    "OpenAlex": search_openalex,
    "DataCite": search_datacite,
    "Europe PMC": search_europe_pmc,
}


def search_named_source(source_name: str, title_query: str) -> list[dict]:
    cache_key = (source_name, title_query)
    if cache_key in _doi_query_cache:
        return list(_doi_query_cache[cache_key])

    search_fn = SOURCE_SEARCH_FUNCTIONS[source_name]
    results = list(search_fn(title_query))
    _doi_query_cache[cache_key] = results
    return list(results)


def search_source_group(title_query: str, source_names: Sequence[str]) -> list[dict]:
    raw_candidates: list[dict] = []
    for idx, source_name in enumerate(source_names):
        if is_stop_requested():
            break
        raw_candidates.extend(search_named_source(source_name, title_query))
        if idx < len(source_names) - 1 and DOI_SLEEP_BETWEEN_REQUESTS > 0:
            interruptible_sleep(DOI_SLEEP_BETWEEN_REQUESTS)
    return raw_candidates


def _collapse_candidates(
    raw_candidates: Sequence[dict],
    title: str,
    authors: str,
    publication_year: Any,
) -> list[Candidate]:
    source_year = parse_year(publication_year)
    candidates_by_doi: dict[str, Candidate] = {}
    sources_by_doi: dict[str, set[str]] = {}

    for item in raw_candidates:
        doi = normalize_doi(item.get("doi"))
        if not doi:
            continue

        matched_title = str(item.get("title") or "").strip()
        matched_year = str(item.get("year") or "").strip()
        matched_journal = str(item.get("journal") or "").strip()
        matched_authors_list = [str(a).strip() for a in item.get("authors") or [] if str(a).strip()]
        source_name = str(item.get("source") or "").strip()

        title_similarity = compute_title_similarity(title, matched_title)
        year_score = compute_year_score(source_year, parse_year(matched_year))
        author_score = compute_author_score(authors, matched_authors_list)
        confidence_score = compute_confidence_score(title_similarity, year_score, author_score)

        candidate = Candidate(
            doi=doi,
            doi_link=make_doi_url(doi),
            source=source_name,
            matched_title=matched_title,
            matched_year=matched_year,
            matched_journal=matched_journal,
            matched_authors=join_authors(matched_authors_list),
            source_count=1,
            title_similarity=title_similarity,
            year_score=year_score,
            author_score=author_score,
            confidence_score=confidence_score,
        )

        sources_by_doi.setdefault(doi, set())
        if source_name:
            sources_by_doi[doi].add(source_name)

        existing = candidates_by_doi.get(doi)
        if existing is None or candidate.confidence_score > existing.confidence_score:
            candidates_by_doi[doi] = candidate

    for doi, candidate in candidates_by_doi.items():
        source_names = sorted(sources_by_doi.get(doi, set()))
        candidate.source = "; ".join(source_names)
        candidate.source_count = len(source_names)

    candidates = list(candidates_by_doi.values())
    candidates.sort(
        key=lambda c: (
            -c.confidence_score,
            -c.source_count,
            -c.title_similarity,
            c.source.lower(),
            c.doi.lower(),
        )
    )
    return candidates


def is_early_stop_candidate(candidate: Optional[Candidate]) -> bool:
    if candidate is None:
        return False
    return (
        candidate.title_similarity >= DOI_EARLY_STOP_MIN_TITLE_SIMILARITY
        and candidate.confidence_score >= DOI_EARLY_STOP_MIN_CONFIDENCE_SCORE
        and candidate.source_count >= DOI_EARLY_STOP_MIN_SOURCE_COUNT
    )


def is_plausible_candidate(candidate: Optional[Candidate]) -> bool:
    if candidate is None:
        return False
    return (
        candidate.title_similarity >= DOI_NEEDS_REVIEW_MIN_TITLE_SIMILARITY
        and candidate.confidence_score >= DOI_NEEDS_REVIEW_MIN_CONFIDENCE_SCORE
    )


def gather_candidates(
    title: str,
    authors: str,
    publication_year: Any,
) -> list[Candidate]:
    raw_candidates: list[dict] = []
    query_variants = build_query_variants(title)

    for idx, query_variant in enumerate(query_variants):
        if is_stop_requested():
            break
        if DOI_ENABLE_SOURCE_CASCADE:
            raw_candidates.extend(search_source_group(query_variant, DOI_PRIMARY_SOURCE_ORDER))
            candidates = _collapse_candidates(raw_candidates, title, authors, publication_year)
            best = candidates[0] if candidates else None
            if is_early_stop_candidate(best):
                return candidates

            raw_candidates.extend(search_source_group(query_variant, DOI_SECONDARY_SOURCE_ORDER))
        else:
            source_order = tuple(DOI_PRIMARY_SOURCE_ORDER) + tuple(DOI_SECONDARY_SOURCE_ORDER)
            raw_candidates.extend(search_source_group(query_variant, source_order))

        candidates = _collapse_candidates(raw_candidates, title, authors, publication_year)
        best = candidates[0] if candidates else None
        if is_early_stop_candidate(best):
            return candidates
        if DOI_ACCEPT_FIRST_VARIANT_IF_PLAUSIBLE and idx == 0 and is_plausible_candidate(best):
            return candidates

        if idx < len(query_variants) - 1 and DOI_SLEEP_BETWEEN_REQUESTS > 0:
            interruptible_sleep(DOI_SLEEP_BETWEEN_REQUESTS)

    return _collapse_candidates(raw_candidates, title, authors, publication_year)


def classify_candidate(best: Optional[Candidate], second: Optional[Candidate]) -> tuple[str, str]:
    if best is None:
        return "unresolved", "No candidates found."

    top2_gap = (
        max(0.0, best.confidence_score - second.confidence_score)
        if second is not None
        else 0.0
    )
    metadata_supported = (
        best.author_score >= DOI_AUTO_APPLY_MIN_AUTHOR_SCORE
        or best.year_score >= DOI_AUTO_APPLY_MIN_YEAR_SCORE
    )
    strong_title = best.title_similarity >= DOI_AUTO_APPLY_MIN_TITLE_SIMILARITY
    strong_confidence = best.confidence_score >= DOI_AUTO_APPLY_MIN_CONFIDENCE_SCORE
    corroborated = best.source_count >= DOI_AUTO_APPLY_MIN_SOURCE_COUNT
    exceptional_title_match = best.title_similarity >= 0.97
    competing_candidate_close = second is not None and top2_gap < DOI_AUTO_APPLY_MIN_TOP2_GAP

    if (
        strong_title
        and strong_confidence
        and metadata_supported
        and not competing_candidate_close
        and (corroborated or exceptional_title_match)
    ):
        return (
            "enriched",
            (
                "Auto-applied: strong title match, metadata support, "
                f"source_count={best.source_count}, confidence_score={best.confidence_score:.2f}"
            ),
        )

    if (
        best.title_similarity >= DOI_NEEDS_REVIEW_MIN_TITLE_SIMILARITY
        and best.confidence_score >= DOI_NEEDS_REVIEW_MIN_CONFIDENCE_SCORE
    ):
        if competing_candidate_close:
            return "needs_review", f"Top candidates are too close (top2_gap={top2_gap:.2f})."
        if not metadata_supported:
            return "needs_review", "Title match is promising but author/year support is weak."
        if not corroborated:
            return "needs_review", "Only one source returned this DOI."
        return "needs_review", "Plausible DOI found but auto-apply safeguards were not met."

    return (
        "unresolved",
        (
            "Below safe thresholds: "
            f"title_similarity={best.title_similarity:.2f}, "
            f"confidence_score={best.confidence_score:.2f}"
        ),
    )


# =========================
# WORKSHEET ANALYSIS
# =========================

def analyze_phase0_scope(ws, col_map: Dict[str, int]) -> dict:
    total_rows_read = 0
    total_existing_doi = 0
    total_missing_doi = 0
    total_eligible = 0
    total_skipped_missing_title = 0
    total_skipped_already_downloaded = 0
    candidate_row_indices: list[int] = []

    for row_idx in range(2, ws.max_row + 1):
        record_id = str(get_optional_cell(ws, row_idx, col_map, "record_id", "") or "").strip()
        if not record_id:
            continue

        if ENABLE_BENCHMARK_MODE and BENCHMARK_RECORD_IDS and record_id not in BENCHMARK_RECORD_IDS:
            continue

        total_rows_read += 1

        if row_has_usable_doi(ws, row_idx, col_map) and not DOI_OVERWRITE_EXISTING:
            total_existing_doi += 1
            continue

        total_missing_doi += 1

        title = str(get_optional_cell(ws, row_idx, col_map, "Title", "") or "").strip()
        if not title:
            total_skipped_missing_title += 1
            continue

        if row_is_already_downloaded(ws, row_idx, col_map):
            total_skipped_already_downloaded += 1
            continue

        candidate_row_indices.append(row_idx)
        total_eligible += 1

        if DOI_MAX_ROWS_TO_PROCESS is not None and total_eligible >= DOI_MAX_ROWS_TO_PROCESS:
            break

    return {
        "total_rows_read": total_rows_read,
        "total_existing_doi": total_existing_doi,
        "total_missing_doi": total_missing_doi,
        "total_eligible": total_eligible,
        "total_skipped_missing_title": total_skipped_missing_title,
        "total_skipped_already_downloaded": total_skipped_already_downloaded,
        "candidate_row_indices": candidate_row_indices,
    }


# =========================
# REPORT FORMATTING
# =========================

def autosize_columns(ws, extra_padding: int = 2, max_width: int = 80) -> None:
    for col_idx in range(1, ws.max_column + 1):
        col_letter = get_column_letter(col_idx)
        max_len = 0

        for row_idx in range(1, ws.max_row + 1):
            value = ws.cell(row=row_idx, column=col_idx).value
            value_str = "" if value is None else str(value)
            if len(value_str) > max_len:
                max_len = len(value_str)

        ws.column_dimensions[col_letter].width = min(max(max_len + extra_padding, 10), max_width)


def style_title_green(ws, start_cell: str, end_cell: str, text: str) -> None:
    ws[start_cell] = text
    ws[start_cell].font = Font(bold=True, size=14, color=WHITE_FONT)
    ws[start_cell].fill = PatternFill(fill_type="solid", fgColor=GREEN_TITLE_FILL)
    ws[start_cell].alignment = CENTER
    ws.merge_cells(f"{start_cell}:{end_cell}")


def style_dark_header_row(ws, row_num: int, headers: list[str], start_col: int = 1) -> None:
    fill = PatternFill(fill_type="solid", fgColor=DARK_HEADER_FILL)
    font = Font(color=WHITE_FONT, bold=True)

    for idx, header in enumerate(headers, start=start_col):
        cell = ws.cell(row=row_num, column=idx, value=header)
        cell.fill = fill
        cell.font = font
        cell.alignment = CENTER


def style_light_header_row(ws, row_num: int, headers: list[str], start_col: int = 1) -> None:
    fill = PatternFill(fill_type="solid", fgColor=LIGHT_BLUE_HEADER_FILL)
    font = Font(bold=True)

    for idx, header in enumerate(headers, start=start_col):
        cell = ws.cell(row=row_num, column=idx, value=header)
        cell.fill = fill
        cell.font = font
        cell.alignment = CENTER_TOP


def format_table_header(ws) -> None:
    for col_idx in range(1, ws.max_column + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = Font(bold=True)
        cell.alignment = CENTER_TOP


def build_phase0_report_headers() -> list[str]:
    return [
        "record_id",
        "Authors",
        "Publication Year",
        "Title",
        "original_doi",
        "phase0_status",
        "matched_doi",
        "confidence_score",
        "title_similarity",
        "year_score",
        "author_score",
        "match_source_count",
        "match_source",
        "matched_title",
        "matched_year",
        "matched_journal",
        "second_candidate_doi",
        "second_candidate_title",
        "second_candidate_confidence",
        "top2_gap",
        "decision_reason",
    ]


def build_phase0_report_headers_enriched() -> list[str]:
    return [
        "record_id",
        "Authors",
        "Publication Year",
        "Title",
        "matched_doi",
        "confidence_score",
        "title_similarity",
        "year_score",
        "author_score",
        "match_source_count",
        "match_source",
        "matched_title",
        "matched_year",
        "matched_journal",
        "decision_reason",
    ]


def build_phase0_report_headers_needs_review() -> list[str]:
    return [
        "record_id",
        "Authors",
        "Publication Year",
        "Title",
        "matched_doi",
        "confidence_score",
        "title_similarity",
        "year_score",
        "author_score",
        "match_source_count",
        "match_source",
        "matched_title",
        "matched_year",
        "matched_journal",
        "second_candidate_doi",
        "second_candidate_title",
        "second_candidate_confidence",
        "top2_gap",
        "decision_reason",
    ]


def build_phase0_report_headers_unresolved() -> list[str]:
    return [
        "record_id",
        "Authors",
        "Publication Year",
        "Title",
        "matched_doi",
        "confidence_score",
        "title_similarity",
        "year_score",
        "author_score",
        "match_source_count",
        "match_source",
        "matched_title",
        "matched_year",
        "matched_journal",
        "decision_reason",
    ]


def row_result_to_values(row: Phase0RowResult) -> list[Any]:
    return [
        row.record_id,
        row.authors,
        row.publication_year,
        row.title,
        row.original_doi,
        row.phase0_status,
        row.matched_doi,
        row.confidence_score,
        row.title_similarity,
        row.year_score,
        row.author_score,
        row.match_source_count,
        row.match_source,
        row.matched_title,
        row.matched_year,
        row.matched_journal,
        row.second_candidate_doi,
        row.second_candidate_title,
        row.second_candidate_confidence,
        row.top2_gap,
        row.decision_reason,
    ]


def row_result_to_values_for_headers(row: Phase0RowResult, headers: Sequence[str]) -> list[Any]:
    value_map = {
        "record_id": row.record_id,
        "Authors": row.authors,
        "Publication Year": row.publication_year,
        "Title": row.title,
        "original_doi": row.original_doi,
        "phase0_status": row.phase0_status,
        "matched_doi": row.matched_doi,
        "confidence_score": row.confidence_score,
        "title_similarity": row.title_similarity,
        "year_score": row.year_score,
        "author_score": row.author_score,
        "match_source_count": row.match_source_count,
        "match_source": row.match_source,
        "matched_title": row.matched_title,
        "matched_year": row.matched_year,
        "matched_journal": row.matched_journal,
        "second_candidate_doi": row.second_candidate_doi,
        "second_candidate_title": row.second_candidate_title,
        "second_candidate_confidence": row.second_candidate_confidence,
        "top2_gap": row.top2_gap,
        "decision_reason": row.decision_reason,
    }
    return [value_map.get(header, "") for header in headers]


def write_table_sheet(ws, rows: list[Phase0RowResult]) -> None:
    headers = build_phase0_report_headers()
    write_table_sheet_with_headers(ws, rows, headers)


def write_table_sheet_with_headers(
    ws,
    rows: list[Phase0RowResult],
    headers: Sequence[str],
) -> None:

    for col_idx, header in enumerate(headers, start=1):
        ws.cell(row=1, column=col_idx, value=header)

    format_table_header(ws)

    current_row = 2
    for row_result in rows:
        values = row_result_to_values_for_headers(row_result, headers)

        for col_idx, value in enumerate(values, start=1):
            cell = ws.cell(row=current_row, column=col_idx, value=value)

            if headers[col_idx - 1] in {
                "confidence_score",
                "title_similarity",
                "year_score",
                "author_score",
                "second_candidate_confidence",
                "top2_gap",
            }:
                cell.number_format = "0.00"

        current_row += 1

    autosize_columns(ws)


def create_phase0_stats_sheet(
    wb: Workbook,
    scope_info: dict,
    processed_rows: list[Phase0RowResult],
    report_generated: bool,
) -> None:
    ws = wb.create_sheet(DOI_REPORT_STATS_SHEET_NAME)

    style_title_green(ws, "A1", "D1", "DOI enrichment - summary")

    ws["A2"] = "Sheet analisada"
    ws["B2"] = "wos_full_text"
    ws["A3"] = "Report gerado"
    ws["B3"] = "Sim" if report_generated else "Não"

    style_dark_header_row(ws, 5, ["Métrica", "Valor"], start_col=1)

    enriched = sum(1 for row in processed_rows if row.phase0_status == "enriched")
    unresolved = sum(1 for row in processed_rows if row.phase0_status == "unresolved")
    search_errors = sum(1 for row in processed_rows if row.phase0_status == "search_error")

    stats_rows = [
        ("Total rows lidas", scope_info["total_rows_read"]),
        ("Com DOI já existente", scope_info["total_existing_doi"]),
        ("Sem DOI detetados", scope_info["total_missing_doi"]),
        ("Elegíveis para procurar", scope_info["total_eligible"]),
        ("Ignorados sem título", scope_info["total_skipped_missing_title"]),
        ("Ignorados já com PDF", scope_info["total_skipped_already_downloaded"]),
        ("Trabalhados na fase 0", len(processed_rows)),
        ("DOIs enriquecidos", enriched),
        ("Não resolvidos", unresolved),
        ("Erros de pesquisa", search_errors),
    ]

    row_idx = 6
    for label, value in stats_rows:
        ws.cell(row=row_idx, column=1, value=label)
        ws.cell(row=row_idx, column=2, value=value)
        ws.cell(row=row_idx, column=1).alignment = CENTER_TOP
        ws.cell(row=row_idx, column=2).alignment = CENTER
        row_idx += 1

    ws.column_dimensions["A"].width = 32
    ws.column_dimensions["B"].width = 14


def create_phase0_stats_sheet_safe(
    wb: Workbook,
    scope_info: dict,
    processed_rows: list[Phase0RowResult],
    report_generated: bool,
) -> None:
    ws = wb.create_sheet(DOI_REPORT_STATS_SHEET_NAME)

    style_title_green(ws, "A1", "D1", "DOI enrichment - summary")

    ws["A2"] = "Sheet analyzed"
    ws["B2"] = "wos_full_text"
    ws["A3"] = "Report generated"
    ws["B3"] = "Yes" if report_generated else "No"

    style_dark_header_row(ws, 5, ["Metric", "Value"], start_col=1)

    enriched = sum(1 for row in processed_rows if row.phase0_status == "enriched")
    needs_review = sum(1 for row in processed_rows if row.phase0_status == "needs_review")
    unresolved = sum(1 for row in processed_rows if row.phase0_status == "unresolved")
    search_errors = sum(1 for row in processed_rows if row.phase0_status == "search_error")
    eligible = scope_info["total_eligible"]
    auto_applied_rate = (enriched / eligible * 100.0) if eligible > 0 else 0.0
    potential_match_rate = ((enriched + needs_review) / eligible * 100.0) if eligible > 0 else 0.0

    stats_rows = [
        ("Total rows read", scope_info["total_rows_read"]),
        ("Existing DOI", scope_info["total_existing_doi"]),
        ("Missing DOI", scope_info["total_missing_doi"]),
        ("Eligible for lookup", scope_info["total_eligible"]),
        ("Skipped without title", scope_info["total_skipped_missing_title"]),
        ("Skipped already with PDF", scope_info["total_skipped_already_downloaded"]),
        ("Processed in phase 0", len(processed_rows)),
        ("Enriched DOIs", enriched),
        ("Needs review", needs_review),
        ("Unresolved", unresolved),
        ("Search errors", search_errors),
        ("Auto-applied rate (%)", round(auto_applied_rate, 2)),
        ("Potential match rate (%)", round(potential_match_rate, 2)),
    ]

    row_idx = 6
    for label, value in stats_rows:
        ws.cell(row=row_idx, column=1, value=label)
        ws.cell(row=row_idx, column=2, value=value)
        ws.cell(row=row_idx, column=1).alignment = CENTER_TOP
        ws.cell(row=row_idx, column=2).alignment = CENTER
        row_idx += 1

    ws.column_dimensions["A"].width = 32
    ws.column_dimensions["B"].width = 14


def create_phase0_report(
    processed_rows: list[Phase0RowResult],
    scope_info: dict,
    output_path,
) -> bool:
    wb = Workbook()
    default_ws = wb.active
    wb.remove(default_ws)

    create_phase0_stats_sheet_safe(
        wb=wb,
        scope_info=scope_info,
        processed_rows=processed_rows,
        report_generated=True,
    )

    all_processed_ws = wb.create_sheet(DOI_REPORT_ALL_PROCESSED_SHEET_NAME)
    write_table_sheet(all_processed_ws, processed_rows)

    enriched_rows = [row for row in processed_rows if row.phase0_status == "enriched"]
    enriched_ws = wb.create_sheet(DOI_REPORT_ENRICHED_SHEET_NAME)
    write_table_sheet_with_headers(
        enriched_ws,
        enriched_rows,
        build_phase0_report_headers_enriched(),
    )

    needs_review_rows = [row for row in processed_rows if row.phase0_status == "needs_review"]
    needs_review_ws = wb.create_sheet(DOI_REPORT_NEEDS_REVIEW_SHEET_NAME)
    write_table_sheet_with_headers(
        needs_review_ws,
        needs_review_rows,
        build_phase0_report_headers_needs_review(),
    )

    unresolved_rows = [
        row for row in processed_rows
        if row.phase0_status in {"unresolved", "search_error"}
    ]
    unresolved_ws = wb.create_sheet(DOI_REPORT_UNRESOLVED_SHEET_NAME)
    write_table_sheet_with_headers(
        unresolved_ws,
        unresolved_rows,
        build_phase0_report_headers_unresolved(),
    )

    return safe_save_workbook(wb, output_path)


# =========================
# CACHE HELPERS
# =========================

def _cache_phase0_run_data(
    session_state: SessionState,
    processed_rows: list[Phase0RowResult],
    scope_info: dict,
) -> None:
    session_state.phase0_last_processed_rows = processed_rows
    session_state.phase0_last_scope_info = scope_info


def _get_cached_phase0_run_data(
    session_state: SessionState,
) -> tuple[list[Phase0RowResult], dict]:
    processed_rows = getattr(session_state, "phase0_last_processed_rows", None)
    scope_info = getattr(session_state, "phase0_last_scope_info", None)

    if processed_rows is None or scope_info is None:
        raise ValueError("No cached phase 0 results available for report generation.")

    return processed_rows, scope_info


def generate_phase0_report_from_cache(session_state: SessionState) -> Optional[str]:
    processed_rows, scope_info = _get_cached_phase0_run_data(session_state)

    report_path = build_report_output_path(
        REPORT_NAME_FASE0_DOI,
        timestamp=session_state.timestamp or None,
    )

    ok = create_phase0_report(
        processed_rows=processed_rows,
        scope_info=scope_info,
        output_path=report_path,
    )
    if not ok:
        raise RuntimeError(f"Could not save phase 0 report: {report_path}")

    session_state.add_generated_report(report_path)
    session_state.phase0_summary.report_path = report_path
    return str(report_path)


# =========================
# MAIN PHASE 0 ENTRYPOINT
# =========================

def run_phase0_doi_enrichment_for_session(
    session_state: SessionState,
    reporter: ReporterType = None,
    generate_report: bool = False,
    stop_event=None,
):
    if not session_state.is_workbook_loaded:
        raise ValueError("SessionState workbook is not loaded.")

    set_current_stop_event(stop_event)

    session_state.current_phase = "phase0"

    ws = session_state.worksheet
    col_map = session_state.col_map
    validate_required_columns(col_map, session_state.sheet_name, DOI_REQUIRED_INPUT_COLUMNS)

    start_time = time.time()
    emit(
        reporter,
        "phase0_start",
        {
            "processed": 0,
            "total": 0,
            "start_time": start_time,
            "total_rows_read": 0,
            "total_existing_doi": 0,
            "total_missing_doi": 0,
            "total_eligible": 0,
            "total_skipped_missing_title": 0,
            "total_skipped_already_downloaded": 0,
            "total_enriched": 0,
            "total_needs_review": 0,
            "total_unresolved": 0,
            "total_errors": 0,
            "last_record_id": "",
            "last_title": "",
            "last_status": "preparing workbook and lookup queue",
            "last_doi": "",
            "last_confidence": "",
        },
    )

    scope_info = analyze_phase0_scope(ws, col_map)
    candidate_row_indices: list[int] = scope_info["candidate_row_indices"]
    total_to_process = len(candidate_row_indices)
    processed = 0
    enriched = 0
    needs_review = 0
    unresolved = 0
    errors = 0

    processed_rows: list[Phase0RowResult] = []

    emit(
        reporter,
        "phase0_start",
        {
            "processed": 0,
            "total": total_to_process,
            "start_time": start_time,
            "total_rows_read": scope_info["total_rows_read"],
            "total_existing_doi": scope_info["total_existing_doi"],
            "total_missing_doi": scope_info["total_missing_doi"],
            "total_eligible": scope_info["total_eligible"],
            "total_skipped_missing_title": scope_info["total_skipped_missing_title"],
            "total_skipped_already_downloaded": scope_info["total_skipped_already_downloaded"],
            "total_enriched": 0,
            "total_needs_review": 0,
            "total_unresolved": 0,
            "total_errors": 0,
            "last_record_id": "",
            "last_title": "",
            "last_status": "starting",
            "last_doi": "",
            "last_confidence": "",
        },
    )

    for row_idx in candidate_row_indices:
        if stop_event is not None and stop_event.is_set():
            break

        record_id = str(get_optional_cell(ws, row_idx, col_map, "record_id", "") or "").strip()
        title = str(get_optional_cell(ws, row_idx, col_map, "Title", "") or "").strip()
        authors = str(get_optional_cell(ws, row_idx, col_map, "Authors", "") or "").strip()
        publication_year = str(get_optional_cell(ws, row_idx, col_map, "Publication Year", "") or "").strip()
        original_doi = str(get_optional_cell(ws, row_idx, col_map, "DOI", "") or "").strip()
        original_doi_link = str(get_optional_cell(ws, row_idx, col_map, "DOI Link", "") or "").strip()

        last_doi = ""
        last_confidence = ""
        last_status = "unresolved"

        try:
            candidates = gather_candidates(title, authors, publication_year)
            best = candidates[0] if candidates else None
            second = candidates[1] if len(candidates) > 1 else None
            top2_gap = (
                max(0.0, best.confidence_score - second.confidence_score)
                if best is not None and second is not None
                else 0.0
            )
            phase0_status, decision_reason = classify_candidate(best, second)

            if phase0_status == "enriched" and best is not None:
                write_doi_enrichment_result(
                    ws=ws,
                    row_idx=row_idx,
                    col_map=col_map,
                    doi=best.doi,
                    doi_link=best.doi_link,
                )
                enriched += 1
                last_status = phase0_status
                last_doi = best.doi
                last_confidence = f"{best.confidence_score:.2f}"
            elif phase0_status == "needs_review":
                needs_review += 1
                last_status = phase0_status
                if best is not None:
                    last_doi = best.doi
                    last_confidence = f"{best.confidence_score:.2f}"
            else:
                unresolved += 1
                last_status = phase0_status
                if best is not None:
                    last_doi = best.doi
                    last_confidence = f"{best.confidence_score:.2f}"

            processed_rows.append(
                Phase0RowResult(
                    record_id=record_id,
                    title=title,
                    authors=authors,
                    publication_year=publication_year,
                    original_doi=original_doi,
                    original_doi_link=original_doi_link,
                    phase0_status=phase0_status,
                    matched_doi=best.doi if best is not None else "",
                    matched_doi_link=best.doi_link if best is not None else "",
                    confidence_score=best.confidence_score if best is not None else 0.0,
                    title_similarity=best.title_similarity if best is not None else 0.0,
                    year_score=best.year_score if best is not None else 0.0,
                    author_score=best.author_score if best is not None else 0.0,
                    match_source_count=best.source_count if best is not None else 0,
                    match_source=best.source if best is not None else "",
                    matched_title=best.matched_title if best is not None else "",
                    matched_year=best.matched_year if best is not None else "",
                    matched_journal=best.matched_journal if best is not None else "",
                    second_candidate_doi=(
                        ""
                        if phase0_status == "enriched"
                        else (second.doi if second is not None else "")
                    ),
                    second_candidate_title=(
                        ""
                        if phase0_status == "enriched"
                        else (second.matched_title if second is not None else "")
                    ),
                    second_candidate_confidence=(
                        0.0
                        if phase0_status == "enriched"
                        else (second.confidence_score if second is not None else 0.0)
                    ),
                    top2_gap=top2_gap,
                    decision_reason=decision_reason,
                )
            )

        except Exception as exc:
            errors += 1
            last_status = "search_error"

            processed_rows.append(
                Phase0RowResult(
                    record_id=record_id,
                    title=title,
                    authors=authors,
                    publication_year=publication_year,
                    original_doi=original_doi,
                    original_doi_link=original_doi_link,
                    phase0_status="search_error",
                    matched_doi="",
                    matched_doi_link="",
                    confidence_score=0.0,
                    title_similarity=0.0,
                    year_score=0.0,
                    author_score=0.0,
                    match_source_count=0,
                    match_source="",
                    matched_title="",
                    matched_year="",
                    matched_journal="",
                    second_candidate_doi="",
                    second_candidate_title="",
                    second_candidate_confidence=0.0,
                    top2_gap=0.0,
                    decision_reason=f"Search error: {exc}",
                )
            )

        processed += 1

        emit(
            reporter,
            "phase0_update",
            {
                "processed": processed,
                "total": total_to_process,
                "start_time": start_time,
                "total_rows_read": scope_info["total_rows_read"],
                "total_existing_doi": scope_info["total_existing_doi"],
                "total_missing_doi": scope_info["total_missing_doi"],
                "total_eligible": scope_info["total_eligible"],
                "total_skipped_missing_title": scope_info["total_skipped_missing_title"],
                "total_skipped_already_downloaded": scope_info["total_skipped_already_downloaded"],
                "total_enriched": enriched,
                "total_needs_review": needs_review,
                "total_unresolved": unresolved,
                "total_errors": errors,
                "last_record_id": record_id,
                "last_title": title,
                "last_status": last_status,
                "last_doi": last_doi,
                "last_confidence": last_confidence,
            },
        )

        if DOI_SLEEP_BETWEEN_REQUESTS > 0:
            interruptible_sleep(DOI_SLEEP_BETWEEN_REQUESTS)

    _cache_phase0_run_data(
        session_state=session_state,
        processed_rows=processed_rows,
        scope_info=scope_info,
    )

    report_path = None
    effective_generate_report = generate_report or DOI_GENERATE_REPORT_BY_DEFAULT
    if effective_generate_report:
        report_path_str = generate_phase0_report_from_cache(session_state)
        report_path = report_path_str

    refresh_session_records(session_state)

    session_state.phase0_summary = Phase0Summary(
        total_rows_read=scope_info["total_rows_read"],
        total_existing_doi=scope_info["total_existing_doi"],
        total_missing_doi=scope_info["total_missing_doi"],
        total_eligible=scope_info["total_eligible"],
        total_skipped_missing_title=scope_info["total_skipped_missing_title"],
        total_skipped_already_downloaded=scope_info["total_skipped_already_downloaded"],
        total_enriched=enriched,
        total_needs_review=needs_review,
        total_unresolved=unresolved,
        total_errors=errors,
        report_path=report_path,
    )

    emit(
        reporter,
        "phase0_end",
        {
            "processed": processed,
            "total": total_to_process,
            "start_time": start_time,
            "total_rows_read": scope_info["total_rows_read"],
            "total_existing_doi": scope_info["total_existing_doi"],
            "total_missing_doi": scope_info["total_missing_doi"],
            "total_eligible": scope_info["total_eligible"],
            "total_skipped_missing_title": scope_info["total_skipped_missing_title"],
            "total_skipped_already_downloaded": scope_info["total_skipped_already_downloaded"],
            "total_enriched": enriched,
            "total_needs_review": needs_review,
            "total_unresolved": unresolved,
            "total_errors": errors,
            "last_record_id": "",
            "last_title": "",
            "last_status": "finished",
            "last_doi": "",
            "last_confidence": "",
            "report_path": str(report_path) if report_path else "",
        },
    )

    return {
        "processed": processed,
        "total": total_to_process,
        "total_enriched": enriched,
        "total_needs_review": needs_review,
        "total_unresolved": unresolved,
        "total_errors": errors,
        "report_path": str(report_path) if report_path else "",
    }
