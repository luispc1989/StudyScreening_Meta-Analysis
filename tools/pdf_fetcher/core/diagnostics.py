from __future__ import annotations

import re
import threading
import time
import unicodedata
from collections import Counter, defaultdict
from concurrent.futures import FIRST_COMPLETED, ThreadPoolExecutor, wait
from typing import Callable, Dict, List, Optional, Tuple
from urllib.parse import urljoin, urlparse

import requests
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from requests.adapters import HTTPAdapter
from urllib3.util.retry import Retry

from tools.pdf_fetcher.core.config import (
    DIAGNOSTIC_ERROR_TAB_COLUMNS,
    DIAGNOSTIC_EXCLUDED_FROM_ERROR_RANKING,
    DIAGNOSTIC_MAX_ROWS_TO_PROCESS,
    DIAGNOSTIC_MAX_WORKERS,
    DIAGNOSTIC_OUTPUT_COLUMNS,
    DIAGNOSTIC_OUTPUT_SHEET_NAME,
    DIAGNOSTIC_PROGRESS_REFRESH_EVERY_N,
    DIAGNOSTIC_PROGRESS_REFRESH_EVERY_SECONDS,
    DIAGNOSTIC_REQUEST_TIMEOUT,
    DIAGNOSTIC_REQUIRED_INPUT_COLUMNS_INITIAL,
    DIAGNOSTIC_REQUIRED_INPUT_COLUMNS_SESSION,
    DIAGNOSTIC_STATS_SHEET_NAME,
    HEADERS,
    REPORT_NAME_DIAGNOSTICO_FASE0,
    REPORT_NAME_DIAGNOSTICO_FASE1,
    REPORT_NAME_DIAGNOSTICO_FASE2,
    REPORT_NAME_DIAGNOSTICO_INICIAL,
    SHEET_NAME,
    build_report_output_path,
)
from tools.pdf_fetcher.core.excel_io import build_col_map, get_optional_cell, safe_save_workbook
from tools.pdf_fetcher.core.excel_io import reconcile_pdf_state_with_filesystem
from tools.pdf_fetcher.core.models import DiagnosticSummary, SessionState


ReporterType = Optional[Callable[[str, dict], None]]
StopEventType = Optional[threading.Event]

_thread_local = threading.local()


# =========================
# STYLE CONSTANTS
# =========================

GREEN_TITLE_FILL = "FF1F6E43"
DARK_HEADER_FILL = "FF000000"
BLUE_SECTION_FILL = "FF1F4E78"
LIGHT_BLUE_HEADER_FILL = "FFD9EAF7"
WHITE_FONT = "FFFFFFFF"

CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True, shrink_to_fit=False)
CENTER_TOP = Alignment(horizontal="center", vertical="top", wrap_text=True, shrink_to_fit=False)
LEFT_TOP = Alignment(horizontal="left", vertical="top", wrap_text=True, shrink_to_fit=False)


# =========================
# GENERIC HELPERS
# =========================

def emit(reporter: ReporterType, event_name: str, payload: dict) -> None:
    if reporter is None:
        return
    reporter(event_name, payload)


def is_blank(value) -> bool:
    return value is None or str(value).strip() == ""


def normalize_text(value: str) -> str:
    value = str(value or "").strip()
    value = unicodedata.normalize("NFKD", value)
    value = "".join(ch for ch in value if not unicodedata.combining(ch))
    return value.strip().lower()


def normalize_label(value: str) -> str:
    return normalize_text(value)


def normalize_doi(doi: str) -> Optional[str]:
    if is_blank(doi):
        return None

    doi = str(doi).strip()
    doi = re.sub(r"^https?://(dx\.)?doi\.org/", "", doi, flags=re.IGNORECASE)
    doi = doi.strip()

    if re.match(r"^10\.\d{4,9}/\S+$", doi):
        return doi

    return None


def make_doi_url(doi: str) -> str:
    return f"https://doi.org/{doi}"


def extract_pdf_url_from_html(base_url: str, html: str) -> Optional[str]:
    if not html:
        return None

    patterns = [
        r'<meta[^>]+name=["\']citation_pdf_url["\'][^>]+content=["\']([^"\']+)["\']',
        r'<meta[^>]+content=["\']([^"\']+)["\'][^>]+name=["\']citation_pdf_url["\']',
        r'href=["\']([^"\']+\.pdf(?:\?[^"\']*)?)["\']',
        r'src=["\']([^"\']+\.pdf(?:\?[^"\']*)?)["\']',
    ]

    for pattern in patterns:
        match = re.search(pattern, html, flags=re.IGNORECASE)
        if match:
            return urljoin(base_url, match.group(1))

    return None


def normalize_host(url: str) -> str:
    if is_blank(url):
        return "(missing)"

    try:
        parsed = urlparse(str(url).strip())
        host = (parsed.netloc or "").strip().lower()
        if ":" in host:
            host = host.split(":")[0]
        if host.startswith("www."):
            host = host[4:]
        return host if host else "(missing)"
    except Exception:
        return "(missing)"


def make_safe_sheet_name(name: str) -> str:
    name = re.sub(r'[:\\/?*\[\]]', "_", str(name).strip())
    return name[:31] if len(name) > 31 else name


def response_looks_like_pdf(resp: requests.Response) -> bool:
    content_type = (resp.headers.get("Content-Type") or "").lower()
    if "application/pdf" in content_type:
        return True

    try:
        return resp.content[:5] == b"%PDF-"
    except Exception:
        return False


def classify_html_failure(status_code: Optional[int], html: str, final_url: str = "") -> str:
    text = (html or "").lower()

    security_markers = [
        "captcha",
        "cf-chl",
        "cloudflare",
        "attention required",
        "verify you are human",
        "security check",
        "bot check",
        "checking if the site connection is secure",
        "recaptcha",
        "hcaptcha",
    ]
    if any(marker in text for marker in security_markers):
        return "captcha_or_security_check"

    if status_code == 403:
        if any(x in text for x in ["captcha", "security", "human", "cloudflare"]):
            return "captcha_or_security_check"
        return "paywalled_institutional_access"

    purchase_markers = [
        "purchase pdf",
        "buy article",
        "buy now",
        "purchase access",
        "rent this article",
        "add to cart",
        "purchase this article",
    ]
    if any(marker in text for marker in purchase_markers):
        return "paywalled_purchase_required"

    institutional_markers = [
        "access through your institution",
        "institutional access",
        "login via institution",
        "shibboleth",
        "sign in through your institution",
        "institutional login",
        "check access",
        "get access",
        "subscribe to journal",
        "subscription required",
    ]
    if any(marker in text for marker in institutional_markers):
        return "paywalled_institutional_access"

    if status_code == 404:
        return "not_found"

    if status_code is not None and status_code >= 500:
        return "broken_link"

    metadata_markers = [
        "abstract",
        "references",
        "article info",
        "metrics",
        "supplementary",
        "full text",
        "author information",
    ]
    if any(marker in text for marker in metadata_markers):
        return "manual_check"

    return "manual_check"


def build_retry_strategy() -> Retry:
    return Retry(
        total=1,
        connect=1,
        read=1,
        redirect=2,
        status=1,
        status_forcelist=[429, 500, 502, 503, 504],
        allowed_methods=frozenset(["GET", "HEAD"]),
        backoff_factor=0.3,
        raise_on_status=False,
        respect_retry_after_header=True,
    )


def get_thread_session() -> requests.Session:
    session = getattr(_thread_local, "session", None)
    if session is not None:
        return session

    session = requests.Session()
    retry_strategy = build_retry_strategy()

    adapter = HTTPAdapter(
        max_retries=retry_strategy,
        pool_connections=100,
        pool_maxsize=100,
    )
    session.mount("http://", adapter)
    session.mount("https://", adapter)

    merged_headers = dict(HEADERS)
    merged_headers.setdefault("Accept-Language", "en-US,en;q=0.9")
    merged_headers.setdefault("Connection", "keep-alive")
    session.headers.update(merged_headers)

    _thread_local.session = session
    return session


def try_download_diagnostic(
    url: str,
    session: requests.Session,
) -> Tuple[str, Optional[int], str]:
    """
    Returns:
        (status_label, http_status, final_url_used)
    """
    try:
        resp = session.get(
            url,
            timeout=DIAGNOSTIC_REQUEST_TIMEOUT,
            allow_redirects=True,
            stream=False,
        )
        http_status = resp.status_code
        final_url = resp.url

        if http_status >= 400:
            if http_status == 404:
                return "not_found", http_status, final_url
            if http_status == 403:
                html = ""
                try:
                    html = resp.text or ""
                except Exception:
                    pass
                return classify_html_failure(http_status, html, final_url), http_status, final_url
            return "broken_link", http_status, final_url

        if response_looks_like_pdf(resp):
            return "downloaded", http_status, final_url

        html = ""
        try:
            html = resp.text or ""
        except Exception:
            pass

        pdf_url = extract_pdf_url_from_html(final_url, html)

        if pdf_url:
            pdf_resp = session.get(
                pdf_url,
                timeout=DIAGNOSTIC_REQUEST_TIMEOUT,
                allow_redirects=True,
                stream=False,
            )
            pdf_http_status = pdf_resp.status_code
            pdf_final_url = pdf_resp.url

            if pdf_http_status >= 400:
                if pdf_http_status == 404:
                    return "not_found", pdf_http_status, pdf_final_url
                if pdf_http_status == 403:
                    pdf_html = ""
                    try:
                        pdf_html = pdf_resp.text or ""
                    except Exception:
                        pass
                    return classify_html_failure(pdf_http_status, pdf_html, pdf_final_url), pdf_http_status, pdf_final_url
                return "broken_link", pdf_http_status, pdf_final_url

            if response_looks_like_pdf(pdf_resp):
                return "downloaded", pdf_http_status, pdf_final_url

            pdf_html = ""
            try:
                pdf_html = pdf_resp.text or ""
            except Exception:
                pass
            return classify_html_failure(pdf_http_status, pdf_html, pdf_final_url), pdf_http_status, pdf_final_url

        return classify_html_failure(http_status, html, final_url), http_status, final_url

    except requests.Timeout:
        return "broken_link", None, url
    except requests.RequestException:
        return "broken_link", None, url


def autosize_columns(ws, extra_padding: int = 2, max_width: int = 80) -> None:
    for col_idx in range(1, ws.max_column + 1):
        col_letter = get_column_letter(col_idx)
        max_len = 0

        for row_idx in range(1, ws.max_row + 1):
            value = ws.cell(row=row_idx, column=col_idx).value
            value_str = "" if value is None else str(value)
            if len(value_str) > max_len:
                max_len = len(value_str)

        ws.column_dimensions[col_letter].width = min(max(max_len + extra_padding, 10), max_width)


# =========================
# FORMATTING HELPERS
# =========================

def format_wos_full_text_header(ws) -> None:
    for col_idx in range(1, ws.max_column + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = Font(bold=True)
        cell.alignment = CENTER


def format_error_tab_header(ws) -> None:
    for col_idx in range(1, ws.max_column + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = Font(bold=True)
        cell.alignment = CENTER_TOP


def format_error_tab_body(ws) -> None:
    for row_idx in range(2, ws.max_row + 1):
        for col_idx in range(1, ws.max_column + 1):
            ws.cell(row=row_idx, column=col_idx).alignment = LEFT_TOP
        for col_idx in (1, 6, 8):
            ws.cell(row=row_idx, column=col_idx).alignment = CENTER


def style_title_green(ws, start_cell: str, end_cell: str, text: str) -> None:
    ws[start_cell] = text
    ws[start_cell].font = Font(bold=True, size=14, color=WHITE_FONT)
    ws[start_cell].fill = PatternFill(fill_type="solid", fgColor=GREEN_TITLE_FILL)
    ws[start_cell].alignment = CENTER
    ws.merge_cells(f"{start_cell}:{end_cell}")


def style_dark_header_row(ws, row_num: int, headers: list[str], start_col: int = 1) -> None:
    fill = PatternFill(fill_type="solid", fgColor=DARK_HEADER_FILL)
    font = Font(color=WHITE_FONT, bold=True)

    for idx, header in enumerate(headers, start=start_col):
        cell = ws.cell(row=row_num, column=idx, value=header)
        cell.fill = fill
        cell.font = font
        cell.alignment = CENTER


def style_blue_section_title(ws, row_num: int, col_start: int, col_end: int, text: str) -> None:
    cell = ws.cell(row=row_num, column=col_start, value=text)
    cell.fill = PatternFill(fill_type="solid", fgColor=BLUE_SECTION_FILL)
    cell.font = Font(color=WHITE_FONT, bold=True)
    cell.alignment = CENTER
    ws.merge_cells(start_row=row_num, start_column=col_start, end_row=row_num, end_column=col_end)


def style_blue_light_header_row(ws, row_num: int, headers: list[str], start_col: int = 1) -> None:
    fill = PatternFill(fill_type="solid", fgColor=LIGHT_BLUE_HEADER_FILL)
    font = Font(bold=True)

    for idx, header in enumerate(headers, start=start_col):
        cell = ws.cell(row=row_num, column=idx, value=header)
        cell.fill = fill
        cell.font = font
        cell.alignment = CENTER


def apply_stats_sheet_layout(ws) -> None:
    ws["A2"].font = Font(italic=True)
    ws["A2"].alignment = CENTER
    ws["A3"].font = Font(italic=True)
    ws["A3"].alignment = CENTER

    for row_idx in range(6, ws.max_row + 1):
        for col_idx in range(1, 6):
            ws.cell(row=row_idx, column=col_idx).alignment = CENTER

    for row_idx in range(1, ws.max_row + 1):
        if ws.cell(row=row_idx, column=2).value == "Total de erros":
            ws.cell(row=row_idx, column=2).font = Font(bold=True)
            ws.cell(row=row_idx, column=3).font = Font(bold=True)
            ws.cell(row=row_idx, column=2).alignment = CENTER
            ws.cell(row=row_idx, column=3).alignment = CENTER

    for row_idx in range(6, ws.max_row + 1):
        for col_idx in range(7, 11):
            ws.cell(row=row_idx, column=col_idx).alignment = CENTER

    for row_idx in range(1, ws.max_row + 1):
        for col_idx in range(1, ws.max_column + 1):
            if ws.cell(row=row_idx, column=col_idx).alignment == CENTER:
                continue
            if ws.cell(row=row_idx, column=col_idx).alignment == CENTER_TOP:
                continue
            ws.cell(row=row_idx, column=col_idx).alignment = LEFT_TOP

    widths = {
        "A": 8,
        "B": 34,
        "C": 10,
        "D": 12,
        "E": 18,
        "F": 3,
        "G": 8,
        "H": 30,
        "I": 28,
        "J": 10,
    }
    for col_letter, width in widths.items():
        ws.column_dimensions[col_letter].width = width


# =========================
# DIAGNOSTIC MODE / OUTPUT
# =========================

def resolve_diagnostic_mode(diagnostic_label: str) -> str:
    normalized = normalize_label(diagnostic_label)
    if "inicial" in normalized:
        return "initial"
    return "session"


def resolve_report_prefix(diagnostic_label: str) -> str:
    normalized = normalize_label(diagnostic_label)

    if "inicial" in normalized:
        return REPORT_NAME_DIAGNOSTICO_INICIAL
    if "fase 0" in normalized or "fase0" in normalized:
        return REPORT_NAME_DIAGNOSTICO_FASE0
    if "fase 1" in normalized or "fase1" in normalized:
        return REPORT_NAME_DIAGNOSTICO_FASE1
    if "fase 2" in normalized or "fase2" in normalized:
        return REPORT_NAME_DIAGNOSTICO_FASE2

    return "diagnostico"


def build_diagnostic_title(diagnostic_label: str) -> str:
    return diagnostic_label.upper()


def validate_required_input_columns(col_map: Dict[str, int], sheet_name: str, diagnostic_mode: str) -> None:
    if diagnostic_mode == "initial":
        required = DIAGNOSTIC_REQUIRED_INPUT_COLUMNS_INITIAL
    else:
        required = DIAGNOSTIC_REQUIRED_INPUT_COLUMNS_SESSION

    missing = [c for c in required if c not in col_map]
    if missing:
        raise ValueError(f"Missing required columns in '{sheet_name}' for diagnostic mode '{diagnostic_mode}': {missing}")


# =========================
# REPORT SHEETS
# =========================

def create_stats_sheet(
    wb: Workbook,
    diagnosed_rows: list[dict],
    source_sheet_name: str,
    diagnostic_label: str,
    diagnostic_mode: str,
) -> None:
    ws = wb.create_sheet(DIAGNOSTIC_STATS_SHEET_NAME)

    style_title_green(ws, "A1", "E1", "PDF download status - ranking de erros e top websites")
    ws["A2"] = f"Base analisada: {source_sheet_name} — {diagnostic_label}"

    if diagnostic_mode == "initial":
        ws["A3"] = (
            "Nota: o ranking de erros exclui o estado 'downloaded'. "
            "O campo 'website' corresponde ao host de pdf_source_url."
        )
    else:
        ws["A3"] = (
            "Nota: esta análise considera apenas registos cujo estado atual não é "
            "'downloaded'. O ranking exclui o estado final 'downloaded'. "
            "O campo 'website' corresponde ao host de pdf_source_url."
        )

    error_counter = Counter()
    error_website_counter = defaultdict(Counter)
    combo_counter = Counter()

    total_rows = len(diagnosed_rows)

    for row in diagnosed_rows:
        status = row.get("pdf_download_status")
        if status in DIAGNOSTIC_EXCLUDED_FROM_ERROR_RANKING:
            continue

        website = normalize_host(row.get("pdf_source_url"))
        error_counter[status] += 1
        error_website_counter[status][website] += 1
        combo_counter[(status, website)] += 1

    total_errors = sum(error_counter.values())

    fifth_header = "% total registos" if diagnostic_mode == "initial" else "% base não-downloaded"

    style_dark_header_row(
        ws,
        5,
        ["Rank", "Erro (pdf_download_status)", "N", "% erros", fifth_header],
        start_col=1,
    )

    sorted_errors = sorted(error_counter.items(), key=lambda x: (-x[1], str(x[0]).lower()))

    current_row = 6
    for rank, (error_name, count) in enumerate(sorted_errors, start=1):
        ws.cell(current_row, 1, rank)
        ws.cell(current_row, 2, error_name)
        ws.cell(current_row, 3, count)
        ws.cell(current_row, 4, count / total_errors if total_errors else 0)
        ws.cell(current_row, 5, count / total_rows if total_rows else 0)
        ws.cell(current_row, 4).number_format = "0.0%"
        ws.cell(current_row, 5).number_format = "0.0%"
        current_row += 1

    ws.cell(current_row, 2, "Total de erros")
    ws.cell(current_row, 3, total_errors)
    ws.cell(current_row, 2).font = Font(bold=True)
    ws.cell(current_row, 3).font = Font(bold=True)

    combo_start_col = 7
    style_blue_section_title(ws, 5, combo_start_col, combo_start_col + 3, "Combinações erro + website mais frequentes")
    style_blue_light_header_row(ws, 6, ["Rank", "Erro", "Website", "N"], start_col=combo_start_col)

    combo_sorted = sorted(
        combo_counter.items(),
        key=lambda x: (-x[1], str(x[0][0]).lower(), str(x[0][1]).lower()),
    )

    combo_row = 7
    for rank, ((error_name, website), n) in enumerate(combo_sorted[:15], start=1):
        ws.cell(combo_row, combo_start_col, rank)
        ws.cell(combo_row, combo_start_col + 1, error_name)
        ws.cell(combo_row, combo_start_col + 2, website)
        ws.cell(combo_row, combo_start_col + 3, n)
        combo_row += 1

    section_row = current_row + 4
    for error_name, _ in sorted_errors:
        style_blue_section_title(ws, section_row, 1, 3, f"Top websites — {error_name}")
        section_row += 1
        style_blue_light_header_row(ws, section_row, ["Rank", "Website", "N"], start_col=1)
        section_row += 1

        websites_sorted = sorted(
            error_website_counter[error_name].items(),
            key=lambda x: (-x[1], str(x[0]).lower()),
        )[:10]

        for rank, (website, n) in enumerate(websites_sorted, start=1):
            ws.cell(section_row, 1, rank)
            ws.cell(section_row, 2, website)
            ws.cell(section_row, 3, n)
            section_row += 1

        section_row += 2

    apply_stats_sheet_layout(ws)


def create_error_tabs(
    wb: Workbook,
    diagnosed_rows: list[dict],
) -> None:
    grouped = defaultdict(list)

    for row in diagnosed_rows:
        status = row.get("pdf_download_status")
        if status in DIAGNOSTIC_EXCLUDED_FROM_ERROR_RANKING:
            continue
        grouped[status].append(row)

    ordered_errors = sorted(grouped.keys(), key=lambda x: (-len(grouped[x]), str(x).lower()))

    for error_name in ordered_errors:
        ws = wb.create_sheet(make_safe_sheet_name(error_name))

        for col_idx, header in enumerate(DIAGNOSTIC_ERROR_TAB_COLUMNS, start=1):
            ws.cell(row=1, column=col_idx, value=header)

        format_error_tab_header(ws)

        current_row = 2
        for row in grouped[error_name]:
            values = [
                row.get("record_id", ""),
                row.get("Authors", ""),
                row.get("Title", ""),
                row.get("DOI", ""),
                row.get("DOI Link", ""),
                row.get("pdf_downloaded", ""),
                row.get("pdf_download_status", ""),
                row.get("pdf_source_url", ""),
                row.get("pdf_http_status", ""),
            ]

            for col_idx, value in enumerate(values, start=1):
                cell = ws.cell(row=current_row, column=col_idx, value=value)

                if col_idx == 5 and value:
                    cell.hyperlink = str(value)
                    cell.style = "Hyperlink"

                if col_idx == 8 and value:
                    cell.hyperlink = str(value)
                    cell.style = "Hyperlink"

            current_row += 1

        format_error_tab_body(ws)
        autosize_columns(ws)


# =========================
# DATA EXTRACTION
# =========================

def extract_rows_for_initial_diagnostic(
    ws,
    col_map: Dict[str, int],
    max_rows: Optional[int],
) -> list[dict]:
    rows: list[dict] = []
    processed = 0

    headers = DIAGNOSTIC_REQUIRED_INPUT_COLUMNS_INITIAL

    for row_idx in range(2, ws.max_row + 1):
        record_id = str(get_optional_cell(ws, row_idx, col_map, "record_id", "") or "").strip()
        if not record_id:
            continue

        row_values = {header: get_optional_cell(ws, row_idx, col_map, header, "") for header in headers}
        row_values["_source_row_idx"] = row_idx
        rows.append(row_values)

        processed += 1
        if max_rows is not None and processed >= max_rows:
            break

    return rows


def extract_rows_for_session_diagnostic(
    ws,
    col_map: Dict[str, int],
    max_rows: Optional[int],
) -> list[dict]:
    rows: list[dict] = []
    processed = 0

    headers = DIAGNOSTIC_REQUIRED_INPUT_COLUMNS_SESSION

    for row_idx in range(2, ws.max_row + 1):
        record_id = str(get_optional_cell(ws, row_idx, col_map, "record_id", "") or "").strip()
        if not record_id:
            continue

        row_values = {header: get_optional_cell(ws, row_idx, col_map, header, "") for header in headers}
        row_values["_source_row_idx"] = row_idx
        row_values["_previous_pdf_download_status"] = get_optional_cell(ws, row_idx, col_map, "pdf_download_status", "")
        rows.append(row_values)

        processed += 1
        if max_rows is not None and processed >= max_rows:
            break

    return rows


def diagnose_single_row(row_data: dict) -> dict:
    session = get_thread_session()

    doi_raw_str = str(row_data.get("DOI") or "").strip()
    doi_link = str(row_data.get("DOI Link") or "").strip()
    doi = normalize_doi(doi_raw_str)
    pdf_downloaded = 0
    try:
        pdf_downloaded = 1 if int(str(row_data.get("pdf_downloaded") or "0").strip()) == 1 else 0
    except Exception:
        pdf_downloaded = 0

    if pdf_downloaded == 1:
        diagnostic_status = str(row_data.get("pdf_download_status") or "").strip() or "downloaded"
        diagnostic_http_status = row_data.get("pdf_http_status", "")
        diagnostic_source_url = str(row_data.get("pdf_source_url") or "").strip()

    elif not doi_raw_str and not doi_link:
        diagnostic_status = "missing_doi"
        diagnostic_http_status = ""
        diagnostic_source_url = ""

    elif doi_raw_str and not doi and not doi_link:
        diagnostic_status = "invalid_doi"
        diagnostic_http_status = ""
        diagnostic_source_url = ""

    else:
        candidate_urls: list[str] = []

        if doi_link:
            candidate_urls.append(doi_link)

        if doi:
            doi_url = make_doi_url(doi)
            if doi_url not in candidate_urls:
                candidate_urls.append(doi_url)

        if not candidate_urls:
            diagnostic_status = "invalid_doi"
            diagnostic_http_status = ""
            diagnostic_source_url = ""
        else:
            diagnostic_status = "manual_check"
            diagnostic_http_status = ""
            diagnostic_source_url = ""

            for url in candidate_urls:
                status_label, http_status, final_url = try_download_diagnostic(url, session)
                diagnostic_status = status_label
                diagnostic_http_status = http_status if http_status is not None else ""
                diagnostic_source_url = final_url or url

                if status_label == "downloaded":
                    break

    result = {
        "record_id": row_data.get("record_id", ""),
        "Authors": row_data.get("Authors", ""),
        "Title": row_data.get("Title", ""),
        "DOI": row_data.get("DOI", ""),
        "DOI Link": row_data.get("DOI Link", ""),
        "pdf_downloaded": pdf_downloaded,
        "pdf_download_status": diagnostic_status,
        "pdf_source_url": diagnostic_source_url,
        "pdf_http_status": diagnostic_http_status,
        "_source_row_idx": row_data.get("_source_row_idx"),
    }
    return result


def should_refresh_progress(
    processed: int,
    last_refresh_time: float,
    force: bool = False,
) -> bool:
    if force:
        return True
    if processed % DIAGNOSTIC_PROGRESS_REFRESH_EVERY_N == 0:
        return True
    if (time.time() - last_refresh_time) >= DIAGNOSTIC_PROGRESS_REFRESH_EVERY_SECONDS:
        return True
    return False


# =========================
# MAIN PUBLIC ENTRYPOINT
# =========================

def run_diagnostic_for_session(
    session_state: SessionState,
    diagnostic_label: str,
    reporter: ReporterType = None,
    stop_event: StopEventType = None,
) -> dict:
    if not session_state.is_workbook_loaded:
        raise ValueError("SessionState workbook is not loaded.")

    ws = session_state.worksheet
    col_map = build_col_map(ws)
    reconcile_pdf_state_with_filesystem(ws, col_map)
    col_map = build_col_map(ws)
    session_state.col_map = col_map

    diagnostic_mode = resolve_diagnostic_mode(diagnostic_label)
    validate_required_input_columns(col_map, session_state.sheet_name, diagnostic_mode)

    max_rows = DIAGNOSTIC_MAX_ROWS_TO_PROCESS

    if diagnostic_mode == "initial":
        rows_to_process = extract_rows_for_initial_diagnostic(ws, col_map, max_rows)
    else:
        rows_to_process = extract_rows_for_session_diagnostic(ws, col_map, max_rows)

    total_to_process = len(rows_to_process)
    start_time = time.time()
    title = build_diagnostic_title(diagnostic_label)

    emit(
        reporter,
        "diagnostic_start",
        {
            "processed": 0,
            "total": total_to_process,
            "workers": DIAGNOSTIC_MAX_WORKERS,
            "start_time": start_time,
            "last_record_id": "",
            "last_status": "starting",
            "title": title,
            "stopped_early": False,
        },
    )

    diagnosed_rows: list[dict] = []
    processed = 0
    last_refresh_time = 0.0
    last_record = ""
    last_status = ""

    if total_to_process > 0:
        executor = ThreadPoolExecutor(max_workers=DIAGNOSTIC_MAX_WORKERS)

        try:
            futures = {}
            next_idx = 0

            while (
                next_idx < total_to_process
                and len(futures) < DIAGNOSTIC_MAX_WORKERS
                and not (stop_event is not None and stop_event.is_set())
            ):
                row_data = rows_to_process[next_idx]
                futures[executor.submit(diagnose_single_row, row_data)] = row_data
                next_idx += 1

            while futures:
                done, _ = wait(
                    list(futures.keys()),
                    timeout=0.1,
                    return_when=FIRST_COMPLETED,
                )

                if stop_event is not None and stop_event.is_set():
                    for pending in list(futures.keys()):
                        pending.cancel()
                    futures.clear()
                    break

                if not done:
                    continue

                for future in done:
                    futures.pop(future, None)
                    result = future.result()
                    diagnosed_rows.append(result)

                    processed += 1
                    last_record = str(result.get("record_id", ""))
                    last_status = str(result.get("pdf_download_status", ""))

                    if should_refresh_progress(processed, last_refresh_time):
                        emit(
                            reporter,
                            "diagnostic_update",
                            {
                                "processed": processed,
                                "total": total_to_process,
                                "workers": DIAGNOSTIC_MAX_WORKERS,
                                "start_time": start_time,
                                "last_record_id": last_record,
                                "last_status": last_status,
                                "title": title,
                                "stopped_early": False,
                            },
                        )
                        last_refresh_time = time.time()

                    if stop_event is not None and stop_event.is_set():
                        for pending in list(futures.keys()):
                            pending.cancel()
                        futures.clear()
                        break

                    while (
                        next_idx < total_to_process
                        and len(futures) < DIAGNOSTIC_MAX_WORKERS
                        and not (stop_event is not None and stop_event.is_set())
                    ):
                        row_data = rows_to_process[next_idx]
                        futures[executor.submit(diagnose_single_row, row_data)] = row_data
                        next_idx += 1

                if stop_event is not None and stop_event.is_set():
                    break

        finally:
            executor.shutdown(wait=False, cancel_futures=True)

    diagnosed_rows.sort(key=lambda row: (row.get("_source_row_idx") or 0))

    total_downloaded = sum(1 for row in diagnosed_rows if row.get("pdf_download_status") == "downloaded")
    total_errors = sum(
        1
        for row in diagnosed_rows
        if row.get("pdf_download_status") not in DIAGNOSTIC_EXCLUDED_FROM_ERROR_RANKING
    )

    stopped_early = stop_event is not None and stop_event.is_set()
    if stopped_early:
        emit(
            reporter,
            "diagnostic_end",
            {
                "processed": processed,
                "total": total_to_process,
                "workers": DIAGNOSTIC_MAX_WORKERS,
                "start_time": start_time,
                "last_record_id": last_record,
                "last_status": last_status if last_status else "stopped",
                "title": title,
                "report_path": "",
                "stopped_early": True,
            },
        )

        return {
            "diagnostic_label": diagnostic_label,
            "diagnostic_mode": diagnostic_mode,
            "processed": processed,
            "total": total_to_process,
            "downloaded": total_downloaded,
            "errors": total_errors,
            "report_path": "",
            "sheet_name": DIAGNOSTIC_OUTPUT_SHEET_NAME,
            "stats_sheet_name": DIAGNOSTIC_STATS_SHEET_NAME,
            "stopped_early": True,
        }

    output_wb = Workbook()
    output_ws = output_wb.active
    output_ws.title = DIAGNOSTIC_OUTPUT_SHEET_NAME

    for col_idx, header in enumerate(DIAGNOSTIC_OUTPUT_COLUMNS, start=1):
        output_ws.cell(row=1, column=col_idx, value=header)

    output_row = 2
    for row_values in diagnosed_rows:
        for col_idx, header in enumerate(DIAGNOSTIC_OUTPUT_COLUMNS, start=1):
            cell = output_ws.cell(row=output_row, column=col_idx, value=row_values.get(header, ""))

            if header == "DOI Link" and row_values.get(header):
                cell.hyperlink = str(row_values.get(header))
                cell.style = "Hyperlink"

            if header == "pdf_source_url" and row_values.get(header):
                cell.hyperlink = str(row_values.get(header))
                cell.style = "Hyperlink"

        output_row += 1

    format_wos_full_text_header(output_ws)
    autosize_columns(output_ws)

    create_stats_sheet(
        wb=output_wb,
        diagnosed_rows=diagnosed_rows,
        source_sheet_name=SHEET_NAME,
        diagnostic_label=diagnostic_label,
        diagnostic_mode=diagnostic_mode,
    )

    create_error_tabs(
        wb=output_wb,
        diagnosed_rows=diagnosed_rows,
    )

    report_prefix = resolve_report_prefix(diagnostic_label)
    report_path = build_report_output_path(report_prefix, timestamp=session_state.timestamp or None)

    ok = safe_save_workbook(output_wb, report_path)
    if not ok:
        raise RuntimeError(f"Could not save diagnostic report: {report_path}")

    summary = DiagnosticSummary(
        diagnostic_label=diagnostic_label,
        total_rows_considered=total_to_process,
        total_rows_diagnosed=processed,
        total_downloaded=total_downloaded,
        total_errors=total_errors,
        report_path=report_path,
    )
    session_state.add_diagnostic_summary(summary)

    emit(
        reporter,
        "diagnostic_end",
        {
            "processed": processed,
            "total": total_to_process,
            "workers": DIAGNOSTIC_MAX_WORKERS,
            "start_time": start_time,
            "last_record_id": last_record,
            "last_status": last_status if last_status else "finished",
            "title": title,
            "report_path": str(report_path),
            "stopped_early": False,
        },
    )

    return {
        "diagnostic_label": diagnostic_label,
        "diagnostic_mode": diagnostic_mode,
        "processed": processed,
        "total": total_to_process,
        "downloaded": total_downloaded,
        "errors": total_errors,
        "report_path": str(report_path),
        "sheet_name": DIAGNOSTIC_OUTPUT_SHEET_NAME,
        "stats_sheet_name": DIAGNOSTIC_STATS_SHEET_NAME,
        "stopped_early": False,
    }
