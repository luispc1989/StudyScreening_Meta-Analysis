from __future__ import annotations
from html import escape
import json
import os
import re
import shutil
import subprocess
import sys
import threading
import webbrowser
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from datetime import datetime
from pathlib import Path
from time import time
from typing import Any
from urllib.parse import quote, quote_plus, urlparse

from openpyxl import Workbook, load_workbook
import streamlit as st


REPO_ROOT = Path(__file__).resolve().parents[3]
if str(REPO_ROOT) not in sys.path:
    sys.path.insert(0, str(REPO_ROOT))

from tools.pdf_fetcher.core.config import ACTIVE_PROJECT_ROOT
from tools.pdf_fetcher.core.excel_io import build_col_map, safe_save_workbook
from tools.pdf_fetcher.core.utils import make_doi_url, normalize_doi
from tools.pdf_fetcher.scout.launcher import finalized_marker_path_for

KEYBOARD_NAVIGATION_HTML = Path(__file__).resolve().parent / "assets" / "keyboard_navigation.html"
LINK_ICON_PATH = REPO_ROOT / "tools" / "pdf_fetcher" / "assets" / "link.png"


BRIDGE_HOST = "127.0.0.1"
BRIDGE_PORT = 8765
DOWNLOAD_BRIDGE_STATE: dict[str, Any] = {
    "server_started": False,
    "events": [],
    "active_cases": {},
    "lock": threading.Lock(),
}


DECISION_OPTIONS = [
    "pending_review",
    "downloaded_manual",
    "paywalled",
    "unavailable",
    "wrong_link",
    "article_not_found",
    "other",
]


def set_page() -> None:
    st.set_page_config(
        page_title="S.C.O.U.T. - Semi-assisted Case Opening and User Triage",
        page_icon="SC",
        layout="wide",
    )


def init_state() -> None:
    defaults = {
        "scout_cases": [],
        "scout_current_index": 0,
        "scout_report_path": "",
        "scout_review_path": "",
        "scout_workbook_drafts": {},
        "scout_decisions": {},
        "scout_loaded": False,
        "scout_browser_message": "",
        "scout_download_message": "",
        "scout_download_message_record_id": "",
        "scout_pdf_status_flash_record_id": "",
        "scout_pdf_status_flash_until": 0.0,
        "scout_sync_message": "",
        "scout_download_markers": {},
        "scout_download_baselines": {},
        "scout_pending_auto_saves": {},
        "scout_report_source_expanded": True,
        "scout_auto_load_attempted": False,
    }
    for key, value in defaults.items():
        if key not in st.session_state:
            st.session_state[key] = value


def normalize_header(value: Any) -> str:
    return str(value or "").strip().lower()


def normalize_url(value: str) -> str:
    return str(value or "").strip()


def extract_host(value: str) -> str:
    try:
        return (urlparse(value).netloc or "").lower()
    except Exception:
        return ""


def build_case_signals(case: dict) -> dict[str, Any]:
    urls = [normalize_url(case.get("doi_url", "")), normalize_url(case.get("source_url", ""))]
    doi = str(case.get("doi", "") or "").strip().lower()
    doi_token = doi.replace("/", "").replace(".", "").replace("-", "") if doi else ""
    return {
        "record_id": case["record_id"],
        "title": case["title"],
        "output_file": case["output_file"],
        "doi": doi,
        "doi_token": doi_token,
        "urls": [url for url in urls if url],
        "hosts": sorted({extract_host(url) for url in urls if url}),
        "opened_at": time(),
        "watch_until": time() + 300,
    }


def get_preferred_case_url(case: dict) -> tuple[str, str]:
    doi_url = normalize_url(case.get("doi_url", ""))
    if doi_url:
        return doi_url, "doi"
    source_url = normalize_url(case.get("source_url", ""))
    if source_url:
        return source_url, "source"
    return "", "none"


def find_windows_browser_executable() -> str | None:
    candidates = [
        Path(os.environ.get("PROGRAMFILES", "")) / "Google" / "Chrome" / "Application" / "chrome.exe",
        Path(os.environ.get("PROGRAMFILES(X86)", "")) / "Google" / "Chrome" / "Application" / "chrome.exe",
        Path(os.environ.get("LOCALAPPDATA", "")) / "Google" / "Chrome" / "Application" / "chrome.exe",
        Path(os.environ.get("PROGRAMFILES", "")) / "Microsoft" / "Edge" / "Application" / "msedge.exe",
        Path(os.environ.get("PROGRAMFILES(X86)", "")) / "Microsoft" / "Edge" / "Application" / "msedge.exe",
        Path(os.environ.get("LOCALAPPDATA", "")) / "Microsoft" / "Edge" / "Application" / "msedge.exe",
    ]
    for candidate in candidates:
        if str(candidate) and candidate.exists():
            return str(candidate)
    return None


def start_download_bridge_server() -> None:
    if DOWNLOAD_BRIDGE_STATE["server_started"]:
        return

    class DownloadBridgeHandler(BaseHTTPRequestHandler):
        def _write_json(self, status: int, payload: dict[str, Any]) -> None:
            body = json.dumps(payload).encode("utf-8")
            self.send_response(status)
            self.send_header("Content-Type", "application/json")
            self.send_header("Access-Control-Allow-Origin", "*")
            self.send_header("Access-Control-Allow-Headers", "Content-Type")
            self.send_header("Content-Length", str(len(body)))
            self.end_headers()
            self.wfile.write(body)

        def do_OPTIONS(self) -> None:
            self.send_response(204)
            self.send_header("Access-Control-Allow-Origin", "*")
            self.send_header("Access-Control-Allow-Headers", "Content-Type")
            self.send_header("Access-Control-Allow-Methods", "POST, GET, OPTIONS")
            self.end_headers()

        def do_POST(self) -> None:
            if self.path != "/download-event":
                self._write_json(404, {"ok": False, "error": "not_found"})
                return

            try:
                length = int(self.headers.get("Content-Length", "0"))
                payload = json.loads(self.rfile.read(length).decode("utf-8"))
            except Exception as exc:
                self._write_json(400, {"ok": False, "error": f"invalid_json: {exc}"})
                return

            payload["received_at"] = time()
            payload["matched_record_id"] = None
            payload["status"] = "pending"

            with DOWNLOAD_BRIDGE_STATE["lock"]:
                DOWNLOAD_BRIDGE_STATE["events"].append(payload)

            self._write_json(200, {"ok": True})

        def do_GET(self) -> None:
            if self.path != "/health":
                self._write_json(404, {"ok": False, "error": "not_found"})
                return
            self._write_json(200, {"ok": True})

        def log_message(self, format: str, *args) -> None:
            return

    server = ThreadingHTTPServer((BRIDGE_HOST, BRIDGE_PORT), DownloadBridgeHandler)
    thread = threading.Thread(target=server.serve_forever, daemon=True)
    thread.start()
    DOWNLOAD_BRIDGE_STATE["server_started"] = True


def register_active_case(case: dict) -> None:
    signals = build_case_signals(case)
    with DOWNLOAD_BRIDGE_STATE["lock"]:
        DOWNLOAD_BRIDGE_STATE["active_cases"][case["record_id"]] = signals


def get_active_case_info(record_id: str) -> dict[str, Any] | None:
    with DOWNLOAD_BRIDGE_STATE["lock"]:
        info = DOWNLOAD_BRIDGE_STATE["active_cases"].get(record_id)
        return dict(info) if info else None


def prune_expired_active_cases() -> None:
    now_ts = time()
    with DOWNLOAD_BRIDGE_STATE["lock"]:
        expired = [
            record_id
            for record_id, info in DOWNLOAD_BRIDGE_STATE["active_cases"].items()
            if float(info.get("watch_until") or 0) <= now_ts
        ]
        for record_id in expired:
            DOWNLOAD_BRIDGE_STATE["active_cases"].pop(record_id, None)


def score_event_for_case(event: dict[str, Any], case_info: dict[str, Any]) -> int:
    score = 0
    event_text = " ".join(
        str(event.get(key, "") or "").lower()
        for key in ("finalUrl", "referrer", "tabUrl", "filename", "mime")
    )

    if event.get("mime", "").lower() == "application/pdf":
        score += 2
    if str(event.get("filename", "")).lower().endswith(".pdf"):
        score += 2

    for host in case_info["hosts"]:
        if host and host in event_text:
            score += 4

    doi = case_info["doi"]
    if doi and doi in event_text:
        score += 8

    doi_token = case_info["doi_token"]
    if doi_token and doi_token in event_text.replace("/", "").replace(".", "").replace("-", ""):
        score += 6

    event_time = float(event.get("received_at") or time())
    opened_at = float(case_info.get("opened_at") or 0)
    if event_time >= opened_at:
        score += 3
        delta = event_time - opened_at
        if delta <= 300:
            score += 3
        elif delta <= 1800:
            score += 1
    else:
        score -= 10

    return score


def move_download_to_case(case: dict, source_path: Path) -> Path:
    if not source_path.exists():
        raise FileNotFoundError(f"Downloaded file not found: {source_path}")

    output_path = Path(case["output_file"])
    output_path.parent.mkdir(parents=True, exist_ok=True)
    shutil.move(str(source_path), str(output_path))
    return output_path


def process_download_bridge_events(cases: list[dict]) -> None:
    case_map = {case["record_id"]: case for case in cases}
    decisions = st.session_state["scout_decisions"]

    with DOWNLOAD_BRIDGE_STATE["lock"]:
        events = DOWNLOAD_BRIDGE_STATE["events"]
        active_cases = dict(DOWNLOAD_BRIDGE_STATE["active_cases"])

        for event in events:
            if event.get("status") != "pending":
                continue

            filename = str(event.get("filename", "") or "").strip()
            if not filename.lower().endswith(".pdf"):
                event["status"] = "ignored"
                continue

            scores: list[tuple[int, str]] = []
            for record_id, case_info in active_cases.items():
                if record_id not in case_map:
                    continue
                score = score_event_for_case(event, case_info)
                if score > 0:
                    scores.append((score, record_id))

            scores.sort(reverse=True)
            if not scores:
                event["status"] = "unmatched"
                continue

            best_score, best_record_id = scores[0]
            second_score = scores[1][0] if len(scores) > 1 else -999
            if best_score < 8 or (best_score - second_score) < 3:
                event["status"] = "ambiguous"
                event["candidate_record_ids"] = [record_id for _, record_id in scores[:3]]
                continue

            source_path = Path(filename)
            if not source_path.exists():
                event["status"] = "error"
                event["error"] = f"Downloaded file not found: {source_path}"
                continue

            event["status"] = "detected"
            event["matched_record_id"] = best_record_id
            event["saved_path"] = str(source_path)
            active_cases.pop(best_record_id, None)
            DOWNLOAD_BRIDGE_STATE["active_cases"].pop(best_record_id, None)
            stage_pending_auto_save(case_map[best_record_id], source_path)


def get_bridge_events_for_case(record_id: str) -> list[dict[str, Any]]:
    with DOWNLOAD_BRIDGE_STATE["lock"]:
        result = []
        for idx, event in enumerate(DOWNLOAD_BRIDGE_STATE["events"]):
            event_copy = dict(event)
            event_copy["event_index"] = idx
            if event.get("matched_record_id") == record_id:
                result.append(event_copy)
                continue
            if record_id in event.get("candidate_record_ids", []):
                result.append(event_copy)
        return result


def assign_download_event_to_case(event_index: int, case: dict) -> Path:
    with DOWNLOAD_BRIDGE_STATE["lock"]:
        event = DOWNLOAD_BRIDGE_STATE["events"][event_index]
        source_path = Path(str(event.get("filename", "")))

    saved_path = move_download_to_case(case, source_path)

    with DOWNLOAD_BRIDGE_STATE["lock"]:
        event = DOWNLOAD_BRIDGE_STATE["events"][event_index]
        event["status"] = "detected"
        event["matched_record_id"] = case["record_id"]
        event["saved_path"] = str(saved_path)
        event["candidate_record_ids"] = []
        DOWNLOAD_BRIDGE_STATE["active_cases"].pop(case["record_id"], None)

    stage_pending_auto_save(case, saved_path)
    return saved_path


def load_cases_from_report(report_path: Path) -> list[dict]:
    wb = load_workbook(report_path)
    try:
        ws = wb.active
        headers = [normalize_header(cell.value) for cell in ws[1]]
        has_authors = "authors" in headers
        has_publication_year = "publication_year" in headers
        has_email_address = "email_address" in headers
        rows: list[dict] = []
        for row in ws.iter_rows(min_row=2, values_only=True):
            if not any(row):
                continue
            case = {headers[idx]: row[idx] for idx in range(min(len(headers), len(row)))}
            authors = str(case.get("authors", "") or "").strip() if has_authors else ""
            publication_year = str(case.get("publication_year", "") or "").strip() if has_publication_year else ""
            email_address = str(case.get("email_address", "") or "").strip() if has_email_address else ""
            rows.append(
                {
                    "record_id": str(case.get("record_id", "") or "").strip(),
                    "title": str(case.get("title", "") or "").strip(),
                    "authors": authors,
                    "publication_year": publication_year,
                    "email_address": email_address,
                    "doi": str(case.get("doi", "") or "").strip(),
                    "doi_url": str(case.get("doi_url", "") or "").strip(),
                    "source_url": str(case.get("source_url", "") or "").strip(),
                    "original_doi": str(case.get("doi", "") or "").strip(),
                    "original_doi_url": str(case.get("doi_url", "") or "").strip(),
                    "original_source_url": str(case.get("source_url", "") or "").strip(),
                    "output_file": str(case.get("output_file", "") or "").strip(),
                    "error_type": str(case.get("error_type", "") or "").strip(),
                    "error_message": str(case.get("error_message", "") or "").strip(),
                    "scout_mode": str(case.get("scout_mode", "") or "").strip(),
                    "workbook_path": str(case.get("workbook_path", "") or "").strip(),
                    "sheet_name": str(case.get("sheet_name", "") or "").strip(),
                    "source_row_idx": int(case.get("source_row_idx", 0) or 0),
                }
            )
        return sorted(rows, key=_case_sort_key)
    finally:
        wb.close()


def _case_sort_key(case: dict) -> tuple[int, int | str]:
    record_id = str(case.get("record_id", "") or "").strip()
    try:
        return (0, int(record_id))
    except Exception:
        return (1, record_id.lower())


def normalize_display_title(value: str) -> str:
    title = " ".join(str(value or "").split()).strip()
    if not title:
        return "(Untitled case)"

    letters = [ch for ch in title if ch.isalpha()]
    if not letters:
        return title

    uppercase_ratio = sum(1 for ch in letters if ch.isupper()) / max(len(letters), 1)
    if uppercase_ratio >= 0.75:
        return title.title()

    return title


def review_output_path_for(report_path: Path) -> Path:
    target_dir = report_path.parent / "scout_reviews"
    target_dir.mkdir(parents=True, exist_ok=True)
    return target_dir / f"{report_path.stem}_scout_review.xlsx"


def draft_workbook_path_for(report_path: Path, workbook_path: Path) -> Path:
    target_dir = report_path.parent / "scout_drafts"
    target_dir.mkdir(parents=True, exist_ok=True)
    safe_stem = workbook_path.stem.replace(" ", "_")
    return target_dir / f"{report_path.stem}__{safe_stem}__draft.xlsx"


def ensure_scout_draft_workbooks(cases: list[dict], report_path: Path) -> dict[str, str]:
    draft_paths: dict[str, str] = {}
    seen_sources: set[str] = set()

    for case in cases:
        source_raw = str(case.get("workbook_path", "") or "").strip()
        if not source_raw or source_raw in seen_sources:
            continue
        seen_sources.add(source_raw)

        source_path = Path(source_raw).expanduser()
        draft_path = draft_workbook_path_for(report_path, source_path)
        if not draft_path.exists():
            shutil.copy2(source_path, draft_path)
        draft_paths[source_raw] = str(draft_path)

    return draft_paths


def save_review_workbook(cases: list[dict], decisions: dict[str, dict], review_path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "scout_review"
    ws.append(
        [
            "record_id",
            "title",
            "authors",
            "publication_year",
            "email_address",
            "doi",
            "doi_url",
            "source_url",
            "output_file",
            "original_error_type",
            "original_error_message",
            "scout_mode",
            "workbook_path",
            "sheet_name",
            "source_row_idx",
            "decision",
            "notes",
            "reviewed_at",
            "pdf_saved",
            "staged_doi",
            "staged_doi_url",
            "staged_source_url",
        ]
    )

    for case in cases:
        record_id = case["record_id"]
        review = decisions.get(record_id, {})
        ws.append(
            [
                record_id,
                case["title"],
                case.get("authors", ""),
                case.get("publication_year", ""),
                case.get("email_address", ""),
                case["doi"],
                case["doi_url"],
                case["source_url"],
                case["output_file"],
                case["error_type"],
                case["error_message"],
                case.get("scout_mode", ""),
                case.get("workbook_path", ""),
                case.get("sheet_name", ""),
                case.get("source_row_idx", ""),
                review.get("decision", ""),
                review.get("notes", ""),
                review.get("reviewed_at", ""),
                "yes" if review.get("pdf_saved") else "no",
                review.get("staged_doi", ""),
                review.get("staged_doi_url", ""),
                review.get("staged_source_url", ""),
            ]
        )

    wb.save(review_path)


def load_existing_review(review_path: Path) -> dict[str, dict]:
    if not review_path.exists():
        return {}

    wb = load_workbook(review_path)
    try:
        ws = wb.active
        headers = [normalize_header(cell.value) for cell in ws[1]]
        decisions: dict[str, dict] = {}
        for row in ws.iter_rows(min_row=2, values_only=True):
            if not any(row):
                continue
            data = {headers[idx]: row[idx] for idx in range(min(len(headers), len(row)))}
            record_id = str(data.get("record_id", "") or "").strip()
            if not record_id:
                continue
            decision = str(data.get("decision", "") or "").strip()
            notes = str(data.get("notes", "") or "").strip()
            reviewed_at = str(data.get("reviewed_at", "") or "").strip()
            pdf_saved = str(data.get("pdf_saved", "") or "").strip().lower() == "yes"
            staged_doi = str(data.get("staged_doi", "") or "").strip()
            staged_doi_url = str(data.get("staged_doi_url", "") or "").strip()
            staged_source_url = str(data.get("staged_source_url", "") or "").strip()

            # Ignore empty placeholder rows from the review workbook. They should
            # not count as completed review history when reloading SCOUT.
            if (
                not decision
                and not notes
                and not reviewed_at
                and not pdf_saved
                and not staged_doi
                and not staged_doi_url
                and not staged_source_url
            ):
                continue

            decisions[record_id] = {
                "decision": decision,
                "notes": notes,
                "reviewed_at": reviewed_at,
                "pdf_saved": pdf_saved,
            }
            if staged_doi:
                decisions[record_id]["staged_doi"] = staged_doi
            if staged_doi_url:
                decisions[record_id]["staged_doi_url"] = staged_doi_url
            if staged_source_url:
                decisions[record_id]["staged_source_url"] = staged_source_url
        return decisions
    finally:
        wb.close()


def reconcile_decisions_with_existing_pdfs(cases: list[dict], decisions: dict[str, dict]) -> tuple[dict[str, dict], bool]:
    updated = dict(decisions)
    changed = False

    for case in cases:
        output_file = str(case.get("output_file", "") or "").strip()
        if not output_file:
            continue

        output_path = Path(output_file)
        if not output_path.exists() or not output_path.is_file():
            continue

        record_id = str(case.get("record_id", "") or "").strip()
        if not record_id:
            continue

        existing = dict(updated.get(record_id, {}))
        existing_decision = str(existing.get("decision", "") or "").strip()
        existing_notes = str(existing.get("notes", "") or "").strip()
        existing_pdf_saved = bool(existing.get("pdf_saved", False))
        existing_reviewed_at = str(existing.get("reviewed_at", "") or "").strip()

        if existing_decision == "downloaded_manual" and existing_pdf_saved:
            continue

        updated[record_id] = {
            "decision": "downloaded_manual",
            "notes": existing_notes,
            "reviewed_at": existing_reviewed_at or datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
            "pdf_saved": True,
        }
        changed = True

    return updated, changed


def apply_review_overrides_to_cases(cases: list[dict], decisions: dict[str, dict]) -> None:
    for case in cases:
        review = decisions.get(str(case.get("record_id", "") or "").strip(), {})
        if "staged_doi" in review:
            case["doi"] = str(review.get("staged_doi", "") or "").strip()
        if "staged_doi_url" in review:
            case["doi_url"] = str(review.get("staged_doi_url", "") or "").strip()
        if "staged_source_url" in review:
            case["source_url"] = str(review.get("staged_source_url", "") or "").strip()


def find_resume_index(cases: list[dict], decisions: dict[str, dict]) -> int:
    if not cases:
        return 0

    for idx, case in enumerate(cases):
        review = decisions.get(case["record_id"], {})
        decision = str(review.get("decision", "") or "").strip()
        if not decision or decision == "pending_review":
            return idx

    return max(0, len(cases) - 1)


def load_report(report_path_value: str) -> None:
    report_path = Path(report_path_value).expanduser()
    cases = load_cases_from_report(report_path)
    review_path = review_output_path_for(report_path)
    draft_paths = ensure_scout_draft_workbooks(cases, report_path)
    decisions = load_existing_review(review_path)
    decisions, decisions_changed = reconcile_decisions_with_existing_pdfs(cases, decisions)
    apply_review_overrides_to_cases(cases, decisions)

    with DOWNLOAD_BRIDGE_STATE["lock"]:
        DOWNLOAD_BRIDGE_STATE["events"].clear()
        DOWNLOAD_BRIDGE_STATE["active_cases"].clear()

    if decisions_changed:
        save_review_workbook(cases, decisions, review_path)

    st.session_state["scout_cases"] = cases
    st.session_state["scout_current_index"] = find_resume_index(cases, decisions)
    st.session_state["scout_report_path"] = str(report_path)
    st.session_state["scout_review_path"] = str(review_path)
    st.session_state["scout_workbook_drafts"] = draft_paths
    st.session_state["scout_decisions"] = decisions
    st.session_state["scout_loaded"] = True
    st.session_state["scout_sync_message"] = ""
    st.session_state["scout_download_message"] = ""
    st.session_state["scout_download_message_record_id"] = ""
    st.session_state["scout_pdf_status_flash_record_id"] = ""
    st.session_state["scout_pdf_status_flash_until"] = 0.0


def try_auto_load_from_environment() -> None:
    if st.session_state.get("scout_loaded"):
        return
    if st.session_state.get("scout_auto_load_attempted"):
        return

    st.session_state["scout_auto_load_attempted"] = True

    auto_load_enabled = str(os.environ.get("SCOUT_AUTO_LOAD", "") or "").strip() == "1"
    report_path_value = str(os.environ.get("SCOUT_AUTO_REPORT_PATH", "") or "").strip()
    if not auto_load_enabled or not report_path_value:
        return

    try:
        load_report(report_path_value)
        st.session_state["scout_report_source_expanded"] = False
    except Exception as exc:
        st.session_state["scout_browser_message"] = f"Could not auto-load SCOUT report: {exc}"
        st.session_state["scout_report_source_expanded"] = True


def get_current_case() -> dict | None:
    cases = st.session_state["scout_cases"]
    if not cases:
        return None
    idx = max(0, min(st.session_state["scout_current_index"], len(cases) - 1))
    st.session_state["scout_current_index"] = idx
    return cases[idx]


def start_download_detection(case: dict) -> None:
    downloads_dir = Path.home() / "Downloads"
    baseline: set[str] = set()
    if downloads_dir.exists():
        baseline = {path.name.lower() for path in downloads_dir.glob("*.pdf") if path.is_file()}

    st.session_state["scout_download_markers"][case["record_id"]] = time()
    st.session_state["scout_download_baselines"][case["record_id"]] = sorted(baseline)
    register_active_case(case)


def set_pdf_detected_feedback(case: dict, saved_path: Path, *, automatic: bool) -> None:
    record_id = str(case.get("record_id", "") or "").strip()
    message = ""
    if not automatic:
        message = f"PDF imported and saved to `{saved_path}`."

    st.session_state["scout_download_message"] = message
    st.session_state["scout_download_message_record_id"] = record_id
    st.session_state["scout_pdf_status_flash_record_id"] = record_id
    st.session_state["scout_pdf_status_flash_until"] = time() + 4.0


def stage_pending_auto_save(case: dict, source_path: Path) -> None:
    record_id = str(case.get("record_id", "") or "").strip()
    st.session_state["scout_pending_auto_saves"][record_id] = {
        "source_path": str(source_path),
        "detected_at": time(),
    }
    st.session_state["scout_pdf_status_flash_record_id"] = record_id
    st.session_state["scout_pdf_status_flash_until"] = time() + 4.0


def get_active_scout_mode() -> str:
    cases = st.session_state.get("scout_cases", [])
    if cases:
        return str(cases[0].get("scout_mode", "") or "").strip().lower()
    return ""


def launch_scout_browser(url: str, case: dict) -> None:
    if not url:
        raise ValueError("No URL is available for this case.")

    label = f"{case.get('record_id', '')} - {case.get('title', '')}".strip(" -") or "Current case"
    opened = False
    launch_mode = "browser"

    if sys.platform.startswith("win"):
        browser_executable = find_windows_browser_executable()
        if browser_executable:
            creationflags = subprocess.CREATE_NEW_PROCESS_GROUP | subprocess.DETACHED_PROCESS
            subprocess.Popen(
                [browser_executable, url],
                creationflags=creationflags,
                close_fds=False,
            )
            opened = True
            launch_mode = "browser tab"

    if not opened:
        opened = webbrowser.open_new(url)
        launch_mode = "new tab"

    if not opened:
        raise RuntimeError("The system browser could not be opened for this URL.")

    start_download_detection(case)
    return None


def build_google_search_url(case: dict) -> str:
    title = str(case.get("title", "") or "").strip()
    if not title:
        raise ValueError("No article title is available for this case.")
    return f"https://www.google.com/search?q={quote_plus(title)}"


def extract_email_recipients(value: str) -> list[str]:
    raw = str(value or "").strip()
    if not raw:
        return []

    candidates = re.split(r"[;, \n\r\t]+", raw)
    recipients: list[str] = []
    seen: set[str] = set()
    for candidate in candidates:
        email = candidate.strip().strip(".,")
        if not email or "@" not in email:
            continue
        normalized = email.lower()
        if normalized in seen:
            continue
        seen.add(normalized)
        recipients.append(email)
    return recipients


def build_request_email_mailto(case: dict) -> str:
    recipients = extract_email_recipients(str(case.get("email_address", "") or ""))
    if not recipients:
        raise ValueError("No author email address is available for this case.")

    title = str(case.get("title", "") or "").strip()
    doi = str(case.get("doi", "") or "").strip()
    doi_url = str(case.get("doi_url", "") or "").strip()

    subject = f"Request for full text article: {title}" if title else "Request for full text article"
    body_lines = [
        "Dear Author,",
        "",
        "I hope you are well.",
        "",
        "I am writing to kindly request a copy of the following article for academic research purposes:",
        "",
    ]
    if title:
        body_lines.append(f"Title: {title}")
    if doi:
        body_lines.append(f"DOI: {doi}")
    if doi_url:
        body_lines.append(f"DOI Link: {doi_url}")

    body_lines.extend(
        [
            "",
            "If you are able to share the paper, I would be very grateful.",
            "",
            "Kind regards,",
            "",
        ]
    )
    body = "\r\n".join(body_lines)
    recipients_part = ";".join(recipients)
    return f"mailto:{quote(recipients_part, safe=';@')}" f"?subject={quote(subject)}&body={quote(body)}"


def open_mailto_link(mailto_url: str) -> None:
    opened = False

    if sys.platform.startswith("win"):
        try:
            os.startfile(mailto_url)  # type: ignore[attr-defined]
            opened = True
        except Exception:
            opened = False

    if not opened:
        opened = webbrowser.open_new(mailto_url)

    if not opened:
        raise RuntimeError("The system email client could not be opened for this request.")


def open_url_in_default_browser(url: str) -> None:
    opened = False

    if sys.platform.startswith("win"):
        try:
            os.startfile(url)  # type: ignore[attr-defined]
            opened = True
        except Exception:
            opened = False

    if not opened:
        opened = webbrowser.open_new(url)

    if not opened:
        raise RuntimeError("The system browser could not be opened for this URL.")


def open_url_in_chrome(url: str) -> None:
    if not url:
        raise ValueError("No URL is available.")

    opened = False

    if sys.platform.startswith("win"):
        try:
            creationflags = subprocess.CREATE_NEW_PROCESS_GROUP | subprocess.DETACHED_PROCESS
            subprocess.Popen(
                ["cmd", "/c", "start", "", "chrome", url],
                creationflags=creationflags,
                close_fds=False,
            )
            opened = True
        except Exception:
            opened = False

        if not opened:
            browser_executable = find_windows_browser_executable()
            if browser_executable:
                try:
                    creationflags = subprocess.CREATE_NEW_PROCESS_GROUP | subprocess.DETACHED_PROCESS
                    subprocess.Popen(
                        [browser_executable, "--new-tab", url],
                        creationflags=creationflags,
                        close_fds=False,
                    )
                    opened = True
                except Exception:
                    opened = False

    if not opened:
        open_url_in_default_browser(url)


def render_request_email_button(case: dict) -> None:
    recipients = extract_email_recipients(str(case.get("email_address", "") or ""))
    disabled = not recipients
    help_text = None if recipients else "No author email address is available for this case."

    if st.button(
        "Send Request Email",
        use_container_width=True,
        key=f"send_request_email_sidebar_{case['record_id']}",
        icon=":material/mail:",
        disabled=disabled,
        help=help_text,
    ):
        try:
            open_mailto_link(build_request_email_mailto(case))
            st.session_state["scout_browser_message"] = ""
        except Exception as exc:
            st.session_state["scout_browser_message"] = f"Could not open request email: {exc}"
        st.rerun()


def render_google_search_button(case: dict) -> None:
    try:
        search_url = build_google_search_url(case)
    except Exception as exc:
        st.caption(f"Search unavailable: {exc}")
        return

    if st.button(
        "Search Engine",
        use_container_width=True,
        key=f"search_article_sidebar_{case['record_id']}",
        icon=":material/search:",
    ):
        try:
            open_url_in_chrome(search_url)
            start_download_detection(case)
            st.session_state["scout_browser_message"] = ""
        except Exception as exc:
            st.session_state["scout_browser_message"] = f"Could not open search: {exc}"
        st.rerun()


def render_import_pdf_button(case: dict, existing: dict) -> bool:
    output_file = str(case.get("output_file", "") or "").strip()
    if output_file and Path(output_file).exists():
        disabled = True
        help_text = "A PDF is already saved for this case."
    else:
        candidates = get_new_download_candidates(case)
        record_id = case["record_id"]
        marker_exists = st.session_state.get("scout_download_markers", {}).get(record_id) is not None
        active_info = get_active_case_info(record_id)

        if candidates:
            disabled = False
            help_text = "A new PDF was detected in Downloads for this case."
        elif marker_exists or active_info:
            disabled = True
            help_text = "Watching Downloads for a matching PDF. This button will activate when one is detected."
        else:
            disabled = True
            help_text = "Open the article link first so S.C.O.U.T. can watch Downloads for this case."

    return st.button(
        "Import PDF",
        use_container_width=True,
        key=f"import_download_sidebar_{case['record_id']}",
        icon=":material/picture_as_pdf:",
        disabled=disabled,
        help=help_text,
    )


def render_open_article_button(case: dict) -> bool:
    launch_url, launch_kind = get_preferred_case_url(case)
    return st.button(
        "Open Article Link",
        use_container_width=True,
        key=f"open_article_sidebar_{case['record_id']}",
        icon=":material/link:",
        disabled=not launch_url,
        help=(
            None
            if launch_url
            else "No DOI Link or Source Link is available for this case. Use Search Engine or Send Request Email."
        ),
    )


def render_navigation_button(
    *,
    label: str,
    key: str,
    icon_side: str = "left",
    icon: str | None = None,
    disabled: bool = False,
) -> bool:
    return st.button(
        label,
        use_container_width=True,
        key=key,
        icon=icon,
        icon_position="right" if icon_side == "right" else "left",
        disabled=disabled,
    )


def build_scout_checked_at(decision: str, notes: str) -> str:
    timestamp = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    clean_notes = str(notes or "").strip()
    if clean_notes:
        return f"{timestamp} | SCOUT: {decision} | {clean_notes}"
    return f"{timestamp} | SCOUT: {decision}"


def normalize_inserted_doi_link(value: str) -> tuple[str, str]:
    raw_value = str(value or "").strip()
    if not raw_value:
        raise ValueError("Please provide a DOI link or DOI value.")

    normalized_doi = normalize_doi(raw_value)
    if normalized_doi:
        return normalized_doi, make_doi_url(normalized_doi)

    raise ValueError("The inserted value is not a valid DOI or DOI link.")


def normalize_inserted_source_link(value: str) -> str:
    raw_value = str(value or "").strip()
    if not raw_value:
        raise ValueError("Please provide a source link.")
    if not raw_value.lower().startswith(("http://", "https://")):
        raise ValueError("The inserted source link must start with http:// or https://")
    return raw_value


def _find_case_row(ws, case: dict) -> int:
    source_row_idx = int(case.get("source_row_idx", 0) or 0)
    record_id = str(case.get("record_id", "") or "").strip()
    col_map = build_col_map(ws)
    record_col = col_map.get("record_id")
    if record_col is None:
        raise ValueError("Workbook sheet does not contain 'record_id'.")

    if source_row_idx >= 2:
        source_record_id = str(ws.cell(row=source_row_idx, column=record_col).value or "").strip()
        if not record_id or source_record_id == record_id:
            return source_row_idx

    for row_idx in range(2, ws.max_row + 1):
        value = str(ws.cell(row=row_idx, column=record_col).value or "").strip()
        if value == record_id:
            return row_idx

    raise ValueError(f"Record ID not found in workbook: {record_id}")


def get_case_workbook_paths(case: dict) -> tuple[Path, Path]:
    source_raw = str(case.get("workbook_path", "") or "").strip()
    if not source_raw:
        raise ValueError("SCOUT case is missing workbook location metadata.")

    source_path = Path(source_raw).expanduser()
    draft_map = st.session_state.get("scout_workbook_drafts", {})
    draft_raw = str(draft_map.get(source_raw, "") or "").strip()
    draft_path = Path(draft_raw).expanduser() if draft_raw else source_path
    return source_path, draft_path


def save_scout_workbook_fast(wb, workbook_path: Path, max_attempts: int = 2) -> None:
    workbook_path.parent.mkdir(parents=True, exist_ok=True)
    temp_path = workbook_path.with_name(workbook_path.stem + ".__scout_tmp__.xlsx")

    last_error: Exception | None = None
    for _ in range(max_attempts):
        try:
            wb.save(temp_path)
            try:
                temp_path.replace(workbook_path)
            except PermissionError:
                if workbook_path.exists():
                    workbook_path.unlink()
                temp_path.replace(workbook_path)
            return
        except KeyboardInterrupt:
            raise
        except Exception as exc:
            last_error = exc

    if temp_path.exists():
        try:
            temp_path.unlink()
        except Exception:
            pass
    raise RuntimeError(f"Could not save workbook after SCOUT update: {workbook_path}") from last_error


def sync_case_review_to_workbook(case: dict, decision: str, notes: str, pdf_saved: bool) -> None:
    scout_mode = str(case.get("scout_mode", "") or "").strip().lower()
    if scout_mode == "load_fail_test_cases":
        return

    sheet_name = str(case.get("sheet_name", "") or "").strip()
    if not sheet_name:
        raise ValueError("SCOUT case is missing workbook location metadata.")

    _source_workbook_path, workbook_path = get_case_workbook_paths(case)
    wb = load_workbook(workbook_path)
    try:
        if sheet_name not in wb.sheetnames:
            raise ValueError(f"Sheet not found in workbook: {sheet_name}")

        ws = wb[sheet_name]
        col_map = build_col_map(ws)
        row_idx = _find_case_row(ws, case)

        required_headers = ["pdf_downloaded", "pdf_download_status", "pdf_file_name", "pdf_local_path", "pdf_checked_at"]
        missing = [header for header in required_headers if header not in col_map]
        if missing:
            raise ValueError(f"Workbook is missing required PDF columns: {missing}")

        output_file = str(case.get("output_file", "") or "").strip()
        output_path = Path(output_file) if output_file else None
        pdf_exists = bool(output_path and output_path.exists() and output_path.is_file())

        status_map = {
            "downloaded_manual": "downloaded",
            "paywalled": "scout_paywalled",
            "unavailable": "scout_unavailable",
            "wrong_link": "scout_wrong_link",
            "article_not_found": "scout_article_not_found",
            "other": "scout_other",
        }
        target_status = status_map.get(decision, "scout_pending_review")

        if decision == "downloaded_manual" and (pdf_saved or pdf_exists):
            ws.cell(row=row_idx, column=col_map["pdf_downloaded"], value=1)
            ws.cell(row=row_idx, column=col_map["pdf_download_status"], value="downloaded")
            if output_path is not None:
                ws.cell(row=row_idx, column=col_map["pdf_file_name"], value=output_path.name)
                try:
                    local_path = str(output_path.relative_to(ACTIVE_PROJECT_ROOT))
                except ValueError:
                    local_path = str(output_path)
                ws.cell(row=row_idx, column=col_map["pdf_local_path"], value=local_path)
        else:
            ws.cell(row=row_idx, column=col_map["pdf_downloaded"], value=0)
            ws.cell(row=row_idx, column=col_map["pdf_download_status"], value=target_status)

        ws.cell(row=row_idx, column=col_map["pdf_checked_at"], value=build_scout_checked_at(decision, notes))

        save_scout_workbook_fast(wb, workbook_path)
    finally:
        wb.close()


def apply_case_review_to_worksheet(ws, case: dict, decision: str, notes: str, pdf_saved: bool) -> None:
    col_map = build_col_map(ws)
    row_idx = _find_case_row(ws, case)

    required_headers = ["pdf_downloaded", "pdf_download_status", "pdf_file_name", "pdf_local_path", "pdf_checked_at"]
    missing = [header for header in required_headers if header not in col_map]
    if missing:
        raise ValueError(f"Workbook is missing required PDF columns: {missing}")

    output_file = str(case.get("output_file", "") or "").strip()
    output_path = Path(output_file) if output_file else None
    pdf_exists = bool(output_path and output_path.exists() and output_path.is_file())

    status_map = {
        "downloaded_manual": "downloaded",
        "paywalled": "scout_paywalled",
        "unavailable": "scout_unavailable",
        "wrong_link": "scout_wrong_link",
        "article_not_found": "scout_article_not_found",
        "other": "scout_other",
    }
    target_status = status_map.get(decision, "scout_pending_review")

    if decision == "downloaded_manual" and (pdf_saved or pdf_exists):
        ws.cell(row=row_idx, column=col_map["pdf_downloaded"], value=1)
        ws.cell(row=row_idx, column=col_map["pdf_download_status"], value="downloaded")
        if output_path is not None:
            ws.cell(row=row_idx, column=col_map["pdf_file_name"], value=output_path.name)
            try:
                local_path = str(output_path.relative_to(ACTIVE_PROJECT_ROOT))
            except ValueError:
                local_path = str(output_path)
            ws.cell(row=row_idx, column=col_map["pdf_local_path"], value=local_path)
    else:
        ws.cell(row=row_idx, column=col_map["pdf_downloaded"], value=0)
        ws.cell(row=row_idx, column=col_map["pdf_download_status"], value=target_status)

    ws.cell(row=row_idx, column=col_map["pdf_checked_at"], value=build_scout_checked_at(decision, notes))


def apply_case_link_overrides_to_worksheet(ws, case: dict, review: dict) -> None:
    col_map = build_col_map(ws)
    row_idx = _find_case_row(ws, case)

    staged_doi = str(review.get("staged_doi", "") or "").strip()
    staged_doi_url = str(review.get("staged_doi_url", "") or "").strip()
    staged_source_url = str(review.get("staged_source_url", "") or "").strip()

    if "DOI" in col_map:
        ws.cell(row=row_idx, column=col_map["DOI"], value=staged_doi)
    if "DOI Link" in col_map:
        ws.cell(row=row_idx, column=col_map["DOI Link"], value=staged_doi_url)
    if "pdf_source_url" in col_map:
        ws.cell(row=row_idx, column=col_map["pdf_source_url"], value=staged_source_url)


def update_case_doi_link_in_workbook(case: dict, inserted_link: str) -> tuple[str, str]:
    scout_mode = str(case.get("scout_mode", "") or "").strip().lower()
    if scout_mode == "load_fail_test_cases":
        raise ValueError("Insert does not update the workbook in SCOUT test mode.")

    sheet_name = str(case.get("sheet_name", "") or "").strip()
    if not sheet_name:
        raise ValueError("SCOUT case is missing workbook location metadata.")

    doi_value, doi_link_value = normalize_inserted_doi_link(inserted_link)
    _source_workbook_path, workbook_path = get_case_workbook_paths(case)
    wb = load_workbook(workbook_path)

    try:
        if sheet_name not in wb.sheetnames:
            raise ValueError(f"Sheet not found in workbook: {sheet_name}")

        ws = wb[sheet_name]
        col_map = build_col_map(ws)
        row_idx = _find_case_row(ws, case)

        if "DOI Link" not in col_map:
            raise ValueError("Workbook is missing required column: DOI Link")
        if "DOI" not in col_map:
            raise ValueError("Workbook is missing required column: DOI")

        ws.cell(row=row_idx, column=col_map["DOI Link"], value=doi_link_value)
        if doi_value:
            ws.cell(row=row_idx, column=col_map["DOI"], value=doi_value)

        save_scout_workbook_fast(wb, workbook_path)
    finally:
        wb.close()

    return doi_value, doi_link_value


def update_case_source_link_in_workbook(case: dict, inserted_link: str) -> str:
    scout_mode = str(case.get("scout_mode", "") or "").strip().lower()
    if scout_mode == "load_fail_test_cases":
        raise ValueError("Insert does not update the workbook in SCOUT test mode.")

    sheet_name = str(case.get("sheet_name", "") or "").strip()
    if not sheet_name:
        raise ValueError("SCOUT case is missing workbook location metadata.")

    source_link_value = normalize_inserted_source_link(inserted_link)
    _source_workbook_path, workbook_path = get_case_workbook_paths(case)
    wb = load_workbook(workbook_path)

    try:
        if sheet_name not in wb.sheetnames:
            raise ValueError(f"Sheet not found in workbook: {sheet_name}")

        ws = wb[sheet_name]
        col_map = build_col_map(ws)
        row_idx = _find_case_row(ws, case)

        if "pdf_source_url" not in col_map:
            raise ValueError("Workbook is missing required column: pdf_source_url")

        ws.cell(row=row_idx, column=col_map["pdf_source_url"], value=source_link_value)

        save_scout_workbook_fast(wb, workbook_path)
    finally:
        wb.close()

    return source_link_value


def clear_case_doi_link_in_workbook(case: dict) -> None:
    scout_mode = str(case.get("scout_mode", "") or "").strip().lower()
    if scout_mode == "load_fail_test_cases":
        raise ValueError("Clear does not update the workbook in SCOUT test mode.")

    sheet_name = str(case.get("sheet_name", "") or "").strip()
    if not sheet_name:
        raise ValueError("SCOUT case is missing workbook location metadata.")

    _source_workbook_path, workbook_path = get_case_workbook_paths(case)
    wb = load_workbook(workbook_path)

    try:
        if sheet_name not in wb.sheetnames:
            raise ValueError(f"Sheet not found in workbook: {sheet_name}")

        ws = wb[sheet_name]
        col_map = build_col_map(ws)
        row_idx = _find_case_row(ws, case)

        if "DOI Link" not in col_map or "DOI" not in col_map:
            raise ValueError("Workbook is missing required DOI columns.")

        ws.cell(row=row_idx, column=col_map["DOI Link"], value="")
        ws.cell(row=row_idx, column=col_map["DOI"], value="")

        save_scout_workbook_fast(wb, workbook_path)
    finally:
        wb.close()


def clear_case_source_link_in_workbook(case: dict) -> None:
    scout_mode = str(case.get("scout_mode", "") or "").strip().lower()
    if scout_mode == "load_fail_test_cases":
        raise ValueError("Clear does not update the workbook in SCOUT test mode.")

    sheet_name = str(case.get("sheet_name", "") or "").strip()
    if not sheet_name:
        raise ValueError("SCOUT case is missing workbook location metadata.")

    _source_workbook_path, workbook_path = get_case_workbook_paths(case)
    wb = load_workbook(workbook_path)

    try:
        if sheet_name not in wb.sheetnames:
            raise ValueError(f"Sheet not found in workbook: {sheet_name}")

        ws = wb[sheet_name]
        col_map = build_col_map(ws)
        row_idx = _find_case_row(ws, case)

        if "pdf_source_url" not in col_map:
            raise ValueError("Workbook is missing required column: pdf_source_url")

        ws.cell(row=row_idx, column=col_map["pdf_source_url"], value="")

        save_scout_workbook_fast(wb, workbook_path)
    finally:
        wb.close()


def update_case_doi_link_in_state(case: dict, doi_value: str, doi_link_value: str) -> None:
    case["doi_url"] = doi_link_value
    if doi_value:
        case["doi"] = doi_value


def update_case_source_link_in_state(case: dict, source_link_value: str) -> None:
    case["source_url"] = source_link_value


def clear_case_doi_link_in_state(case: dict) -> None:
    case["doi"] = str(case.get("original_doi", "") or "").strip()
    case["doi_url"] = str(case.get("original_doi_url", "") or "").strip()


def clear_case_source_link_in_state(case: dict) -> None:
    case["source_url"] = str(case.get("original_source_url", "") or "").strip()


def clear_staged_case_link_updates(case: dict, *, clear_doi: bool = False, clear_source: bool = False) -> None:
    record_id = case["record_id"]
    decisions = st.session_state["scout_decisions"]
    existing = dict(decisions.get(record_id, {}))

    if clear_doi:
        existing.pop("staged_doi", None)
        existing.pop("staged_doi_url", None)
    if clear_source:
        existing.pop("staged_source_url", None)

    if existing:
        decisions[record_id] = existing
    else:
        decisions.pop(record_id, None)

    save_review_workbook(
        st.session_state["scout_cases"],
        decisions,
        Path(st.session_state["scout_review_path"]),
    )


def stage_case_link_updates(
    case: dict,
    *,
    doi_value: str | None = None,
    doi_url: str | None = None,
    source_url: str | None = None,
) -> None:
    record_id = case["record_id"]
    decisions = st.session_state["scout_decisions"]
    existing = dict(decisions.get(record_id, {}))

    if doi_value is not None:
        existing["staged_doi"] = doi_value
    if doi_url is not None:
        existing["staged_doi_url"] = doi_url
    if source_url is not None:
        existing["staged_source_url"] = source_url

    decisions[record_id] = existing
    save_review_workbook(
        st.session_state["scout_cases"],
        decisions,
        Path(st.session_state["scout_review_path"]),
    )


def finalize_scout_workbook_updates(cases: list[dict]) -> list[Path]:
    draft_map = st.session_state.get("scout_workbook_drafts", {})
    decisions = st.session_state.get("scout_decisions", {})
    finalized_paths: list[Path] = []
    seen_sources: set[str] = set()

    for case in cases:
        source_raw = str(case.get("workbook_path", "") or "").strip()
        if not source_raw or source_raw in seen_sources:
            continue
        seen_sources.add(source_raw)

        draft_raw = str(draft_map.get(source_raw, "") or "").strip()
        if not draft_raw:
            continue

        source_path = Path(source_raw).expanduser()
        draft_path = Path(draft_raw).expanduser()
        if not draft_path.exists():
            continue

        wb = load_workbook(draft_path)
        try:
            cases_for_workbook = [
                case
                for case in cases
                if str(case.get("workbook_path", "") or "").strip() == source_raw
            ]
            for case in cases_for_workbook:
                record_id = str(case.get("record_id", "") or "").strip()
                review = decisions.get(record_id, {})
                decision = str(review.get("decision", "pending_review") or "").strip() or "pending_review"
                notes = str(review.get("notes", "") or "").strip()
                pdf_saved = bool(review.get("pdf_saved", False))
                sheet_name = str(case.get("sheet_name", "") or "").strip()
                if not sheet_name or sheet_name not in wb.sheetnames:
                    continue
                ws = wb[sheet_name]
                apply_case_review_to_worksheet(ws, case, decision, notes, pdf_saved)
                apply_case_link_overrides_to_worksheet(ws, case, review)

            ok = safe_save_workbook(wb, source_path)
            if not ok:
                raise RuntimeError(f"Could not finalize SCOUT workbook: {source_path}")
        finally:
            wb.close()

        shutil.copy2(source_path, draft_path)
        finalized_paths.append(source_path)

    return finalized_paths


def mark_current_scout_session_finalized() -> None:
    report_path_raw = str(st.session_state.get("scout_report_path", "") or "").strip()
    if not report_path_raw:
        return

    report_path = Path(report_path_raw).expanduser()
    marker_path = finalized_marker_path_for(report_path)
    marker_path.parent.mkdir(parents=True, exist_ok=True)
    marker_path.write_text(datetime.now().strftime("%Y-%m-%d %H:%M:%S"), encoding="utf-8")


def save_case_review(case: dict, decision: str, notes: str, pdf_saved: bool) -> None:
    record_id = case["record_id"]
    decisions = st.session_state["scout_decisions"]
    decisions[record_id] = {
        "decision": decision,
        "notes": notes.strip(),
        "reviewed_at": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        "pdf_saved": pdf_saved,
    }
    save_review_workbook(
        st.session_state["scout_cases"],
        decisions,
        Path(st.session_state["scout_review_path"]),
    )


def clear_case_review(case: dict) -> None:
    record_id = case["record_id"]
    decisions = st.session_state["scout_decisions"]
    decisions.pop(record_id, None)
    save_review_workbook(
        st.session_state["scout_cases"],
        decisions,
        Path(st.session_state["scout_review_path"]),
    )


def case_has_staged_overrides(review: dict) -> bool:
    return any(
        key in review and str(review.get(key, "") or "").strip()
        for key in ("staged_doi", "staged_doi_url", "staged_source_url")
    )


def sync_case_review_from_widgets(case: dict, existing: dict) -> None:
    record_id = case["record_id"]
    decision_value = st.session_state.get(f"decision_{record_id}", existing.get("decision", "pending_review"))
    notes_value = st.session_state.get(f"notes_{record_id}", existing.get("notes", ""))
    pdf_saved_value = bool(existing.get("pdf_saved", False))
    notes_clean = str(notes_value or "").strip()
    has_staged_overrides = case_has_staged_overrides(existing)

    if decision_value == "pending_review":
        if not notes_clean and not has_staged_overrides and not pdf_saved_value:
            if record_id in st.session_state["scout_decisions"]:
                clear_case_review(case)
            return

        if (
            existing.get("decision") != "pending_review"
            or str(existing.get("notes", "") or "").strip() != notes_clean
            or bool(existing.get("pdf_saved", False)) != pdf_saved_value
        ):
            save_case_review(case, "pending_review", notes_clean, pdf_saved_value)
        elif record_id not in st.session_state["scout_decisions"]:
            save_case_review(case, "pending_review", notes_clean, pdf_saved_value)
        return

    if (
        existing.get("decision") != decision_value
        or str(existing.get("notes", "") or "").strip() != notes_clean
        or bool(existing.get("pdf_saved", False)) != pdf_saved_value
    ):
        save_case_review(case, decision_value, notes_clean, pdf_saved_value)


def flush_case_review_on_navigation(case: dict) -> None:
    record_id = case["record_id"]
    decisions = st.session_state["scout_decisions"]
    existing = decisions.get(record_id, {})
    current_decision = str(st.session_state.get(f"decision_{record_id}", existing.get("decision", "pending_review")) or "").strip()
    current_notes = str(st.session_state.get(f"notes_{record_id}", existing.get("notes", "")) or "").strip()
    current_pdf_saved = bool(existing.get("pdf_saved", False))

    existing_decision = str(existing.get("decision", "") or "").strip()
    existing_notes = str(existing.get("notes", "") or "").strip()
    existing_pdf_saved = bool(existing.get("pdf_saved", False))
    has_staged_overrides = case_has_staged_overrides(existing)

    if current_decision == "pending_review":
        if record_id not in decisions and not current_notes and not has_staged_overrides:
            return
        save_case_review(case, "pending_review", current_notes, current_pdf_saved)
        return

    if (
        existing_decision != current_decision
        or existing_notes != current_notes
        or existing_pdf_saved != current_pdf_saved
    ):
        save_case_review(case, current_decision, current_notes, current_pdf_saved)


def ensure_case_review_matches_existing_pdf(case: dict) -> None:
    output_file = str(case.get("output_file", "") or "").strip()
    if not output_file:
        return

    output_path = Path(output_file)
    if not output_path.exists() or not output_path.is_file():
        return

    record_id = case["record_id"]
    existing = st.session_state["scout_decisions"].get(record_id, {})
    existing_decision = str(existing.get("decision", "") or "").strip()
    existing_notes = str(existing.get("notes", "") or "").strip()
    existing_pdf_saved = bool(existing.get("pdf_saved", False))

    if existing_decision == "downloaded_manual" and existing_pdf_saved:
        st.session_state[f"decision_{record_id}"] = "downloaded_manual"
        return

    st.session_state[f"decision_{record_id}"] = "downloaded_manual"
    save_case_review(case, "downloaded_manual", existing_notes, True)


def get_new_download_candidates(case: dict) -> list[Path]:
    downloads_dir = Path.home() / "Downloads"
    if not downloads_dir.exists():
        return []

    record_id = case["record_id"]
    marker = st.session_state.get("scout_download_markers", {}).get(record_id)
    if marker is None:
        return []

    baseline_names = {
        str(name or "").strip().lower()
        for name in st.session_state.get("scout_download_baselines", {}).get(record_id, [])
        if str(name or "").strip()
    }

    all_pdfs = [path for path in downloads_dir.glob("*.pdf") if path.is_file()]
    new_by_name = [path for path in all_pdfs if path.name.lower() not in baseline_names]
    if new_by_name:
        return sorted(
            new_by_name,
            key=lambda path: path.stat().st_mtime,
            reverse=True,
        )

    return sorted(
        [
            path for path in all_pdfs
            if path.stat().st_mtime >= marker
        ],
        key=lambda path: path.stat().st_mtime,
        reverse=True,
    )


def import_latest_downloaded_pdf(case: dict) -> Path:
    candidates = get_new_download_candidates(case)
    record_id = case["record_id"]
    if st.session_state.get("scout_download_markers", {}).get(record_id) is None:
        raise RuntimeError("Open the DOI or source link for this case first so S.C.O.U.T. can watch for a new download.")
    if not candidates:
        raise FileNotFoundError("No new PDF downloads were found for this case since the link was opened.")

    if len(candidates) > 1:
        candidate_names = ", ".join(path.name for path in sorted(candidates, key=lambda path: path.stat().st_mtime, reverse=True)[:3])
        raise RuntimeError(
            "More than one new PDF was found in Downloads for this case. "
            f"Candidates: {candidate_names}. Use manual upload instead."
        )

    latest_pdf = candidates[0]
    output_path = Path(case["output_file"])
    output_path.parent.mkdir(parents=True, exist_ok=True)
    shutil.move(str(latest_pdf), str(output_path))
    st.session_state["scout_download_markers"].pop(record_id, None)
    st.session_state["scout_download_baselines"].pop(record_id, None)
    with DOWNLOAD_BRIDGE_STATE["lock"]:
        DOWNLOAD_BRIDGE_STATE["active_cases"].pop(record_id, None)
    return output_path


def try_auto_import_download_for_case(case: dict) -> Path | None:
    output_file = str(case.get("output_file", "") or "").strip()
    if output_file and Path(output_file).exists():
        return None

    record_id = str(case.get("record_id", "") or "").strip()
    pending_auto_saves = st.session_state.get("scout_pending_auto_saves", {})
    if record_id in pending_auto_saves:
        return None

    candidates = get_new_download_candidates(case)
    if len(candidates) != 1:
        return None

    stage_pending_auto_save(case, candidates[0])
    return candidates[0]


def try_auto_import_downloads(cases: list[dict]) -> tuple[dict | None, Path | None]:
    for case in cases:
        saved_path = try_auto_import_download_for_case(case)
        if saved_path is not None:
            return case, saved_path
    return None, None


def process_pending_auto_saves(cases: list[dict]) -> tuple[dict | None, Path | None]:
    pending_auto_saves = st.session_state.get("scout_pending_auto_saves", {})
    if not pending_auto_saves:
        return None, None

    case_map = {str(case.get("record_id", "") or "").strip(): case for case in cases}

    for record_id, info in list(pending_auto_saves.items()):
        case = case_map.get(record_id)
        if case is None:
            pending_auto_saves.pop(record_id, None)
            continue

        source_path = Path(str(info.get("source_path", "") or "").strip())
        detected_at = float(info.get("detected_at", 0.0) or 0.0)
        if not source_path.exists():
            pending_auto_saves.pop(record_id, None)
            continue
        if (time() - detected_at) < 5.0:
            continue

        saved_path = move_download_to_case(case, source_path)
        st.session_state["scout_download_markers"].pop(record_id, None)
        st.session_state["scout_download_baselines"].pop(record_id, None)
        with DOWNLOAD_BRIDGE_STATE["lock"]:
            DOWNLOAD_BRIDGE_STATE["active_cases"].pop(record_id, None)

        st.session_state[f"decision_{record_id}"] = "downloaded_manual"
        current_notes = st.session_state.get(
            f"notes_{record_id}",
            st.session_state["scout_decisions"].get(record_id, {}).get("notes", ""),
        )
        save_case_review(case, "downloaded_manual", current_notes, True)
        set_pdf_detected_feedback(case, saved_path, automatic=True)
        pending_auto_saves.pop(record_id, None)
        return case, saved_path

    return None, None


def build_progress_stats(cases: list[dict], decisions: dict[str, dict]) -> dict[str, int]:
    stats = {
        "pending": 0,
        "downloaded": 0,
        "paywalled": 0,
        "unavailable": 0,
        "other_done": 0,
        "pdf_saved": 0,
    }

    for case in cases:
        review = decisions.get(case["record_id"], {})
        decision = str(review.get("decision", "") or "").strip()
        pdf_saved = bool(review.get("pdf_saved", False))

        if not decision or decision == "pending_review":
            stats["pending"] += 1
        elif decision == "downloaded_manual":
            stats["downloaded"] += 1
        elif decision == "paywalled":
            stats["paywalled"] += 1
        elif decision in {"unavailable", "article_not_found", "wrong_link"}:
            stats["unavailable"] += 1
        else:
            stats["other_done"] += 1

        if pdf_saved:
            stats["pdf_saved"] += 1

    return stats


def render_header() -> None:
    st.markdown(
        """
        <style>
        :root {
            --scout-hero-bg: color-mix(in srgb, var(--primary-color) 18%, var(--secondary-background-color) 82%);
            --scout-summary-bg: color-mix(in srgb, var(--primary-color) 11%, var(--secondary-background-color) 89%);
            --scout-panel-bg: var(--secondary-background-color);
            --scout-hero-border: color-mix(in srgb, var(--primary-color) 34%, transparent);
            --scout-summary-border: color-mix(in srgb, var(--primary-color) 24%, transparent);
            --scout-panel-border: color-mix(in srgb, var(--text-color) 12%, transparent);
            --scout-hero-shadow: 0 10px 30px color-mix(in srgb, var(--primary-color) 12%, transparent);
            --scout-summary-shadow: 0 10px 28px color-mix(in srgb, var(--primary-color) 8%, transparent);
        }
        header[data-testid="stHeader"] {
            height: auto;
        }
        [data-testid="stToolbar"] {
            visibility: visible;
        }
        #MainMenu {
            visibility: visible;
        }
        .stAppDeployButton {
            display: none;
        }
        .block-container {
            padding-top: 3.25rem;
        }
        .scout-card {
            padding: 0.9rem 1.1rem 0.95rem 1.1rem;
            border: 1px solid var(--scout-hero-border);
            border-radius: 18px;
            background: var(--scout-hero-bg);
            box-shadow: var(--scout-hero-shadow);
            margin-top: 0;
            margin-bottom: 0.75rem;
        }
        .scout-title {
            font-size: 1.95rem;
            font-weight: 800;
            margin: 0.1rem 0 0.2rem 0;
        }
        .scout-title-accent {
            color: #7dd3fc;
        }
        .scout-title-secondary {
            color: inherit;
            font-size: 1.15rem;
            font-weight: 700;
            margin: 0 0 0.45rem 0;
            line-height: 1.25;
        }
        .scout-sub {
            color: #aab4bf;
            margin: 0;
        }
        .scout-stats {
            padding: 0.95rem 1.1rem 1rem 1.1rem;
            border: 1px solid var(--scout-summary-border);
            border-radius: 18px;
            background: var(--scout-summary-bg);
            box-shadow: var(--scout-summary-shadow);
            margin-bottom: 1rem;
        }
        .scout-panel {
            padding: 0.95rem 1.05rem 1rem 1.05rem;
            border: 1px solid var(--scout-panel-border);
            border-radius: 18px;
            background: var(--scout-panel-bg);
            margin-bottom: 1rem;
        }
        .scout-panel-title {
            color: inherit;
            font-size: 1rem;
            font-weight: 700;
            margin-bottom: 0.8rem;
        }
        .scout-stats-grid {
            display: grid;
            grid-template-columns: repeat(3, minmax(0, 1fr));
            gap: 1rem;
        }
        .scout-stat-label {
            color: inherit;
            opacity: 0.78;
            font-size: 0.95rem;
            font-weight: 600;
            margin-bottom: 0.45rem;
        }
        .scout-stat-value {
            color: inherit;
            font-size: 2rem;
            font-weight: 700;
            line-height: 1;
        }
        .scout-section-heading {
            color: #67b8ff;
            font-size: 1rem;
            font-weight: 800;
            margin: 0 0 1rem 0;
            padding-bottom: 0.6rem;
            border-bottom: 1px solid rgba(120, 130, 150, 0.28);
            letter-spacing: 0.01em;
        }
        .scout-status-pulse {
            display: inline-flex;
            align-items: center;
            gap: 0.35rem;
            margin: 0 0 0.95rem 0;
            padding: 0.38rem 0.72rem;
            border-radius: 999px;
            background: rgba(34, 197, 94, 0.12);
            color: #16a34a;
            font-size: 0.92rem;
            font-weight: 800;
            animation: scoutPulse 1s ease-in-out infinite;
        }
        @keyframes scoutPulse {
            0% { transform: scale(1); box-shadow: 0 0 0 0 rgba(34, 197, 94, 0.24); }
            50% { transform: scale(1.02); box-shadow: 0 0 0 10px rgba(34, 197, 94, 0.0); }
            100% { transform: scale(1); box-shadow: 0 0 0 0 rgba(34, 197, 94, 0.0); }
        }
        .scout-article-title {
            color: inherit;
            font-size: clamp(1.7rem, 2.3vw, 2.5rem);
            font-weight: 800;
            line-height: 1.1;
            letter-spacing: -0.03em;
            margin: 0 0 1rem 0;
            width: 100%;
        }
        .scout-meta-list {
            display: grid;
            gap: 0.65rem;
            margin-bottom: 0.75rem;
            min-width: 0;
        }
        .scout-meta-row {
            color: inherit;
            font-size: 0.98rem;
            line-height: 1.45;
            display: grid;
            grid-template-columns: auto minmax(0, 1fr);
            align-items: start;
            column-gap: 0.35rem;
            min-width: 0;
        }
        .scout-meta-label {
            font-weight: 700;
        }
        .scout-meta-value {
            font-weight: 400;
            color: inherit;
            text-decoration: none;
            font-family: inherit;
            white-space: normal;
            overflow-wrap: anywhere;
            word-break: break-word;
            min-width: 0;
        }
        .scout-meta-value-code {
            font-family: var(--font-monospace, "SFMono-Regular", Consolas, monospace);
            font-size: 0.92em;
            color: #198f6b;
            background: rgba(120, 130, 150, 0.08);
            border-radius: 0.35rem;
            padding: 0.08rem 0.32rem;
            display: inline-block;
            text-decoration: none;
            white-space: normal;
            overflow-wrap: anywhere;
            word-break: break-word;
            max-width: 100%;
            min-width: 0;
        }
        .scout-stat-list {
            display: grid;
            gap: 0.55rem;
            max-width: 24rem;
            min-height: 18.5rem;
            padding-bottom: 0.45rem;
        }
        .scout-stat-row {
            display: grid;
            grid-template-columns: minmax(0, 1fr) auto;
            align-items: center;
            column-gap: 0.85rem;
            padding: 0.1rem 0;
        }
        .scout-stat-row-label {
            font-size: 0.95rem;
            font-weight: 600;
            color: inherit;
            line-height: 1.25;
        }
        .scout-stat-row-value {
            min-width: 2rem;
            text-align: center;
            font-size: 1rem;
            font-weight: 700;
            color: #4f8f63;
            line-height: 1;
            background: rgba(120, 130, 150, 0.08);
            border-radius: 0.45rem;
            padding: 0.18rem 0.42rem;
            display: inline-flex;
            align-items: center;
            justify-content: center;
            justify-self: end;
        }
        </style>
        """,
        unsafe_allow_html=True,
    )
    st.markdown(
        """
        <div class="scout-card">
            <div class="scout-title"><span class="scout-title-accent">S.C.O.U.T.</span></div>
            <div class="scout-title-secondary">Semi-assisted Case Opening and User Triage</div>
            <p class="scout-sub">
                Review failed cases one by one, open the article link in your browser,
                match PDF downloads back to the current case, and classify unresolved cases with
                structured decisions.
            </p>
        </div>
        """,
        unsafe_allow_html=True,
    )
def render_loader() -> None:
    if st.session_state.get("scout_loaded") and get_active_scout_mode() == "pending_downloads":
        return

    default_path = r"C:\Users\Luís Pinto Coelho\Desktop\resolver_tests\wiley\reports\wiley_failed_cases_2026-04-06_14-46-15.xlsx"
    with st.sidebar:
        with st.expander("Report source", expanded=st.session_state["scout_report_source_expanded"]):
            report_path_value = st.text_input("Failed-cases Excel path", value=default_path)
            if st.button("Load failed cases", use_container_width=True):
                try:
                    load_report(report_path_value)
                    st.session_state["scout_report_source_expanded"] = False
                    st.success("Failed-case report loaded.")
                    st.rerun()
                except Exception as exc:
                    st.session_state["scout_report_source_expanded"] = True
                    st.error(f"Could not load report: {exc}")


def render_sidebar(case: dict, total_cases: int) -> None:
    idx = st.session_state["scout_current_index"]
    decisions = st.session_state["scout_decisions"]
    existing = decisions.get(case["record_id"], {})

    with st.sidebar:
        st.markdown("# :material/explore: Navigator")
        st.caption(f"Case {idx + 1} of {total_cases}")

        st.progress((idx + 1) / total_cases if total_cases else 0.0)

        col_prev, col_next = st.columns(2)
        with col_prev:
            if render_navigation_button(
                label="Previous",
                key=f"previous_sidebar_{case['record_id']}",
                icon_side="right",
                icon=":material/chevron_left:",
                disabled=idx <= 0,
            ):
                flush_case_review_on_navigation(case)
                st.session_state["scout_current_index"] = max(0, idx - 1)
                st.rerun()
        with col_next:
            if render_navigation_button(
                label="Next",
                key=f"next_sidebar_{case['record_id']}",
                icon=":material/chevron_right:",
                disabled=idx >= total_cases - 1,
            ):
                flush_case_review_on_navigation(case)
                st.session_state["scout_current_index"] = min(total_cases - 1, idx + 1)
                st.rerun()

        with st.expander(":material/tune: Advanced Navigation", expanded=False):
            col_first, col_last = st.columns(2)
            with col_first:
                if st.button("First", use_container_width=True, disabled=idx <= 0, icon=":material/first_page:"):
                    flush_case_review_on_navigation(case)
                    st.session_state["scout_current_index"] = 0
                    st.rerun()
            with col_last:
                if st.button("Last", use_container_width=True, disabled=idx >= total_cases - 1, icon=":material/last_page:"):
                    flush_case_review_on_navigation(case)
                    st.session_state["scout_current_index"] = max(0, total_cases - 1)
                    st.rerun()

            first_pending_index = find_resume_index(st.session_state["scout_cases"], decisions)
            if st.button(
                "Go to first pending",
                use_container_width=True,
                icon=":material/resume:",
                disabled=first_pending_index == idx,
            ):
                flush_case_review_on_navigation(case)
                st.session_state["scout_current_index"] = first_pending_index
                st.rerun()

            with st.form(key="scout_go_to_case_form", clear_on_submit=False):
                go_to_case = st.number_input(
                    "Go to case",
                    min_value=1,
                    max_value=max(1, total_cases),
                    value=idx + 1,
                    step=1,
                    format="%d",
                    key="scout_go_to_case",
                )
                go_to_case_submitted = st.form_submit_button("Go", use_container_width=True, icon=":material/arrow_forward:")

            if go_to_case_submitted and int(go_to_case) != idx + 1:
                flush_case_review_on_navigation(case)
                st.session_state["scout_current_index"] = int(go_to_case) - 1
                st.rerun()

        if render_open_article_button(case):
            launch_url, _launch_url_kind = get_preferred_case_url(case)
            try:
                launch_scout_browser(launch_url, case)
                st.session_state["scout_browser_message"] = ""
            except Exception as exc:
                st.session_state["scout_browser_message"] = f"Could not open article link: {exc}"
            st.rerun()

        render_google_search_button(case)
        render_request_email_button(case)

        with st.expander(":material/add_link: Insert", expanded=False):
            insert_mode = st.radio(
                "Insert type",
                ["DOI / DOI link", "Source link"],
                key=f"insert_mode_{case['record_id']}",
                label_visibility="collapsed",
            )

            if insert_mode == "DOI / DOI link":
                link_input = st.text_input(
                    "New DOI / DOI link",
                    value="",
                    key=f"insert_doi_link_{case['record_id']}",
                    placeholder="https://doi.org/... or 10.xxxx/...",
                )
                col_save_doi, col_clear_doi = st.columns(2)
                with col_save_doi:
                    if st.button("Save", use_container_width=True, key=f"save_inserted_doi_link_{case['record_id']}", icon=":material/save:"):
                        try:
                            doi_value, doi_link_value = normalize_inserted_doi_link(link_input)
                            stage_case_link_updates(case, doi_value=doi_value, doi_url=doi_link_value)
                            update_case_doi_link_in_state(case, doi_value, doi_link_value)
                            st.success("DOI / DOI Link staged in SCOUT.")
                            st.rerun()
                        except Exception as exc:
                            st.error(f"Could not update DOI / DOI link: {exc}")
                with col_clear_doi:
                    if st.button("Clear", use_container_width=True, key=f"clear_inserted_doi_link_{case['record_id']}", icon=":material/delete:"):
                        try:
                            clear_staged_case_link_updates(case, clear_doi=True)
                            clear_case_doi_link_in_state(case)
                            st.success("DOI / DOI Link cleared in SCOUT.")
                            st.rerun()
                        except Exception as exc:
                            st.error(f"Could not clear DOI / DOI link: {exc}")
            else:
                source_input = st.text_input(
                    "New source link",
                    value="",
                    key=f"insert_source_link_{case['record_id']}",
                    placeholder="https://publisher-site/article...",
                )
                col_save_source, col_clear_source = st.columns(2)
                with col_save_source:
                    if st.button("Save", use_container_width=True, key=f"save_inserted_source_link_{case['record_id']}", icon=":material/save:"):
                        try:
                            source_link_value = normalize_inserted_source_link(source_input)
                            stage_case_link_updates(case, source_url=source_link_value)
                            update_case_source_link_in_state(case, source_link_value)
                            st.success("Source link staged in SCOUT.")
                            st.rerun()
                        except Exception as exc:
                            st.error(f"Could not update source link: {exc}")
                with col_clear_source:
                    if st.button("Clear", use_container_width=True, key=f"clear_inserted_source_link_{case['record_id']}", icon=":material/delete:"):
                        try:
                            clear_staged_case_link_updates(case, clear_source=True)
                            clear_case_source_link_in_state(case)
                            st.success("Source link cleared in SCOUT.")
                            st.rerun()
                        except Exception as exc:
                            st.error(f"Could not clear source link: {exc}")

        st.header(":material/rule: Decision")
        default_decision = existing.get("decision", "pending_review")
        decision_index = DECISION_OPTIONS.index(default_decision) if default_decision in DECISION_OPTIONS else 0
        decision = st.radio(
            "Decision",
            DECISION_OPTIONS,
            index=decision_index,
            key=f"decision_{case['record_id']}",
            label_visibility="collapsed",
        )
        notes = st.text_area(
            ":material/stylus_note: Notes",
            value=existing.get("notes", ""),
            height=140,
            key=f"notes_{case['record_id']}",
        )
        sync_case_review_from_widgets(case, existing)
        progress_stats = build_progress_stats(st.session_state["scout_cases"], decisions)
        if st.button(":material/bar_chart: Statistics", use_container_width=True, key=f"statistics_sidebar_{case['record_id']}"):
            render_statistics_dialog(progress_stats)
        finalize_disabled = progress_stats["pending"] > 0
        finalize_help = (
            "Finalize becomes available after all cases have been reviewed."
            if finalize_disabled
            else "Sync the SCOUT draft back to the workbook and verify the PDF folder."
        )
        if st.button(
            ":material/save: Finalize To Workbook",
            use_container_width=True,
            key=f"finalize_sidebar_{case['record_id']}",
            disabled=finalize_disabled,
            help=finalize_help,
        ):
            try:
                flush_case_review_on_navigation(case)
                finalized_paths = finalize_scout_workbook_updates(st.session_state["scout_cases"])
                if finalized_paths:
                    mark_current_scout_session_finalized()
                    st.session_state["scout_sync_message"] = (
                        f"SCOUT draft synced to workbook for {len(finalized_paths)} file(s). "
                        "PDF folder verification was applied during finalization."
                    )
                else:
                    st.session_state["scout_sync_message"] = "No SCOUT draft changes were available to finalize."
                st.rerun()
            except Exception as exc:
                st.error(f"Could not finalize workbook updates: {exc}")
        sync_message = str(st.session_state.get("scout_sync_message", "") or "").strip()
        if sync_message:
            st.caption(sync_message)



def render_case(case: dict) -> None:
    decisions = st.session_state["scout_decisions"]
    existing = decisions.get(case["record_id"], {})
    bridge_events = get_bridge_events_for_case(case["record_id"])
    launch_url, _launch_url_kind = get_preferred_case_url(case)
    has_side_content = bool(bridge_events)

    if has_side_content:
        col_main, col_side = st.columns([4.8, 1.2], vertical_alignment="top")
    else:
        col_main, = st.columns([1], vertical_alignment="top")
        col_side = None

    with col_main:
        with st.container(border=True):
            st.markdown('<div class="scout-section-heading">Article Information</div>', unsafe_allow_html=True)
            display_title = normalize_display_title(case["title"])
            st.markdown(
                f'<div class="scout-article-title">{escape(display_title)}</div>',
                unsafe_allow_html=True,
            )
            def render_meta_row(label: str, value: str, *, code_style: bool = False) -> str:
                css_class = "scout-meta-value-code" if code_style else "scout-meta-value"
                return (
                    f'<div class="scout-meta-row"><span class="scout-meta-label">{escape(label)}:</span> '
                    f'<span class="{css_class}">{escape(str(value or "-"))}</span></div>'
                )

            meta_html = "".join(
                [
                    render_meta_row("Record ID", case.get("record_id") or "-", code_style=True),
                    render_meta_row("Authors", case.get("authors") or "-", code_style=True),
                    render_meta_row("Publication Year", case.get("publication_year") or "-", code_style=True),
                    render_meta_row("E-mail Address", case.get("email_address") or "-", code_style=True),
                    render_meta_row("DOI", case.get("doi") or "-", code_style=True),
                    render_meta_row("DOI Link", case.get("doi_url") or "-", code_style=True),
                    render_meta_row("Source Link", case.get("source_url") or "-", code_style=True),
                ]
            )
            st.markdown(f'<div class="scout-meta-list">{meta_html}</div>', unsafe_allow_html=True)

            browser_message = st.session_state.get("scout_browser_message", "")
            if browser_message:
                st.info(browser_message)
            download_message = str(st.session_state.get("scout_download_message", "") or "").strip()
            download_message_record_id = str(st.session_state.get("scout_download_message_record_id", "") or "").strip()
            if download_message and download_message_record_id == str(case.get("record_id", "") or "").strip():
                st.success(download_message)

            output_path = Path(case["output_file"]) if case["output_file"] else None
            candidates = get_new_download_candidates(case)
            marker_exists = st.session_state.get("scout_download_markers", {}).get(case["record_id"]) is not None
            active_info = get_active_case_info(case["record_id"])
            pending_auto_save = st.session_state.get("scout_pending_auto_saves", {}).get(case["record_id"], {})
            detected_pending = bool(pending_auto_save)
            flash_record_id = str(st.session_state.get("scout_pdf_status_flash_record_id", "") or "").strip()
            flash_until = float(st.session_state.get("scout_pdf_status_flash_until", 0.0) or 0.0)
            flash_active = flash_record_id == str(case["record_id"]) and flash_until > time()

            with st.container(border=True):
                st.markdown('<div class="scout-section-heading">Error Information</div>', unsafe_allow_html=True)
                error_text = (
                    f"**Original error:** `{case.get('error_type') or '-'}`\n\n"
                    f"**Error detail:** `{case.get('error_message') or '-'}`"
                )
                st.info(error_text)

            with st.container(border=True):
                st.markdown('<div class="scout-section-heading">PDF Status</div>', unsafe_allow_html=True)
                if flash_active:
                    st.markdown('<div class="scout-status-pulse">PDF detected</div>', unsafe_allow_html=True)
                if output_path and output_path.exists():
                    st.success(f"Saved PDF detected at `{output_path}`")
                elif detected_pending:
                    remaining_save = max(0, int(5 - (time() - float(pending_auto_save.get("detected_at", 0.0) or 0.0))))
                    st.success(f"PDF detected. Saving automatically in {remaining_save}s.")
                elif candidates:
                    candidate_names = ", ".join(path.name for path in candidates[:3])
                    st.warning(f"PDF detected in Downloads and ready to import: `{candidate_names}`")
                elif marker_exists or active_info:
                    st.info("Detecting PDF download...")
                elif output_path:
                    st.info("No PDF has been detected for this case yet.")
                else:
                    st.info("This case does not have an expected PDF output path yet.")

    if col_side is not None:
        with col_side:
            if bridge_events:
                with st.container(border=True):
                    st.markdown('<div class="scout-section-heading">Recent Downloads</div>', unsafe_allow_html=True)
                    for event in bridge_events[-3:]:
                        status = event.get("status", "pending")
                        filename = Path(str(event.get("filename", "") or "")).name or "(unknown file)"
                        st.caption(f"`{status}` - `{filename}`")
                        if status == "matched" and event.get("saved_path"):
                            st.write(f"Saved to `{event['saved_path']}`")
                        if status == "ambiguous":
                            if st.button("Use this download for current case", use_container_width=True, key=f"assign_event_{event['event_index']}"):
                                try:
                                    saved_path = assign_download_event_to_case(event["event_index"], case)
                                    st.success(f"Download assigned to `{saved_path}`")
                                    st.rerun()
                                except Exception as exc:
                                    st.error(f"Could not assign download: {exc}")


def render_summary() -> None:
    decisions = st.session_state["scout_decisions"]
    cases = st.session_state["scout_cases"]
    total = len(cases)
    reviewed = sum(
        1
        for case in cases
        if str(decisions.get(case["record_id"], {}).get("decision", "") or "").strip()
        and str(decisions.get(case["record_id"], {}).get("decision", "") or "").strip() != "pending_review"
    )
    downloaded = sum(1 for value in decisions.values() if value.get("decision") == "downloaded_manual")

    st.markdown(
        f"""
        <div class="scout-stats">
            <div class="scout-stats-grid">
                <div>
                    <div class="scout-stat-label">Cases</div>
                    <div class="scout-stat-value">{total}</div>
                </div>
                <div>
                    <div class="scout-stat-label">Reviewed</div>
                    <div class="scout-stat-value">{reviewed}</div>
                </div>
                <div>
                    <div class="scout-stat-label">Manual downloads</div>
                    <div class="scout-stat-value">{downloaded}</div>
                </div>
            </div>
        </div>
        """,
        unsafe_allow_html=True,
    )


@st.dialog("Statistics")
def render_statistics_dialog(progress_stats: dict[str, int]) -> None:
    st.markdown(
        f"""
        <div class="scout-stat-list">
            <div class="scout-stat-row">
                <div class="scout-stat-row-label">Pending</div>
                <div class="scout-stat-row-value">{progress_stats['pending']}</div>
            </div>
            <div class="scout-stat-row">
                <div class="scout-stat-row-label">Downloaded</div>
                <div class="scout-stat-row-value">{progress_stats['downloaded']}</div>
            </div>
            <div class="scout-stat-row">
                <div class="scout-stat-row-label">Paywalled</div>
                <div class="scout-stat-row-value">{progress_stats['paywalled']}</div>
            </div>
            <div class="scout-stat-row">
                <div class="scout-stat-row-label">Unavailable / wrong link</div>
                <div class="scout-stat-row-value">{progress_stats['unavailable']}</div>
            </div>
            <div class="scout-stat-row">
                <div class="scout-stat-row-label">Other completed</div>
                <div class="scout-stat-row-value">{progress_stats['other_done']}</div>
            </div>
            <div class="scout-stat-row">
                <div class="scout-stat-row-label">PDFs saved</div>
                <div class="scout-stat-row-value">{progress_stats['pdf_saved']}</div>
            </div>
        </div>
        """,
        unsafe_allow_html=True,
    )


def should_auto_refresh_case(case: dict) -> tuple[bool, int]:
    output_file = str(case.get("output_file", "") or "").strip()
    if output_file and Path(output_file).exists():
        return False, 0

    active_info = get_active_case_info(case["record_id"])
    if not active_info:
        return False, 0

    remaining = int(max(0, float(active_info.get("watch_until") or 0) - time()))
    if remaining <= 0:
        return False, 0

    return True, remaining


@st.fragment(run_every=1)
def render_controlled_auto_refresh(case: dict) -> None:
    should_refresh, remaining = should_auto_refresh_case(case)
    cases = st.session_state.get("scout_cases", [])
    process_download_bridge_events(cases)
    detected_case, detected_path = try_auto_import_downloads(cases)
    if detected_path is not None and detected_case is not None:
        st.rerun()
    saved_case, saved_path = process_pending_auto_saves(cases)
    if saved_path is not None and saved_case is not None:
        st.rerun()

    if not should_refresh:
        return

    st.caption(f"Watching Downloads for this case... {remaining}s remaining")


def render_keyboard_navigation() -> None:
    return


def main() -> None:
    set_page()
    init_state()
    try_auto_load_from_environment()
    start_download_bridge_server()
    prune_expired_active_cases()
    render_header()
    render_loader()

    if not st.session_state["scout_loaded"]:
        st.info("Load a failed-case Excel report to start reviewing cases.")
        return

    cases = st.session_state["scout_cases"]
    process_download_bridge_events(cases)
    auto_saved_case, auto_saved_path = try_auto_import_downloads(cases)
    if auto_saved_path is not None and auto_saved_case is not None:
        st.rerun()
    saved_case, saved_path = process_pending_auto_saves(cases)
    if saved_path is not None and saved_case is not None:
        st.rerun()

    case = get_current_case()
    if case is None:
        st.warning("No cases available in the loaded report.")
        return

    ensure_case_review_matches_existing_pdf(case)
    render_keyboard_navigation()
    render_sidebar(case, len(cases))
    render_summary()
    render_case(case)
    render_controlled_auto_refresh(case)


if __name__ == "__main__":
    main()
