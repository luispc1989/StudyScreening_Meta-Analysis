from __future__ import annotations

import os
import socket
import subprocess
import sys
import time
import webbrowser
from datetime import datetime
from pathlib import Path
from typing import Any, Literal

from openpyxl import Workbook
from openpyxl import load_workbook

from tools.pdf_fetcher.core.config import REPORTS_OUTPUT_DIR
from tools.pdf_fetcher.core.excel_io import (
    build_col_map,
    get_optional_cell,
    reconcile_pdf_state_with_filesystem,
)
from tools.pdf_fetcher.core.models import SessionState
from tools.pdf_fetcher.core.utils import build_pdf_filename, make_doi_url, normalize_doi


ScoutMode = Literal["pending_downloads", "load_fail_test_cases"]
SCOUT_STREAMLIT_PORT = 8511

SCOUT_REPORT_HEADERS = [
    "record_id",
    "title",
    "authors",
    "publication_year",
    "email_address",
    "doi",
    "doi_url",
    "source_url",
    "output_file",
    "error_type",
    "error_message",
    "scout_mode",
    "workbook_path",
    "sheet_name",
    "source_row_idx",
]

SCOUT_METADATA_HEADERS = ["authors", "publication_year", "email_address"]


def review_output_path_for(report_path: Path) -> Path:
    target_dir = report_path.parent / "scout_reviews"
    target_dir.mkdir(parents=True, exist_ok=True)
    return target_dir / f"{report_path.stem}_scout_review.xlsx"


def finalized_marker_path_for(report_path: Path) -> Path:
    target_dir = report_path.parent / "scout_reviews"
    target_dir.mkdir(parents=True, exist_ok=True)
    return target_dir / f"{report_path.stem}_finalized.flag"


def _is_local_port_available(port: int) -> bool:
    with socket.socket(socket.AF_INET, socket.SOCK_STREAM) as sock:
        sock.setsockopt(socket.SOL_SOCKET, socket.SO_REUSEADDR, 1)
        try:
            sock.bind(("127.0.0.1", port))
        except OSError:
            return False
    return True


def _pick_available_port(preferred_port: int = SCOUT_STREAMLIT_PORT, max_attempts: int = 25) -> int:
    for port in range(preferred_port, preferred_port + max_attempts):
        if _is_local_port_available(port):
            return port

    with socket.socket(socket.AF_INET, socket.SOCK_STREAM) as sock:
        sock.bind(("127.0.0.1", 0))
        return int(sock.getsockname()[1])


def _wait_for_server_start(
    port: int,
    process: subprocess.Popen,
    timeout_seconds: float = 20.0,
) -> bool:
    deadline = time.time() + timeout_seconds
    while time.time() < deadline:
        if process.poll() is not None:
            return False

        with socket.socket(socket.AF_INET, socket.SOCK_STREAM) as sock:
            sock.settimeout(0.5)
            if sock.connect_ex(("127.0.0.1", port)) == 0:
                return True

        time.sleep(0.25)

    return False


def _coerce_binary_flag(value) -> int:
    try:
        return 1 if int(str(value or "0").strip()) == 1 else 0
    except Exception:
        return 0


def _status_matches_mode(status: str, mode: ScoutMode) -> bool:
    normalized = str(status or "").strip().lower()
    if mode == "pending_downloads":
        return True
    if mode == "load_fail_test_cases":
        return normalized.startswith("download_failed_") or "load_fail" in normalized or "failed" in normalized
    return False


def build_scout_cases_from_session(
    session_state: SessionState,
    mode: ScoutMode = "pending_downloads",
) -> list[dict]:
    if not session_state.is_workbook_loaded:
        raise ValueError("SessionState workbook is not loaded.")

    ws = session_state.worksheet
    col_map = build_col_map(ws)
    reconcile_pdf_state_with_filesystem(ws, col_map)
    col_map = build_col_map(ws)

    cases: list[dict] = []

    for row_idx in range(2, ws.max_row + 1):
        record_id = str(get_optional_cell(ws, row_idx, col_map, "record_id", "") or "").strip()
        if not record_id:
            continue

        pdf_downloaded = _coerce_binary_flag(get_optional_cell(ws, row_idx, col_map, "pdf_downloaded", 0))
        status = str(get_optional_cell(ws, row_idx, col_map, "pdf_download_status", "") or "").strip()

        if pdf_downloaded == 1:
            continue
        if not _status_matches_mode(status, mode):
            continue

        title = str(get_optional_cell(ws, row_idx, col_map, "Title", "") or "").strip()
        authors = str(get_optional_cell(ws, row_idx, col_map, "Authors", "") or "").strip()
        publication_year = str(get_optional_cell(ws, row_idx, col_map, "Publication Year", "") or "").strip()
        email_address = str(
            get_optional_cell(
                ws,
                row_idx,
                col_map,
                "E-mail Address",
                get_optional_cell(ws, row_idx, col_map, "E-mail Adress", ""),
            )
            or ""
        ).strip()
        doi = str(get_optional_cell(ws, row_idx, col_map, "DOI", "") or "").strip()
        doi_link = str(get_optional_cell(ws, row_idx, col_map, "DOI Link", "") or "").strip()
        normalized_doi = normalize_doi(doi) or normalize_doi(doi_link) or ""
        doi_url = doi_link or (make_doi_url(normalized_doi) if normalized_doi else "")
        source_url = str(get_optional_cell(ws, row_idx, col_map, "pdf_source_url", "") or "").strip()
        error_message = str(get_optional_cell(ws, row_idx, col_map, "pdf_checked_at", "") or "").strip()
        output_file = str((session_state.tool_root / "PDFs" / build_pdf_filename(record_id, title)) if session_state.tool_root else "")

        cases.append(
            {
                "record_id": record_id,
                "title": title,
                "authors": authors,
                "publication_year": publication_year,
                "email_address": email_address,
                "doi": normalized_doi or doi,
                "doi_url": doi_url,
                "source_url": source_url,
                "output_file": output_file,
                "error_type": status or ("pending_download" if mode == "pending_downloads" else "load_fail_test_case"),
                "error_message": error_message,
                "scout_mode": mode,
                "workbook_path": str(session_state.current_workbook_path or session_state.workbook_path or ""),
                "sheet_name": str(session_state.sheet_name or ""),
                "source_row_idx": row_idx,
            }
        )

    return cases


def _report_matches_session(report_path: Path, session_state: SessionState, mode: ScoutMode) -> bool:
    if not report_path.exists():
        return False

    workbook_path = str(session_state.current_workbook_path or session_state.workbook_path or "").strip()
    if not workbook_path:
        return False

    try:
        wb = load_workbook(report_path, read_only=True, data_only=True)
    except Exception:
        return False

    try:
        ws = wb.active
        headers = [str(cell.value or "").strip() for cell in ws[1]]
        normalized_headers = [header.strip().lower() for header in headers]
        header_map = {normalized_headers[idx]: idx for idx in range(len(normalized_headers))}

        workbook_idx = header_map.get("workbook_path")
        mode_idx = header_map.get("scout_mode")
        if workbook_idx is None or mode_idx is None:
            return False

        for row in ws.iter_rows(min_row=2, values_only=True):
            if not any(row):
                continue
            row_workbook = str(row[workbook_idx] or "").strip()
            row_mode = str(row[mode_idx] or "").strip().lower()
            if row_workbook == workbook_path and row_mode == mode:
                return True
        return False
    finally:
        wb.close()


def ensure_report_metadata_columns(report_path: Path) -> None:
    if not report_path.exists():
        return

    wb = load_workbook(report_path)
    try:
        ws = wb.active
        headers = [str(cell.value or "").strip() for cell in ws[1]]
        normalized_headers = [header.strip().lower() for header in headers]
        header_map = {normalized_headers[idx]: idx + 1 for idx in range(len(normalized_headers))}

        missing_headers = [header for header in SCOUT_METADATA_HEADERS if header not in header_map]
        if not missing_headers:
            return

        for header in missing_headers:
            ws.cell(row=1, column=ws.max_column + 1, value=header)

        headers = [str(cell.value or "").strip() for cell in ws[1]]
        normalized_headers = [header.strip().lower() for header in headers]
        header_map = {normalized_headers[idx]: idx + 1 for idx in range(len(normalized_headers))}

        workbook_idx = header_map.get("workbook_path")
        sheet_idx = header_map.get("sheet_name")
        row_idx_col = header_map.get("source_row_idx")
        record_id_idx = header_map.get("record_id")
        if not workbook_idx or not sheet_idx or not row_idx_col or not record_id_idx:
            wb.save(report_path)
            return

        workbook_cache: dict[str, tuple[Any, dict[str, int]]] = {}

        def get_sheet(workbook_path_raw: str, sheet_name: str):
            cache_key = f"{workbook_path_raw}::{sheet_name}"
            if cache_key in workbook_cache:
                return workbook_cache[cache_key]

            workbook_path = Path(workbook_path_raw).expanduser()
            if not workbook_path.exists():
                workbook_cache[cache_key] = (None, {})
                return workbook_cache[cache_key]

            try:
                source_wb = load_workbook(workbook_path, read_only=True, data_only=True)
            except Exception:
                workbook_cache[cache_key] = (None, {})
                return workbook_cache[cache_key]

            if sheet_name not in source_wb.sheetnames:
                source_wb.close()
                workbook_cache[cache_key] = (None, {})
                return workbook_cache[cache_key]

            source_ws = source_wb[sheet_name]
            source_col_map = build_col_map(source_ws)
            workbook_cache[cache_key] = (source_ws, source_col_map)
            return workbook_cache[cache_key]

        try:
            for excel_row_idx in range(2, ws.max_row + 1):
                workbook_path_raw = str(ws.cell(row=excel_row_idx, column=workbook_idx).value or "").strip()
                sheet_name = str(ws.cell(row=excel_row_idx, column=sheet_idx).value or "").strip()
                source_row_idx_raw = ws.cell(row=excel_row_idx, column=row_idx_col).value
                record_id = str(ws.cell(row=excel_row_idx, column=record_id_idx).value or "").strip()

                if not workbook_path_raw or not sheet_name:
                    continue

                source_ws, source_col_map = get_sheet(workbook_path_raw, sheet_name)
                if source_ws is None:
                    continue

                try:
                    source_row_idx = int(source_row_idx_raw or 0)
                except Exception:
                    source_row_idx = 0

                record_col = source_col_map.get("record_id")
                if source_row_idx >= 2 and record_col:
                    source_record_id = str(source_ws.cell(row=source_row_idx, column=record_col).value or "").strip()
                    if source_record_id != record_id:
                        source_row_idx = 0

                if source_row_idx < 2 and record_col:
                    for candidate_row_idx in range(2, source_ws.max_row + 1):
                        candidate_record_id = str(source_ws.cell(row=candidate_row_idx, column=record_col).value or "").strip()
                        if candidate_record_id == record_id:
                            source_row_idx = candidate_row_idx
                            break

                if source_row_idx < 2:
                    continue

                ws.cell(
                    row=excel_row_idx,
                    column=header_map["authors"],
                    value=str(get_optional_cell(source_ws, source_row_idx, source_col_map, "Authors", "") or "").strip(),
                )
                ws.cell(
                    row=excel_row_idx,
                    column=header_map["publication_year"],
                    value=str(get_optional_cell(source_ws, source_row_idx, source_col_map, "Publication Year", "") or "").strip(),
                )
                ws.cell(
                    row=excel_row_idx,
                    column=header_map["email_address"],
                    value=str(
                        get_optional_cell(
                            source_ws,
                            source_row_idx,
                            source_col_map,
                            "E-mail Address",
                            get_optional_cell(source_ws, source_row_idx, source_col_map, "E-mail Adress", ""),
                        )
                        or ""
                    ).strip(),
                )
        finally:
            for source_ws, _source_col_map in workbook_cache.values():
                if source_ws is not None:
                    source_ws.parent.close()

        wb.save(report_path)
    finally:
        wb.close()


def find_resumable_scout_report(session_state: SessionState, mode: ScoutMode) -> Path | None:
    repo_root = Path(__file__).resolve().parents[3]
    scout_report_dirs = [
        REPORTS_OUTPUT_DIR / "SCOUT",
        repo_root / "tools" / "pdf_fetcher" / "runtime" / "reports" / "SCOUT",
    ]

    candidates: list[Path] = []
    for scout_reports_dir in scout_report_dirs:
        if not scout_reports_dir.exists():
            continue
        candidates.extend(sorted(scout_reports_dir.glob(f"scout_{mode}_*.xlsx")))

    resumable: list[tuple[Path, float]] = []
    for report_path in candidates:
        review_path = review_output_path_for(report_path)
        marker_path = finalized_marker_path_for(report_path)
        if marker_path.exists() or not review_path.exists():
            continue
        if not _report_matches_session(report_path, session_state, mode):
            continue
        try:
            review_mtime = review_path.stat().st_mtime
        except OSError:
            review_mtime = 0.0
        resumable.append((report_path, review_mtime))

    if not resumable:
        return None

    return max(resumable, key=lambda item: item[1])[0]


def export_scout_cases_report(
    session_state: SessionState,
    mode: ScoutMode = "pending_downloads",
    timestamp: str | None = None,
) -> Path:
    resumable_report = find_resumable_scout_report(session_state, mode)
    if resumable_report is not None:
        ensure_report_metadata_columns(resumable_report)
        return resumable_report

    cases = build_scout_cases_from_session(session_state, mode=mode)
    if not cases:
        raise ValueError(f"No SCOUT cases are available for mode '{mode}'.")

    effective_timestamp = (
        timestamp
        or getattr(session_state, "timestamp", "")
        or datetime.now().strftime("%Y-%m-%d_%H-%M-%S")
    )

    wb = Workbook()
    ws = wb.active
    ws.title = "scout_cases"
    ws.append(SCOUT_REPORT_HEADERS)

    for case in cases:
        ws.append([case.get(header, "") for header in SCOUT_REPORT_HEADERS])

    report_candidates: list[Path] = []
    base_name = f"scout_{mode}_{effective_timestamp}"
    fallback_stamp = datetime.now().strftime("%Y-%m-%d_%H-%M-%S")
    repo_root = Path(__file__).resolve().parents[3]
    scout_report_dirs = [
        REPORTS_OUTPUT_DIR / "SCOUT",
        repo_root / "tools" / "pdf_fetcher" / "runtime" / "reports" / "SCOUT",
    ]

    for scout_reports_dir in scout_report_dirs:
        try:
            scout_reports_dir.mkdir(parents=True, exist_ok=True)
        except OSError:
            continue

        report_candidates.append(scout_reports_dir / f"{base_name}.xlsx")
        report_candidates.append(scout_reports_dir / f"{base_name}_{fallback_stamp}.xlsx")
        for idx in range(1, 6):
            report_candidates.append(scout_reports_dir / f"{base_name}_{fallback_stamp}_{idx}.xlsx")

    last_error: Exception | None = None
    for report_path in report_candidates:
        try:
            wb.save(report_path)
            return report_path
        except PermissionError as exc:
            last_error = exc
            continue

    if last_error is not None:
        raise last_error
    raise RuntimeError("Could not save SCOUT report.")


def launch_scout_with_report(report_path: Path) -> str:
    repo_root = Path(__file__).resolve().parents[3]
    python_candidates = [
        repo_root / ".venv" / "Scripts" / "python.exe",
        repo_root / "venv" / "Scripts" / "python.exe",
    ]
    python_exe = next((candidate for candidate in python_candidates if candidate.exists()), None)
    if python_exe is None:
        raise FileNotFoundError("Python virtual environment not found in .venv or venv.")

    env = os.environ.copy()
    env["SCOUT_AUTO_REPORT_PATH"] = str(report_path)
    env["SCOUT_AUTO_LOAD"] = "1"
    env["BROWSER"] = ""
    selected_port = _pick_available_port()
    env["SCOUT_SERVER_PORT"] = str(selected_port)

    creationflags = 0
    if sys.platform.startswith("win"):
        creationflags = subprocess.CREATE_NEW_CONSOLE

    process = subprocess.Popen(
        [
            str(python_exe),
            "-m",
            "streamlit",
            "run",
            "tools\\pdf_fetcher\\scout\\app.py",
            "--server.port",
            str(selected_port),
            "--server.headless",
            "true",
            "--browser.gatherUsageStats",
            "false",
        ],
        cwd=str(repo_root),
        env=env,
        creationflags=creationflags,
    )

    scout_url = f"http://localhost:{selected_port}"
    if not _wait_for_server_start(selected_port, process):
        exit_code = process.poll()
        raise RuntimeError(
            f"Streamlit did not start correctly on port {selected_port}."
            + (f" Process exited with code {exit_code}." if exit_code is not None else "")
        )

    opened = webbrowser.open_new_tab(scout_url)
    if not opened and sys.platform.startswith("win"):
        subprocess.Popen(
            ["cmd", "/c", "start", "", scout_url],
            cwd=str(repo_root),
            creationflags=creationflags,
        )
    return scout_url
