from pathlib import Path
from datetime import datetime
import re
import sys

from openpyxl import Workbook

PROJECT_ROOT = Path(__file__).resolve().parents[6]
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

from tools.pdf_fetcher.core.config import SHEET_NAME, build_current_workbook_path
from tools.pdf_fetcher.core.excel_io import build_col_map, build_record, get_optional_cell, load_workbook_and_sheet
from tools.pdf_fetcher.core.utils import build_pdf_filename, normalize_doi

import wiley


PUBLISHER = "Wiley"
OUTPUT_DIR = Path.home() / "Desktop" / "resolver_tests" / "wiley"
OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
REPORTS_DIR = OUTPUT_DIR / "reports"
REPORTS_DIR.mkdir(parents=True, exist_ok=True)


def build_output_file(record_id: str, title: str, doi: str) -> Path:
    try:
        file_name = build_pdf_filename(record_id, title)
    except Exception:
        file_name = wiley.minimal_safe_filename_from_doi(doi)
    return OUTPUT_DIR / file_name


def is_wiley_candidate(doi_raw: str, doi_link_raw: str, source_url: str = "") -> bool:
    values = [
        str(doi_raw or "").strip(),
        str(doi_link_raw or "").strip(),
        str(source_url or "").strip(),
    ]

    normalized_doi = normalize_doi(doi_raw)
    normalized_doi_link = normalize_doi(doi_link_raw)
    if normalized_doi:
        values.append(normalized_doi)
    if normalized_doi_link:
        values.append(normalized_doi_link)

    lowered = [v.lower() for v in values if v]
    for value in lowered:
        if "onlinelibrary.wiley.com" in value:
            return True
        if "wiley.com" in value:
            return True
        if value.startswith("10.1111/"):
            return True
        if "/10.1111/" in value:
            return True
    return False


def collect_wiley_cases_from_current() -> list[dict]:
    workbook_path = build_current_workbook_path()
    wb, ws = load_workbook_and_sheet(workbook_path, SHEET_NAME)

    try:
        col_map = build_col_map(ws)
        cases: list[dict] = []

        for row_idx in range(2, ws.max_row + 1):
            record = build_record(ws, row_idx, col_map)
            if record is None:
                continue

            source_url = str(get_optional_cell(ws, row_idx, col_map, "pdf_source_url", "") or "").strip()
            if not is_wiley_candidate(record.doi_raw, record.doi_link_raw, source_url):
                continue

            doi = normalize_doi(record.doi_raw) or normalize_doi(record.doi_link_raw)
            if not doi:
                continue

            cases.append(
                {
                    "row_idx": row_idx,
                    "record_id": record.record_id,
                    "title": record.title,
                    "doi": doi,
                    "article_url": f"https://doi.org/{doi}",
                    "source_url": source_url,
                    "output_file": build_output_file(record.record_id, record.title, doi),
                }
            )

        return cases
    finally:
        wb.close()


def process_case(case: dict) -> str:
    wiley.ARTICLE_URL = case["article_url"]
    wiley.SOURCE_URL = case["source_url"] or case["article_url"]
    wiley.DOI = case["doi"]
    wiley.RECORD_ID = case["record_id"]
    wiley.TITLE = case["title"]
    wiley.OUTPUT_DIR = OUTPUT_DIR

    output_file = case["output_file"]
    if output_file.exists() and output_file.stat().st_size > 0:
        print(f"[{case['record_id']}] duplicate_pdf -> {output_file.name}")
        return "duplicate_pdf"

    wiley.main()
    return "downloaded"


def save_failed_cases_report(failed_cases: list[dict]) -> Path | None:
    if not failed_cases:
        return None

    wb = Workbook()
    ws = wb.active
    ws.title = "wiley_failed_cases"
    ws.append(
        [
            "record_id",
            "title",
            "doi",
            "doi_url",
            "source_url",
            "output_file",
            "error_type",
            "error_message",
        ]
    )

    for case in failed_cases:
        ws.append(
            [
                case["record_id"],
                case["title"],
                case["doi"],
                case["article_url"],
                case["source_url"],
                str(case["output_file"]),
                case["error_type"],
                case["error_message"],
            ]
        )

    timestamp = datetime.now().strftime("%Y-%m-%d_%H-%M-%S")
    report_path = REPORTS_DIR / f"wiley_failed_cases_{timestamp}.xlsx"
    wb.save(report_path)
    return report_path


def main():
    current_workbook = build_current_workbook_path()
    cases = collect_wiley_cases_from_current()

    print(f"Current workbook: {current_workbook}")
    print(f"Wiley candidates found: {len(cases)}")

    if not cases:
        print("No Wiley candidates found in Current workbook.")
        return

    downloaded = 0
    duplicates = 0
    failed = 0
    failed_cases: list[dict] = []

    for idx, case in enumerate(cases, start=1):
        print(f"\n[{idx}/{len(cases)}] Processing Wiley candidate...")
        try:
            status = process_case(case)
            if status == "downloaded":
                downloaded += 1
            elif status == "duplicate_pdf":
                duplicates += 1
        except KeyboardInterrupt:
            raise
        except Exception as exc:
            failed += 1
            failed_cases.append(
                {
                    **case,
                    "error_type": type(exc).__name__,
                    "error_message": str(exc),
                }
            )
            print(f"[{case['record_id']}] failed -> {type(exc).__name__}: {exc}")

    report_path = save_failed_cases_report(failed_cases)

    print("\n" + "=" * 78)
    print("Wiley batch finished")
    print(f"Candidates : {len(cases)}")
    print(f"Downloaded : {downloaded}")
    print(f"Duplicates : {duplicates}")
    print(f"Failed     : {failed}")
    if report_path:
        print(f"Report     : {report_path}")


if __name__ == "__main__":
    main()
