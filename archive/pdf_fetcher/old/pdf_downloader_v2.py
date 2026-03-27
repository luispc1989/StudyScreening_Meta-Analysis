from __future__ import annotations

import os
import re
import sys
import time
from concurrent.futures import FIRST_COMPLETED, ThreadPoolExecutor, wait
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Dict, List, Optional, Tuple
from urllib.parse import urljoin

import requests
from openpyxl import load_workbook
from playwright.sync_api import sync_playwright, TimeoutError as PlaywrightTimeoutError
from requests.adapters import HTTPAdapter
from urllib3.util.retry import Retry


# =========================
# CONFIGURATION
# =========================

EXTERNAL_WOS_DIR = Path(
    r"C:\Users\Luís Pinto Coelho\Desktop\Dissertação\metanalysis development\2 - literature search & study screening\data\wos"
)

EXCEL_INPUT_DIR = EXTERNAL_WOS_DIR / "excel_files"
PDF_BASE_DIR = EXTERNAL_WOS_DIR / "pdf_files"

PREFERRED_WORKBOOK_PATH = EXCEL_INPUT_DIR / "wos_workbook_03.03.2026_v2.xlsx"

if PREFERRED_WORKBOOK_PATH.exists():
    WORKBOOK_PATH = PREFERRED_WORKBOOK_PATH
else:
    candidates = sorted(
        [p for p in EXCEL_INPUT_DIR.glob("*.xlsx") if not p.name.startswith("~$")],
        key=lambda p: p.stat().st_mtime,
        reverse=True,
    )

    if not candidates:
        raise FileNotFoundError(f"No .xlsx workbook found in: {EXCEL_INPUT_DIR}")

    WORKBOOK_PATH = candidates[0]

OUTPUT_WORKBOOK_PATH = EXCEL_INPUT_DIR / f"{WORKBOOK_PATH.stem}_pdf_downloaded.xlsx"
SHEET_NAME = "wos_full_text"

# Phase 1 - base download
MAX_WORKERS = 4
CONNECT_TIMEOUT = 8
READ_TIMEOUT = 15
SLEEP_BETWEEN_REQUESTS = 0.05
MAX_ROWS_TO_PROCESS: Optional[int] = None

# Benchmark mode
ENABLE_BENCHMARK_MODE = False
BENCHMARK_RECORD_IDS: set[str] = set()

# Save cadence
SAVE_EVERY_N_ROWS = 100
SPECIALIZED_SAVE_EVERY_N_ROWS = 10

# Dashboard refresh
REFRESH_DASHBOARD_EVERY_N_COMPLETIONS = 10

# Force simple terminal mode if desired
FORCE_SIMPLE_TERMINAL = False

# Phase 2 - specialized retry
ENABLE_SPECIALIZED_RESOLVERS = True

# Frontiers / Playwright
ENABLE_FRONTIERS_RESOLVER = True
FRONTIERS_HEADLESS = False
FRONTIERS_NAVIGATION_TIMEOUT_MS = 30000
FRONTIERS_DOWNLOAD_TIMEOUT_MS = 40000
FRONTIERS_POST_LOAD_WAIT_MS = 1000
FRONTIERS_ALLOWED_DOMAIN_PATTERN = r"(^|\.)frontiersin\.org$"

HEADERS = {
    "User-Agent": (
        "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
        "AppleWebKit/537.36 (KHTML, like Gecko) "
        "Chrome/122.0 Safari/537.36"
    ),
    "Accept": (
        "text/html,application/xhtml+xml,application/xml;q=0.9,"
        "application/pdf,*/*;q=0.8"
    ),
}

REQUIRED_INPUT_COLUMNS = [
    "record_id",
    "Title",
    "DOI",
    "DOI Link",
]

REQUIRED_OUTPUT_COLUMNS = [
    "pdf_downloaded",
    "pdf_download_status",
    "pdf_file_name",
    "pdf_source_url",
    "pdf_local_path",
    "pdf_http_status",
    "pdf_checked_at",
]

REQUIRED_COLUMNS = REQUIRED_INPUT_COLUMNS + REQUIRED_OUTPUT_COLUMNS

BASE_AVAILABLE_STATUSES = {"downloaded", "duplicate_pdf"}
BASE_DOWNLOADED_NOW_STATUSES = {"downloaded"}
SPECIALIZED_DOWNLOADED_NOW_STATUSES = {
    "downloaded_frontiers",
}


# =========================
# DATA MODELS
# =========================

@dataclass
class Record:
    row_idx: int
    record_id: str
    title: str
    doi_raw: str
    doi_link_raw: str
    file_name: str
    pdf_path: Path
    relative_path: str


@dataclass
class DownloadResult:
    pdf_downloaded: int
    pdf_download_status: str
    pdf_file_name: str
    pdf_source_url: str
    pdf_local_path: str
    pdf_http_status: Optional[int]
    pdf_checked_at: str


@dataclass
class RetryTask:
    record: Record
    resolver_name: str


# =========================
# TERMINAL LAYOUT ENGINE
# =========================

FRAME_WIDTH = 78

_TERMINAL_MODE = "simple"  # "ansi" or "simple"
_PHASE1_RENDERED = False
_PHASE2_RENDERED = False
_PHASE1_LAST_HEIGHT = 0
_PHASE2_LAST_HEIGHT = 0


def is_git_bash() -> bool:
    env = os.environ
    return (
        "MINGW" in env.get("MSYSTEM", "").upper()
        or "git\\usr\\bin" in env.get("SHELL", "").lower()
        or "bash.exe" in env.get("TERM_PROGRAM", "").lower()
    )


def enable_windows_vt_mode() -> bool:
    if os.name != "nt":
        return True

    try:
        import ctypes
        kernel32 = ctypes.windll.kernel32
        handle = kernel32.GetStdHandle(-11)  # STD_OUTPUT_HANDLE
        mode = ctypes.c_uint32()
        if kernel32.GetConsoleMode(handle, ctypes.byref(mode)):
            ENABLE_VIRTUAL_TERMINAL_PROCESSING = 0x0004
            new_mode = mode.value | ENABLE_VIRTUAL_TERMINAL_PROCESSING
            return bool(kernel32.SetConsoleMode(handle, new_mode))
    except Exception:
        return False

    return False


def init_terminal_mode() -> None:
    global _TERMINAL_MODE

    if FORCE_SIMPLE_TERMINAL:
        _TERMINAL_MODE = "simple"
        return

    if is_git_bash():
        _TERMINAL_MODE = "simple"
        return

    ansi_ok = sys.stdout.isatty() and enable_windows_vt_mode()
    _TERMINAL_MODE = "ansi" if ansi_ok else "simple"


def fit_line(text: str, width: int = FRAME_WIDTH) -> str:
    text = str(text or "")
    if len(text) > width:
        return text[:width]
    return text.ljust(width)


def normalize_block(lines: List[str]) -> List[str]:
    return [fit_line(line) for line in lines]


def clear_screen_simple() -> None:
    os.system("cls" if os.name == "nt" else "clear")


def _clear_current_line() -> None:
    sys.stdout.write("\033[2K\r")


def _move_up(lines: int) -> None:
    if lines > 0:
        sys.stdout.write(f"\033[{lines}A")


def _render_block_in_place(lines: List[str], already_rendered: bool, last_height: int) -> int:
    lines = normalize_block(lines)
    new_height = len(lines)

    if _TERMINAL_MODE != "ansi":
        sys.stdout.write("\n".join(lines) + "\n")
        sys.stdout.flush()
        return new_height

    if already_rendered:
        _move_up(last_height)

    total_lines_to_draw = max(last_height, new_height)

    for i in range(total_lines_to_draw):
        _clear_current_line()

        if i < new_height:
            sys.stdout.write(lines[i])

        if i < total_lines_to_draw - 1:
            sys.stdout.write("\n")

    sys.stdout.flush()
    return new_height


def render_simple_combined(phase1_lines: List[str], phase2_lines: Optional[List[str]] = None) -> None:
    clear_screen_simple()
    sys.stdout.write("\n".join(normalize_block(phase1_lines)))
    sys.stdout.write("\n")
    if phase2_lines:
        sys.stdout.write("\n")
        sys.stdout.write("\n".join(normalize_block(phase2_lines)))
        sys.stdout.write("\n")
    sys.stdout.flush()


def render_phase1_block(lines: List[str]) -> None:
    global _PHASE1_RENDERED, _PHASE1_LAST_HEIGHT

    if _TERMINAL_MODE == "simple":
        render_simple_combined(lines, None)
        _PHASE1_RENDERED = True
        _PHASE1_LAST_HEIGHT = len(lines)
        return

    _PHASE1_LAST_HEIGHT = _render_block_in_place(
        lines=lines,
        already_rendered=_PHASE1_RENDERED,
        last_height=_PHASE1_LAST_HEIGHT,
    )
    _PHASE1_RENDERED = True


def freeze_phase1_and_start_phase2(phase1_lines: List[str], phase2_lines: List[str]) -> None:
    global _PHASE2_RENDERED, _PHASE2_LAST_HEIGHT

    if _TERMINAL_MODE == "simple":
        render_simple_combined(phase1_lines, phase2_lines)
        _PHASE2_RENDERED = True
        _PHASE2_LAST_HEIGHT = len(phase2_lines)
        return

    if _PHASE2_RENDERED:
        _PHASE2_LAST_HEIGHT = _render_block_in_place(
            lines=phase2_lines,
            already_rendered=True,
            last_height=_PHASE2_LAST_HEIGHT,
        )
        return

    sys.stdout.write("\n")
    sys.stdout.flush()

    _PHASE2_LAST_HEIGHT = _render_block_in_place(
        lines=phase2_lines,
        already_rendered=False,
        last_height=0,
    )
    _PHASE2_RENDERED = True


def render_phase2_block(phase1_lines: List[str], phase2_lines: List[str]) -> None:
    global _PHASE2_RENDERED, _PHASE2_LAST_HEIGHT

    if _TERMINAL_MODE == "simple":
        render_simple_combined(phase1_lines, phase2_lines)
        _PHASE2_RENDERED = True
        _PHASE2_LAST_HEIGHT = len(phase2_lines)
        return

    if not _PHASE2_RENDERED:
        freeze_phase1_and_start_phase2(phase1_lines, phase2_lines)
        return

    _PHASE2_LAST_HEIGHT = _render_block_in_place(
        lines=phase2_lines,
        already_rendered=True,
        last_height=_PHASE2_LAST_HEIGHT,
    )


# =========================
# GENERIC HELPERS
# =========================

def now_str() -> str:
    return datetime.now().strftime("%Y-%m-%d %H:%M:%S")


def is_blank(value) -> bool:
    return value is None or str(value).strip() == ""


def sanitize_filename(text: str, max_len: int = 140) -> str:
    text = str(text or "").strip()

    text = re.sub(r'[\\/:*?"<>|]', " ", text)
    text = re.sub(r"[\r\n\t]+", " ", text)
    text = re.sub(r"\s+", " ", text).strip()
    text = text.replace(" ", "_")
    text = re.sub(r"_+", "_", text).strip("_")

    if not text:
        text = "untitled"

    return text[:max_len].rstrip("_")


def build_pdf_filename(record_id: str, title: str) -> str:
    safe_title = sanitize_filename(title)
    return f"{record_id}__{safe_title}.pdf"


def normalize_doi(doi: str) -> Optional[str]:
    if is_blank(doi):
        return None

    doi = str(doi).strip()
    doi = re.sub(r"^doi:\s*", "", doi, flags=re.IGNORECASE)
    doi = re.sub(r"^https?://(dx\.)?doi\.org/", "", doi, flags=re.IGNORECASE)
    doi = doi.strip()

    if re.match(r"^10\.\d{4,9}/\S+$", doi):
        return doi.rstrip(" .;,:)")

    match = re.search(r"(10\.\d{4,9}/\S+)", doi, flags=re.IGNORECASE)
    if match:
        return match.group(1).rstrip(" .;,:)")

    return None


def make_doi_url(doi: str) -> str:
    return f"https://doi.org/{doi}"


def extract_host(url: str) -> str:
    match = re.search(r"^https?://([^/]+)", str(url or "").strip(), re.IGNORECASE)
    return match.group(1).lower() if match else ""


def url_matches_domain_pattern(url: str, pattern: str) -> bool:
    host = extract_host(url)
    if not host:
        return False
    return re.search(pattern, host) is not None


def response_looks_like_pdf(resp: requests.Response) -> bool:
    content_type = (resp.headers.get("Content-Type") or "").lower()
    if "application/pdf" in content_type:
        return True

    try:
        return resp.content[:5] == b"%PDF-"
    except Exception:
        return False


def save_pdf_bytes(pdf_path: Path, content: bytes) -> bool:
    if not content:
        return False

    pdf_path.parent.mkdir(parents=True, exist_ok=True)
    pdf_path.write_bytes(content)
    return pdf_path.exists() and pdf_path.stat().st_size > 0


def extract_pdf_url_from_html(base_url: str, html: str) -> Optional[str]:
    if not html:
        return None

    patterns = [
        r'<meta[^>]+name=["\']citation_pdf_url["\'][^>]+content=["\']([^"\']+)["\']',
        r'<meta[^>]+content=["\']([^"\']+)["\'][^>]+name=["\']citation_pdf_url["\']',
        r'href=["\']([^"\']+\.pdf(?:\?[^"\']*)?)["\']',
        r'src=["\']([^"\']+\.pdf(?:\?[^"\']*)?)["\']',
    ]

    for pattern in patterns:
        match = re.search(pattern, html, flags=re.IGNORECASE)
        if match:
            return urljoin(base_url, match.group(1))

    return None


def classify_html_failure(status_code: Optional[int], html: str) -> str:
    text = (html or "").lower()

    if status_code == 403:
        return "paywalled"

    paywall_markers = [
        "purchase pdf",
        "buy article",
        "access through your institution",
        "institutional access",
        "subscribe to journal",
        "get access",
        "login via institution",
        "rent this article",
    ]
    if any(marker in text for marker in paywall_markers):
        return "paywalled"

    metadata_markers = [
        "abstract",
        "references",
        "metrics",
        "article info",
    ]
    if any(marker in text for marker in metadata_markers):
        return "metadata_only"

    if status_code == 404:
        return "not_found"

    if status_code is not None and status_code >= 500:
        return "broken_link"

    return "manual_check"


def format_seconds(seconds: float) -> str:
    seconds = max(0, int(seconds))
    hours, rem = divmod(seconds, 3600)
    minutes, secs = divmod(rem, 60)
    return f"{hours:02d}:{minutes:02d}:{secs:02d}"


def deduplicate_urls(urls: List[str]) -> List[str]:
    seen = set()
    unique_urls: List[str] = []

    for url in urls:
        cleaned = str(url or "").strip()
        if not cleaned:
            continue
        if cleaned not in seen:
            seen.add(cleaned)
            unique_urls.append(cleaned)

    return unique_urls


def build_session() -> requests.Session:
    session = requests.Session()

    retry = Retry(
        total=2,
        read=2,
        connect=2,
        backoff_factor=0.3,
        status_forcelist=[429, 500, 502, 503, 504],
        allowed_methods=["GET", "HEAD", "OPTIONS"],
        raise_on_status=False,
    )

    adapter = HTTPAdapter(
        max_retries=retry,
        pool_connections=20,
        pool_maxsize=20,
    )

    session.mount("http://", adapter)
    session.mount("https://", adapter)

    return session


def make_result(
    record: Record,
    downloaded: int,
    status: str,
    source_url: str = "",
    local_path: str = "",
    http_status: Optional[int] = None,
) -> DownloadResult:
    return DownloadResult(
        pdf_downloaded=downloaded,
        pdf_download_status=status,
        pdf_file_name=record.file_name,
        pdf_source_url=source_url,
        pdf_local_path=local_path,
        pdf_http_status=http_status,
        pdf_checked_at=now_str(),
    )


def safe_save_workbook(wb, output_path: Path, max_attempts: int = 2) -> bool:
    for attempt in range(1, max_attempts + 1):
        try:
            temp_path = output_path.with_name(output_path.stem + ".__tmp__.xlsx")
            wb.save(temp_path)
            if output_path.exists():
                try:
                    output_path.unlink()
                except PermissionError:
                    pass
            temp_path.replace(output_path)
            return True
        except KeyboardInterrupt:
            raise
        except Exception:
            if attempt == max_attempts:
                return False
            time.sleep(0.5)
    return False


# =========================
# EXCEL HELPERS
# =========================

def build_col_map(ws) -> Dict[str, int]:
    col_map: Dict[str, int] = {}
    for col_idx in range(1, ws.max_column + 1):
        header = ws.cell(row=1, column=col_idx).value
        if header is not None and str(header).strip():
            col_map[str(header).strip()] = col_idx
    return col_map


def validate_required_columns(col_map: Dict[str, int], sheet_name: str) -> None:
    missing = [c for c in REQUIRED_COLUMNS if c not in col_map]
    if missing:
        raise ValueError(f"Missing required columns in '{sheet_name}': {missing}")


def get_cell(ws, row_idx: int, col_map: Dict[str, int], header: str):
    return ws.cell(row=row_idx, column=col_map[header]).value


def set_cell(ws, row_idx: int, col_map: Dict[str, int], header: str, value) -> None:
    ws.cell(row=row_idx, column=col_map[header], value=value)


def build_record(ws, row_idx: int, col_map: Dict[str, int]) -> Optional[Record]:
    record_id = str(get_cell(ws, row_idx, col_map, "record_id") or "").strip()
    if not record_id:
        return None

    if ENABLE_BENCHMARK_MODE and BENCHMARK_RECORD_IDS and record_id not in BENCHMARK_RECORD_IDS:
        return None

    title = str(get_cell(ws, row_idx, col_map, "Title") or "").strip()
    doi_raw = str(get_cell(ws, row_idx, col_map, "DOI") or "").strip()
    doi_link_raw = str(get_cell(ws, row_idx, col_map, "DOI Link") or "").strip()

    file_name = build_pdf_filename(record_id, title)
    pdf_path = PDF_BASE_DIR / file_name
    relative_path = f"pdf_files/{file_name}"

    return Record(
        row_idx=row_idx,
        record_id=record_id,
        title=title,
        doi_raw=doi_raw,
        doi_link_raw=doi_link_raw,
        file_name=file_name,
        pdf_path=pdf_path,
        relative_path=relative_path,
    )


def write_result(ws, row_idx: int, col_map: Dict[str, int], result: DownloadResult) -> None:
    set_cell(ws, row_idx, col_map, "pdf_downloaded", result.pdf_downloaded)
    set_cell(ws, row_idx, col_map, "pdf_download_status", result.pdf_download_status)
    set_cell(ws, row_idx, col_map, "pdf_file_name", result.pdf_file_name)
    set_cell(ws, row_idx, col_map, "pdf_source_url", result.pdf_source_url)
    set_cell(ws, row_idx, col_map, "pdf_local_path", result.pdf_local_path)
    set_cell(
        ws,
        row_idx,
        col_map,
        "pdf_http_status",
        result.pdf_http_status if result.pdf_http_status is not None else "",
    )
    set_cell(ws, row_idx, col_map, "pdf_checked_at", result.pdf_checked_at)


def collect_records(ws, col_map: Dict[str, int], max_rows_to_process: Optional[int]) -> List[Record]:
    records: List[Record] = []
    logical_seen = 0

    for row_idx in range(2, ws.max_row + 1):
        record = build_record(ws, row_idx, col_map)
        if record is None:
            continue

        records.append(record)
        logical_seen += 1

        if max_rows_to_process is not None and logical_seen >= max_rows_to_process:
            break

    return records


# =========================
# PHASE 1 - BASE DOWNLOAD
# =========================

def try_download_from_url(
    url: str,
    session: requests.Session,
) -> Tuple[bool, str, Optional[int], Optional[bytes], str]:
    try:
        resp = session.get(
            url,
            headers=HEADERS,
            timeout=(CONNECT_TIMEOUT, READ_TIMEOUT),
            allow_redirects=True,
        )
        http_status = resp.status_code
        final_url = resp.url

        if http_status >= 400:
            if http_status == 403:
                return False, "paywalled", http_status, None, final_url
            if http_status == 404:
                return False, "not_found", http_status, None, final_url
            return False, "broken_link", http_status, None, final_url

        if response_looks_like_pdf(resp):
            return True, "downloaded", http_status, resp.content, final_url

        html = resp.text or ""
        pdf_url = extract_pdf_url_from_html(final_url, html)

        if pdf_url:
            pdf_resp = session.get(
                pdf_url,
                headers=HEADERS,
                timeout=(CONNECT_TIMEOUT, READ_TIMEOUT),
                allow_redirects=True,
            )
            pdf_http_status = pdf_resp.status_code
            pdf_final_url = pdf_resp.url

            if pdf_http_status >= 400:
                if pdf_http_status == 403:
                    return False, "paywalled", pdf_http_status, None, pdf_final_url
                if pdf_http_status == 404:
                    return False, "not_found", pdf_http_status, None, pdf_final_url
                return False, "broken_link", pdf_http_status, None, pdf_final_url

            if response_looks_like_pdf(pdf_resp):
                return True, "downloaded", pdf_http_status, pdf_resp.content, pdf_final_url

            return False, "manual_check", pdf_http_status, None, pdf_final_url

        return False, classify_html_failure(http_status, html), http_status, None, final_url

    except requests.Timeout:
        return False, "broken_link", None, None, url
    except requests.RequestException:
        return False, "broken_link", None, None, url


def build_candidate_urls(record: Record) -> List[str]:
    urls: List[str] = []

    doi_link = str(record.doi_link_raw or "").strip()
    doi = normalize_doi(record.doi_raw)

    if doi_link:
        urls.append(doi_link)

    if doi:
        urls.append(make_doi_url(doi))

    return deduplicate_urls(urls)


def try_base_download(record: Record, session: requests.Session) -> DownloadResult:
    if record.pdf_path.exists() and record.pdf_path.stat().st_size > 0:
        return make_result(
            record=record,
            downloaded=1,
            status="duplicate_pdf",
            local_path=record.relative_path,
        )

    candidate_urls = build_candidate_urls(record)

    if not candidate_urls:
        return make_result(
            record=record,
            downloaded=0,
            status="invalid_doi",
        )

    final_status = "not_found"
    final_http_status: Optional[int] = None
    final_source_url = ""

    for url in candidate_urls:
        success, status_label, http_status, pdf_bytes, final_url = try_download_from_url(url, session)

        final_status = status_label
        final_http_status = http_status
        final_source_url = final_url or url

        if success and pdf_bytes:
            saved = save_pdf_bytes(record.pdf_path, pdf_bytes)
            if saved:
                return make_result(
                    record=record,
                    downloaded=1,
                    status="downloaded",
                    source_url=final_source_url,
                    local_path=record.relative_path,
                    http_status=http_status,
                )

            final_status = "manual_check"
            final_http_status = http_status
            final_source_url = final_url or url

    return make_result(
        record=record,
        downloaded=0,
        status=final_status,
        source_url=final_source_url,
        http_status=final_http_status,
    )


def process_record_phase1(record: Record) -> Tuple[Record, DownloadResult]:
    session = build_session()
    try:
        result = try_base_download(record, session)
        return record, result
    finally:
        session.close()


# =========================
# SPECIALIZED CLASSIFIER DETECTION
# =========================

def is_frontiers_candidate_from_values(
    doi_raw: str,
    doi_link_raw: str,
    source_url: str = "",
) -> bool:
    values = [
        str(doi_raw or "").strip(),
        str(doi_link_raw or "").strip(),
        str(source_url or "").strip(),
    ]

    normalized_doi = normalize_doi(doi_raw)
    normalized_doi_link = normalize_doi(doi_link_raw)

    if normalized_doi:
        values.append(normalized_doi)

    if normalized_doi_link:
        values.append(normalized_doi_link)

    lowered = [v.lower() for v in values if v]

    for value in lowered:
        if "frontiersin.org" in value:
            return True
        if "frontiers" in value:
            return True
        if value.startswith("10.3389/"):
            return True
        if "/10.3389/" in value:
            return True

    return False


def detect_specialized_resolver_name(record: Record, phase1_result: DownloadResult) -> Optional[str]:
    if ENABLE_FRONTIERS_RESOLVER and is_frontiers_candidate_from_values(
        doi_raw=record.doi_raw,
        doi_link_raw=record.doi_link_raw,
        source_url=phase1_result.pdf_source_url,
    ):
        return "frontiers"

    return None


def should_retry_in_phase2(result: DownloadResult) -> bool:
    return result.pdf_download_status != "downloaded"


# =========================
# FRONTIERS SPECIALIZED RESOLVER
# =========================

def looks_like_pdf_url(href: str) -> bool:
    href_l = str(href or "").lower()
    return (
        "/pdf" in href_l
        or href_l.endswith(".pdf")
        or "download-a-pdf" in href_l
    )


def is_pdf_bytes_bytes(data: bytes) -> bool:
    return bool(data) and data.startswith(b"%PDF")


def extract_doi_from_page(page, fallback_url: str | None = None) -> str:
    meta_selectors = [
        "meta[name='citation_doi']",
        "meta[name='dc.Identifier']",
        "meta[name='DC.Identifier']",
        "meta[name='dc.identifier']",
        "meta[name='DC.identifier']",
        "meta[name='prism.doi']",
        "meta[property='og:doi']",
    ]

    for sel in meta_selectors:
        try:
            loc = page.locator(sel).first
            if loc.count() > 0:
                content = loc.get_attribute("content")
                if content:
                    match = re.search(r"(10\.\d{4,9}/\S+)", content.strip(), re.I)
                    if match:
                        return match.group(1).rstrip(" .;,:)")
        except Exception:
            pass

    doi_link_selectors = [
        "a[href*='doi.org/']",
        "a[href*='dx.doi.org/']",
    ]

    for sel in doi_link_selectors:
        try:
            count = page.locator(sel).count()
            for i in range(count):
                href = page.locator(sel).nth(i).get_attribute("href")
                if not href:
                    continue
                match = re.search(r"doi\.org/(10\.\d{4,9}/\S+)", href, re.I)
                if match:
                    return match.group(1).rstrip(" .;,:)")
        except Exception:
            pass

    try:
        html = page.content()
        match = re.search(r"(10\.\d{4,9}/[^\s\"<>&]+)", html, re.I)
        if match:
            return match.group(1).rstrip(" .;,:)")
    except Exception:
        pass

    if fallback_url:
        match = re.search(r"(10\.\d{4,9}/\S+)", fallback_url, re.I)
        if match:
            return match.group(1).rstrip(" .;,:)")

    raise RuntimeError("Could not detect DOI from Frontiers page.")


def find_frontiers_pdf_href(page) -> Optional[str]:
    selectors = [
        "a[data-event='download-a-pdf']",
        "a.DownloadArticleButton_action",
        "a[href*='/pdf']",
        "a[href$='.pdf']",
        "li.ToolbarDownload a",
        "a.ToolbarDownload_action",
        "nav.Toolbar a[data-event='download-a-pdf']",
        "a:has-text('PDF')",
        "a:has-text('Download')",
    ]

    for sel in selectors:
        try:
            count = page.locator(sel).count()
            for i in range(count):
                href = page.locator(sel).nth(i).get_attribute("href")
                if href and looks_like_pdf_url(href):
                    return href
        except Exception:
            pass

    return None


def try_frontiers_resolver(record: Record) -> DownloadResult:
    doi = normalize_doi(record.doi_raw)
    if not doi:
        doi = normalize_doi(record.doi_link_raw)

    if not doi:
        return make_result(
            record=record,
            downloaded=0,
            status="invalid_doi",
        )

    doi_url = make_doi_url(doi)

    if record.pdf_path.exists() and record.pdf_path.stat().st_size > 0:
        return make_result(
            record=record,
            downloaded=1,
            status="duplicate_pdf",
            source_url=doi_url,
            local_path=record.relative_path,
        )

    with sync_playwright() as p:
        browser = p.chromium.launch(
            headless=FRONTIERS_HEADLESS,
            args=["--disable-blink-features=AutomationControlled"],
        )

        context = browser.new_context(accept_downloads=True)
        page = context.new_page()

        try:
            page.goto(
                doi_url,
                wait_until="domcontentloaded",
                timeout=FRONTIERS_NAVIGATION_TIMEOUT_MS,
            )
            page.wait_for_load_state("networkidle")
            page.wait_for_timeout(FRONTIERS_POST_LOAD_WAIT_MS)

            final_page_url = page.url

            if not url_matches_domain_pattern(final_page_url, FRONTIERS_ALLOWED_DOMAIN_PATTERN):
                return make_result(
                    record=record,
                    downloaded=0,
                    status="not_frontiers",
                    source_url=final_page_url,
                )

            pdf_href = find_frontiers_pdf_href(page)
            _detected_doi = extract_doi_from_page(page, fallback_url=final_page_url)

            if not pdf_href:
                return make_result(
                    record=record,
                    downloaded=0,
                    status="pdf_link_not_found_frontiers",
                    source_url=final_page_url,
                )

            pdf_url = urljoin(final_page_url, pdf_href)

            try:
                response = context.request.get(
                    pdf_url,
                    headers={"Referer": final_page_url},
                    timeout=FRONTIERS_DOWNLOAD_TIMEOUT_MS,
                )

                if response.status == 200:
                    data = response.body()
                    if is_pdf_bytes_bytes(data):
                        if save_pdf_bytes(record.pdf_path, data):
                            return make_result(
                                record=record,
                                downloaded=1,
                                status="downloaded_frontiers",
                                source_url=pdf_url,
                                local_path=record.relative_path,
                                http_status=response.status,
                            )
            except Exception:
                pass

            for sel in [
                "a[data-event='download-a-pdf']",
                "a.DownloadArticleButton_action",
                "a[href*='/pdf']",
                "a[href$='.pdf']",
                "li.ToolbarDownload a",
                "a.ToolbarDownload_action",
                "nav.Toolbar a[data-event='download-a-pdf']",
                "a:has-text('PDF')",
                "a:has-text('Download')",
            ]:
                try:
                    count = page.locator(sel).count()
                    for i in range(count):
                        loc = page.locator(sel).nth(i)
                        href = loc.get_attribute("href")
                        if not href or not looks_like_pdf_url(href):
                            continue

                        with page.expect_download(timeout=FRONTIERS_DOWNLOAD_TIMEOUT_MS) as download_info:
                            loc.click()

                        download = download_info.value
                        download.save_as(str(record.pdf_path))

                        if record.pdf_path.exists() and record.pdf_path.stat().st_size > 0:
                            return make_result(
                                record=record,
                                downloaded=1,
                                status="downloaded_frontiers",
                                source_url=pdf_url,
                                local_path=record.relative_path,
                            )
                except PlaywrightTimeoutError:
                    continue
                except Exception:
                    continue

            return make_result(
                record=record,
                downloaded=0,
                status="download_failed_frontiers",
                source_url=final_page_url,
            )

        finally:
            browser.close()


def process_record_phase2(record: Record, resolver_name: str) -> DownloadResult:
    if resolver_name == "frontiers":
        return try_frontiers_resolver(record)

    return make_result(
        record=record,
        downloaded=0,
        status="unsupported_specialized_resolver",
    )


# =========================
# DASHBOARDS
# =========================

def build_phase1_lines(
    processed: int,
    total: int,
    workers: int,
    start_time: float,
    last_record_id: str,
    last_status: str,
    pdfs_available_global: int,
    downloaded_now_total: int,
) -> List[str]:
    remaining_count = max(0, total - processed)
    failures_count = max(0, processed - pdfs_available_global)
    elapsed = time.time() - start_time

    if processed > 0:
        estimated_total = (elapsed / processed) * total
        estimated_remaining = max(0.0, estimated_total - elapsed)
    else:
        estimated_total = 0.0
        estimated_remaining = 0.0

    success_rate_global = (pdfs_available_global / total * 100) if total > 0 else 0.0

    return [
        "=" * FRAME_WIDTH,
        "PDF DOWNLOAD STATUS",
        "=" * FRAME_WIDTH,
        "Fase atual              : 1/2 - base download",
        "-" * FRAME_WIDTH,
        f"Artigos diagnosticados  : {processed}",
        f"Artigos por diagnosticar: {remaining_count}",
        f"Total a diagnosticar    : {total}",
        f"Workers paralelos       : {workers}",
        "-" * FRAME_WIDTH,
        f"PDFs disponíveis        : {pdfs_available_global}",
        f"PDFs descarregados agora: {downloaded_now_total}",
        f"Falhas                  : {failures_count}",
        f"Taxa de sucesso global  : {success_rate_global:6.2f}%",
        "-" * FRAME_WIDTH,
        f"Tempo decorrido         : {format_seconds(elapsed)}",
        f"Tempo estimado total    : {format_seconds(estimated_total)}",
        f"Tempo restante estimado : {format_seconds(estimated_remaining)}",
        "-" * FRAME_WIDTH,
        f"Último registo          : {last_record_id or '-'}",
        f"Último estado           : {last_status or '-'}",
        "=" * FRAME_WIDTH,
    ]


def build_phase2_lines(
    retry_processed: int,
    retry_total: int,
    phase2_start_time: float,
    resolver_name: str,
    last_record_id: str,
    last_doi: str,
    last_status: str,
    pdfs_available_global: int,
    downloaded_now_total: int,
    recovered_this_phase: int,
    total_records: int,
) -> List[str]:
    retry_remaining = max(0, retry_total - retry_processed)
    phase2_still_without_pdf = max(0, retry_processed - recovered_this_phase)
    elapsed = time.time() - phase2_start_time

    if retry_processed > 0:
        estimated_total = (elapsed / retry_processed) * retry_total
        estimated_remaining = max(0.0, estimated_total - elapsed)
    else:
        estimated_total = 0.0
        estimated_remaining = 0.0

    success_rate_global = (pdfs_available_global / total_records * 100) if total_records > 0 else 0.0

    return [
        "=" * FRAME_WIDTH,
        "PDF DOWNLOAD STATUS",
        "=" * FRAME_WIDTH,
        "Fase atual              : 2/2 - specialized retry",
        f"Resolver atual          : {resolver_name or '-'}",
        "-" * FRAME_WIDTH,
        f"Artigos em retry        : {retry_total}",
        f"Artigos tratados        : {retry_processed}",
        f"Artigos por tratar      : {retry_remaining}",
        "-" * FRAME_WIDTH,
        f"PDFs disponíveis        : {pdfs_available_global}",
        f"PDFs descarregados agora: {downloaded_now_total}",
        f"Recuperados nesta fase  : {recovered_this_phase}",
        f"Ainda sem PDF           : {phase2_still_without_pdf}",
        f"Taxa de sucesso global  : {success_rate_global:6.2f}%",
        "-" * FRAME_WIDTH,
        f"Tempo decorrido fase 2  : {format_seconds(elapsed)}",
        f"Tempo estimado fase 2   : {format_seconds(estimated_total)}",
        f"Tempo restante fase 2   : {format_seconds(estimated_remaining)}",
        "-" * FRAME_WIDTH,
        f"Último registo          : {last_record_id or '-'}",
        f"Último DOI              : {last_doi or '-'}",
        f"Último estado           : {last_status or '-'}",
        "=" * FRAME_WIDTH,
    ]


# =========================
# PHASE 2 TASK BUILDING
# =========================

def build_retry_tasks(
    records: List[Record],
    phase1_results_by_row: Dict[int, DownloadResult],
) -> List[RetryTask]:
    tasks: List[RetryTask] = []

    if not ENABLE_SPECIALIZED_RESOLVERS:
        return tasks

    for record in records:
        phase1_result = phase1_results_by_row.get(record.row_idx)
        if phase1_result is None:
            continue

        if not should_retry_in_phase2(phase1_result):
            continue

        resolver_name = detect_specialized_resolver_name(record, phase1_result)
        if resolver_name is None:
            continue

        tasks.append(RetryTask(record=record, resolver_name=resolver_name))

    return tasks


# =========================
# MAIN PROCESS
# =========================

def main() -> None:
    init_terminal_mode()

    if not WORKBOOK_PATH.exists():
        raise FileNotFoundError(f"Workbook not found: {WORKBOOK_PATH}")

    PDF_BASE_DIR.mkdir(parents=True, exist_ok=True)

    wb = load_workbook(WORKBOOK_PATH)
    if SHEET_NAME not in wb.sheetnames:
        raise ValueError(f"Sheet not found: {SHEET_NAME}")

    ws = wb[SHEET_NAME]
    col_map = build_col_map(ws)
    validate_required_columns(col_map, SHEET_NAME)

    records = collect_records(ws, col_map, MAX_ROWS_TO_PROCESS)
    total_to_process = len(records)

    if total_to_process == 0:
        raise ValueError("No records found to process.")

    phase1_results_by_row: Dict[int, DownloadResult] = {}

    pdfs_available_global = 0
    downloaded_now_total = 0

    phase1_start_time = time.time()
    phase1_processed = 0
    phase1_last_record_id = ""
    phase1_last_status = "starting"

    phase1_lines = build_phase1_lines(
        processed=0,
        total=total_to_process,
        workers=MAX_WORKERS,
        start_time=phase1_start_time,
        last_record_id="",
        last_status="starting",
        pdfs_available_global=0,
        downloaded_now_total=0,
    )
    render_phase1_block(phase1_lines)

    executor = ThreadPoolExecutor(max_workers=MAX_WORKERS)

    try:
        futures = {}
        next_idx = 0

        while next_idx < total_to_process and len(futures) < MAX_WORKERS:
            record = records[next_idx]
            futures[executor.submit(process_record_phase1, record)] = record
            next_idx += 1
            if SLEEP_BETWEEN_REQUESTS > 0:
                time.sleep(SLEEP_BETWEEN_REQUESTS)

        while futures:
            done, _ = wait(list(futures.keys()), return_when=FIRST_COMPLETED)

            for future in done:
                _original_record = futures.pop(future)
                record, result = future.result()

                write_result(ws, record.row_idx, col_map, result)
                phase1_results_by_row[record.row_idx] = result

                phase1_processed += 1
                phase1_last_record_id = record.record_id
                phase1_last_status = result.pdf_download_status

                if result.pdf_download_status in BASE_AVAILABLE_STATUSES:
                    pdfs_available_global += 1

                if result.pdf_download_status in BASE_DOWNLOADED_NOW_STATUSES:
                    downloaded_now_total += 1

                if (
                    phase1_processed == 1
                    or phase1_processed % REFRESH_DASHBOARD_EVERY_N_COMPLETIONS == 0
                    or phase1_processed == total_to_process
                ):
                    phase1_lines = build_phase1_lines(
                        processed=phase1_processed,
                        total=total_to_process,
                        workers=MAX_WORKERS,
                        start_time=phase1_start_time,
                        last_record_id=phase1_last_record_id,
                        last_status=phase1_last_status,
                        pdfs_available_global=pdfs_available_global,
                        downloaded_now_total=downloaded_now_total,
                    )
                    render_phase1_block(phase1_lines)

                if phase1_processed % SAVE_EVERY_N_ROWS == 0:
                    safe_save_workbook(wb, OUTPUT_WORKBOOK_PATH)

                while next_idx < total_to_process and len(futures) < MAX_WORKERS:
                    next_record = records[next_idx]
                    futures[executor.submit(process_record_phase1, next_record)] = next_record
                    next_idx += 1
                    if SLEEP_BETWEEN_REQUESTS > 0:
                        time.sleep(SLEEP_BETWEEN_REQUESTS)

        safe_save_workbook(wb, OUTPUT_WORKBOOK_PATH)

        phase1_lines = build_phase1_lines(
            processed=phase1_processed,
            total=total_to_process,
            workers=MAX_WORKERS,
            start_time=phase1_start_time,
            last_record_id=phase1_last_record_id,
            last_status=phase1_last_status,
            pdfs_available_global=pdfs_available_global,
            downloaded_now_total=downloaded_now_total,
        )
        render_phase1_block(phase1_lines)

        retry_tasks = build_retry_tasks(records, phase1_results_by_row)
        retry_total = len(retry_tasks)

        if retry_total == 0:
            phase2_lines = build_phase2_lines(
                retry_processed=0,
                retry_total=0,
                phase2_start_time=time.time(),
                resolver_name="none",
                last_record_id="",
                last_doi="",
                last_status="no_retry_tasks",
                pdfs_available_global=pdfs_available_global,
                downloaded_now_total=downloaded_now_total,
                recovered_this_phase=0,
                total_records=total_to_process,
            )
            freeze_phase1_and_start_phase2(phase1_lines, phase2_lines)

            print()
            print(f"PDFs disponíveis        : {pdfs_available_global}")
            print(f"PDFs descarregados agora: {downloaded_now_total}")
            print(f"Saved workbook          : {OUTPUT_WORKBOOK_PATH}")
            print(f"PDF folder              : {PDF_BASE_DIR}")
            return

        phase2_start_time = time.time()
        retry_processed = 0
        recovered_this_phase = 0
        phase2_last_record_id = ""
        phase2_last_status = "starting"
        phase2_last_doi = ""
        phase2_current_resolver = "starting"

        phase2_lines = build_phase2_lines(
            retry_processed=0,
            retry_total=retry_total,
            phase2_start_time=phase2_start_time,
            resolver_name=phase2_current_resolver,
            last_record_id="",
            last_doi="",
            last_status="starting",
            pdfs_available_global=pdfs_available_global,
            downloaded_now_total=downloaded_now_total,
            recovered_this_phase=0,
            total_records=total_to_process,
        )
        freeze_phase1_and_start_phase2(phase1_lines, phase2_lines)

        for task in retry_tasks:
            record = task.record
            resolver_name = task.resolver_name

            phase2_current_resolver = resolver_name
            phase2_last_record_id = record.record_id
            phase2_last_doi = normalize_doi(record.doi_raw) or normalize_doi(record.doi_link_raw) or ""

            result = process_record_phase2(record, resolver_name)
            write_result(ws, record.row_idx, col_map, result)
            phase1_results_by_row[record.row_idx] = result

            retry_processed += 1
            phase2_last_status = result.pdf_download_status

            if result.pdf_download_status in SPECIALIZED_DOWNLOADED_NOW_STATUSES:
                downloaded_now_total += 1
                recovered_this_phase += 1
                pdfs_available_global += 1
            elif result.pdf_download_status == "duplicate_pdf":
                pdfs_available_global += 1

            if (
                retry_processed == 1
                or retry_processed % REFRESH_DASHBOARD_EVERY_N_COMPLETIONS == 0
                or retry_processed == retry_total
            ):
                phase2_lines = build_phase2_lines(
                    retry_processed=retry_processed,
                    retry_total=retry_total,
                    phase2_start_time=phase2_start_time,
                    resolver_name=phase2_current_resolver,
                    last_record_id=phase2_last_record_id,
                    last_doi=phase2_last_doi,
                    last_status=phase2_last_status,
                    pdfs_available_global=pdfs_available_global,
                    downloaded_now_total=downloaded_now_total,
                    recovered_this_phase=recovered_this_phase,
                    total_records=total_to_process,
                )
                render_phase2_block(phase1_lines, phase2_lines)

            if retry_processed % SPECIALIZED_SAVE_EVERY_N_ROWS == 0:
                safe_save_workbook(wb, OUTPUT_WORKBOOK_PATH)

        safe_save_workbook(wb, OUTPUT_WORKBOOK_PATH)

        phase2_lines = build_phase2_lines(
            retry_processed=retry_processed,
            retry_total=retry_total,
            phase2_start_time=phase2_start_time,
            resolver_name=phase2_current_resolver if phase2_current_resolver else "finished",
            last_record_id=phase2_last_record_id,
            last_doi=phase2_last_doi,
            last_status=phase2_last_status if phase2_last_status else "finished",
            pdfs_available_global=pdfs_available_global,
            downloaded_now_total=downloaded_now_total,
            recovered_this_phase=recovered_this_phase,
            total_records=total_to_process,
        )
        render_phase2_block(phase1_lines, phase2_lines)

        print()
        print(f"PDFs disponíveis        : {pdfs_available_global}")
        print(f"PDFs descarregados agora: {downloaded_now_total}")
        print(f"Saved workbook          : {OUTPUT_WORKBOOK_PATH}")
        print(f"PDF folder              : {PDF_BASE_DIR}")

    except KeyboardInterrupt:
        print("\nInterrupção pedida. A terminar workers...")
        executor.shutdown(wait=False, cancel_futures=True)

        try:
            ok = safe_save_workbook(wb, OUTPUT_WORKBOOK_PATH)
            if ok:
                print("Progresso guardado com sucesso.")
            else:
                print("Não foi possível guardar o workbook de forma segura.")
        except KeyboardInterrupt:
            print("Interrupção forçada durante a gravação final.")

        print(f"Saved workbook          : {OUTPUT_WORKBOOK_PATH}")

    finally:
        try:
            executor.shutdown(wait=False, cancel_futures=True)
        except Exception:
            pass


if __name__ == "__main__":
    main()