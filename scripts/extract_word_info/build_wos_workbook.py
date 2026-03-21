# ============================================================
# build_wos_workbook.py
# ============================================================
# WEB OF SCIENCE (WoS) -> CLEAN + NORMALIZE + SCREENING WORKBOOK
# XLSX-only outputs. Minimal formatting:
#   - Only headers in bold on all sheets
#   - Keep: data validation for Label, conditional formatting, hidden columns by default
#   - NO auto-fit, NO column width changes, NO borders
#
# FIXES IN THIS VERSION (v1.3.8)
# -----------------------------
# 1) screening_view Abstract/Title stays within the cell:
#    - Wrap Text ON for Title and Abstract
#    - Fixed row height => text is clipped (does not expand or block Label)
#
# 2) Conditional formatting works (full + screening_view):
#    - Green row if Label = 1
#    - Red row if Label = 0
#
# 3) full sheet updates Label/Reason/Comment automatically from screening_view:
#    - screening_view keeps hidden key in column A: source_record_id
#    - Lookup formulas use XLOOKUP with INDEX/MATCH fallback
#
# 4) IMPORTANT: In full, Label/Reason/Comment stay BLANK until edited:
#    - screening_view initial values are "" (empty string) not None
#    - formulas propagate "" so Excel does not show 0 by default
#
# Screening view requirements:
#   - Visible order: record_id, DOI, DOI Link, Title, Abstract, Label, Reason to Exclude, Comment
#   - Technical key source_record_id exists (hidden) as first column
#   - Freeze panes at C2 (keeps hidden key + record_id fixed)
#
# Full sheet requirements:
#   - Keep E-mail Address column
#   - DOI adjacent to DOI Link
# ============================================================

from __future__ import annotations

from pathlib import Path
from typing import List, Dict, Tuple, Any, Optional
from datetime import datetime
import hashlib
import json
import platform
import re
import sys
import traceback

import pandas as pd

# ============================================================
# USER INPUT (EDIT ONLY THIS SECTION)
# ============================================================

INPUT_DIR = Path(
    r"C:\Users\Luís Pinto Coelho\Desktop\Dissertação\StudyScreening_Meta-Analysis\data\web_of_science\raw_data"
)

OUTPUT_RUNS_DIR = Path(
    r"C:\Users\Luís Pinto Coelho\Desktop\Dissertação\metanalysis development\2 - literature search & study screening\data\wos_test\clean_data"
)

FILE_PATTERNS = ["*savedrecs*.xls*", "*savedrecs*.csv*"]

EXPORT_CSV = False  # XLSX only
INCLUDE_ISSUE = True
ENABLE_DEDUPLICATION = True
DEDUPLICATION_MODE = "doi_then_title_year"  # "strict_doi" or "doi_then_title_year"
KEEP_POLICY = "best_record"                 # "first" or "best_record"
CONTINUE_ON_FILE_ERROR = True

# ============================================================
# ADVANCED OPTIONS
# ============================================================

SCRIPT_VERSION = "wos_workbook_builder_v1.3.8"

SOURCE_NAME = "WoS"
FULL_SHEET_NAME = "wos_full"
SCREENING_SHEET_NAME = "wos_screening_view"
SEARCH_PROTOCOL_SHEET_NAME = "wos_search_protocol"

DOI_VALID_PATTERN = r"^10\.\d{4,9}/\S+$"
MIN_YEAR = 1900
MAX_YEAR = datetime.now().year + 1

# Fixed row height used in screening view to prevent Excel expanding wrapped rows
SCREENING_VIEW_ROW_HEIGHT = 15

TARGET_COLUMNS = [
    "Authors",
    "Publication Year",
    "Title",
    "Abstract",
    "Author Keywords",
    "Source Title",
    "Volume",
    "Issue",
    "Begin Page",
    "End Page",
    "Article Number",
    "E-mail Address",
    "DOI",
    "DOI Link",
]

TRACE_COLUMNS = [
    "source_record_id",
    "master_id",
    "source",
    "source_priority",
    "run_id",
    "source_path",
    "source_row_index",
    "date_imported",
]

SCREENING_COLUMNS = [
    "Label",
    "Reason to Exclude",
    "Comment",
]

ALIASES: Dict[str, List[str]] = {
    "Authors": ["authors", "author(s)"],
    "Publication Year": ["publication year", "year published", "year"],
    "Title": ["title", "article title"],
    "Abstract": ["abstract"],
    "Author Keywords": ["author keywords", "keywords"],
    "Source Title": ["source title", "journal", "publication name", "source"],
    "Volume": ["volume"],
    "Issue": ["issue", "number"],
    "Begin Page": ["begin page", "start page", "page begin"],
    "End Page": ["end page", "page end"],
    "Article Number": ["article number"],
    "E-mail Address": [
        "e-mail address", "e-mail addresses", "email address", "email addresses",
        "author email address", "E-mail Addresses"
    ],
    "DOI": ["doi"],
    "DOI Link": ["doi link"],
    "Pages": ["pages"],
    "Addresses": ["addresses", "affiliations", "author information", "reprint address"],
}

# Default hidden columns in wos_full (E-mail Address must remain visible)
FULL_HIDDEN_COLUMNS_DEFAULT = {
    "source_record_id",
    "master_id",
    "source_priority",
    "run_id",
    "source_path",
    "source_row_index",
    "date_imported",
    "Article Number",
}

# ============================================================
# BASIC HELPERS
# ============================================================

def friendly_print(msg: str) -> None:
    print(msg)

def now_run_id() -> str:
    return datetime.now().strftime("%Y-%m-%d_%H%M%S")

def source_priority_value(source: str) -> int:
    return {"WoS": 1, "GS": 2, "Grey": 3}.get(source, 99)

def write_json(path: Path, data: Dict[str, Any]) -> None:
    path.write_text(json.dumps(data, indent=2, ensure_ascii=False), encoding="utf-8")

def sha256_file(path: Path, chunk_size: int = 1024 * 1024) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        while True:
            chunk = f.read(chunk_size)
            if not chunk:
                break
            h.update(chunk)
    return h.hexdigest()

def validate_user_inputs() -> None:
    errs = []
    if not INPUT_DIR.exists():
        errs.append(f"INPUT_DIR does not exist: {INPUT_DIR}")
    if DEDUPLICATION_MODE not in {"strict_doi", "doi_then_title_year"}:
        errs.append("DEDUPLICATION_MODE must be 'strict_doi' or 'doi_then_title_year'")
    if KEEP_POLICY not in {"first", "best_record"}:
        errs.append("KEEP_POLICY must be 'first' or 'best_record'")
    if errs:
        raise ValueError("Configuration error(s):\n- " + "\n- ".join(errs))

# ============================================================
# FILE DISCOVERY / MANIFEST
# ============================================================

def list_input_files(input_dir: Path, patterns: List[str]) -> List[Path]:
    files: List[Path] = []
    for pattern in patterns:
        for f in input_dir.glob(pattern):
            if f.is_file():
                files.append(f)
    return sorted(set(files), key=lambda p: p.name.lower())

def build_input_manifest(files: List[Path]) -> pd.DataFrame:
    rows = []
    for f in files:
        st = f.stat()
        rows.append({
            "file_name": f.name,
            "file_path": str(f.resolve()),
            "suffix": f.suffix.lower(),
            "size_bytes": st.st_size,
            "modified_time_iso": datetime.fromtimestamp(st.st_mtime).isoformat(timespec="seconds"),
            "sha256": sha256_file(f),
        })
    return pd.DataFrame(rows)

# ============================================================
# READERS
# ============================================================

def read_csv_any(path: Path) -> pd.DataFrame:
    encodings = ["utf-8-sig", "utf-8", "latin-1"]
    last_err = None
    for enc in encodings:
        try:
            return pd.read_csv(
                path, sep=None, engine="python", encoding=enc,
                dtype=str, na_filter=False, low_memory=False
            )
        except Exception as e:
            last_err = e
    raise RuntimeError(f"Failed to read CSV '{path.name}'. Last error: {last_err}")

def read_excel_any(path: Path) -> pd.DataFrame:
    suf = path.suffix.lower()
    if suf == ".xlsx":
        return pd.read_excel(path, engine="openpyxl", dtype=str)
    elif suf == ".xls":
        return pd.read_excel(path, engine="xlrd", dtype=str)  # requires xlrd installed
    raise ValueError(f"Unsupported Excel file type: {path.suffix}")

def read_any(path: Path) -> pd.DataFrame:
    suf = path.suffix.lower()
    if suf in {".csv", ".tsv"}:
        return read_csv_any(path)
    if suf in {".xls", ".xlsx"}:
        return read_excel_any(path)
    raise ValueError(f"Unsupported file type: {path.suffix}")

# ============================================================
# CLEANING / NORMALIZATION
# ============================================================

def normalize_columns(df: pd.DataFrame) -> pd.DataFrame:
    lower_to_orig = {str(c).lower().strip(): c for c in df.columns}
    col_map: Dict[str, str] = {}

    for target, aliases in ALIASES.items():
        for alt in aliases:
            k = alt.lower().strip()
            if k in lower_to_orig:
                col_map[lower_to_orig[k]] = target
                break

    for c in df.columns:
        if c not in col_map and c in TARGET_COLUMNS:
            col_map[c] = c

    return df.rename(columns=col_map)

def tidy_strings(df: pd.DataFrame, cols: List[str]) -> pd.DataFrame:
    for c in cols:
        if c in df.columns:
            s = df[c].astype("string")
            s = s.str.replace("\u00A0", " ", regex=False).str.strip()
            s = s.replace({"": pd.NA, "nan": pd.NA, "None": pd.NA, "<NA>": pd.NA})
            df[c] = s
    return df

def normalize_whitespace_all(df: pd.DataFrame) -> pd.DataFrame:
    for c in df.columns:
        if pd.api.types.is_string_dtype(df[c]) or df[c].dtype == "object":
            s = df[c].astype("string").str.replace("\u00A0", " ", regex=False).str.strip()
            s = s.replace({"": pd.NA, "nan": pd.NA, "None": pd.NA, "<NA>": pd.NA})
            df[c] = s
    return df

def ensure_email_column(df: pd.DataFrame) -> pd.DataFrame:
    if "E-mail Address" not in df.columns:
        df["E-mail Address"] = pd.NA
    else:
        df["E-mail Address"] = (
            df["E-mail Address"].astype("string").str.strip()
            .replace({"": pd.NA, "nan": pd.NA, "None": pd.NA, "<NA>": pd.NA})
        )
    return df

def normalize_doi_fields(df: pd.DataFrame) -> pd.DataFrame:
    if "DOI" in df.columns:
        s = df["DOI"].astype("string").str.strip().str.lower()
        s = s.str.replace(r"^https?://(dx\.)?doi\.org/", "", regex=True)
        s = s.replace({"": pd.NA, "nan": pd.NA, "none": pd.NA, "<NA>": pd.NA})
        df["DOI"] = s

    if "DOI Link" not in df.columns:
        df["DOI Link"] = pd.NA
    else:
        s = df["DOI Link"].astype("string").str.strip()
        s = s.replace({"": pd.NA, "0": pd.NA, "nan": pd.NA, "none": pd.NA, "<NA>": pd.NA})
        df["DOI Link"] = s

    if "DOI" in df.columns:
        need = df["DOI Link"].isna() & df["DOI"].notna()
        df.loc[need, "DOI Link"] = "https://doi.org/" + df.loc[need, "DOI"].astype("string")

    return df

def derive_pages_from_pages_column(df: pd.DataFrame) -> pd.DataFrame:
    if "Pages" not in df.columns:
        return df

    need_begin = ("Begin Page" not in df.columns) or df["Begin Page"].isna().all()
    need_end = ("End Page" not in df.columns) or df["End Page"].isna().all()
    if not (need_begin or need_end):
        return df

    pat = re.compile(r"^\s*([A-Za-z0-9]+)\s*[-–—]\s*([A-Za-z0-9]+)\s*$")
    begins, ends = [], []
    for val in df["Pages"].astype("string").fillna(""):
        v = str(val).strip()
        if not v or v.lower() in {"nan", "none", "<na>"}:
            begins.append(pd.NA)
            ends.append(pd.NA)
            continue
        m = pat.match(v)
        if m:
            begins.append(m.group(1))
            ends.append(m.group(2))
        else:
            begins.append(v)
            ends.append(pd.NA)

    pages_df = pd.DataFrame({"Begin Page": begins, "End Page": ends})
    if "Begin Page" not in df.columns or df["Begin Page"].isna().all():
        df["Begin Page"] = pages_df["Begin Page"]
    if "End Page" not in df.columns or df["End Page"].isna().all():
        df["End Page"] = pages_df["End Page"]
    return df

def clean_page_columns(df: pd.DataFrame) -> pd.DataFrame:
    for c in ["Begin Page", "End Page"]:
        if c in df.columns:
            s = df[c].astype("string").str.replace("\u00A0", " ", regex=False).str.strip()
            s = s.replace({
                "": pd.NA, "+": pd.NA, "-": pd.NA, "–": pd.NA, "—": pd.NA,
                "N/A": pd.NA, "n/a": pd.NA, "na": pd.NA, "nan": pd.NA, "None": pd.NA, "<NA>": pd.NA
            })
            df[c] = s
    return df

def cast_numeric_like_safe(df: pd.DataFrame, cols: List[str]) -> pd.DataFrame:
    for c in cols:
        if c not in df.columns:
            continue
        s = df[c].astype("string").str.strip()
        mask_num = s.str.fullmatch(r"\d+").fillna(False)
        out = pd.Series([None] * len(df), index=df.index, dtype=object)
        if mask_num.any():
            out.loc[mask_num] = pd.to_numeric(s.loc[mask_num], errors="coerce", downcast="integer")
        if (~mask_num).any():
            txt = s.loc[~mask_num].replace({"": pd.NA, "nan": pd.NA, "None": pd.NA, "<NA>": pd.NA})
            out.loc[~mask_num] = txt.astype(object)
        df[c] = out
    return df

def valid_doi_mask(series: pd.Series) -> pd.Series:
    s = series.astype("string").str.strip().str.lower()
    s = s.str.replace(r"^https?://(dx\.)?doi\.org/", "", regex=True)
    return s.str.fullmatch(DOI_VALID_PATTERN).fillna(False)

def normalize_title_for_key(value: Any) -> Optional[str]:
    if value is None or pd.isna(value):
        return None
    t = str(value).strip().lower().replace("\u00A0", " ")
    t = re.sub(r"\s+", " ", t)
    t = re.sub(r"[^\w\s]", "", t)
    t = t.strip()
    return t or None

# ============================================================
# TRACEABILITY / SCHEMA
# ============================================================

def add_traceability_columns(df: pd.DataFrame, run_id: str, file_path: Path) -> pd.DataFrame:
    src_upper = "WOS"
    source_row_index = (df.index + 2).astype("Int64")

    df["source_record_id"] = [f"{src_upper}_{file_path.stem}_r{int(i):06d}" for i in source_row_index]
    df["master_id"] = pd.NA
    df["source"] = SOURCE_NAME
    df["source_priority"] = source_priority_value(SOURCE_NAME)
    df["run_id"] = run_id
    df["source_path"] = str(file_path.resolve())
    df["source_row_index"] = source_row_index
    df["date_imported"] = datetime.now().strftime("%Y-%m-%d")

    for c in SCREENING_COLUMNS:
        if c not in df.columns:
            df[c] = pd.NA

    return df

def select_and_order_full_schema(df: pd.DataFrame, include_issue: bool = True) -> pd.DataFrame:
    cols = TRACE_COLUMNS + TARGET_COLUMNS.copy() + SCREENING_COLUMNS
    if not include_issue and "Issue" in cols:
        cols.remove("Issue")

    for c in cols:
        if c not in df.columns:
            df[c] = pd.NA

    return df[cols]

# ============================================================
# SINGLE FILE PROCESSING
# ============================================================

def process_file(path: Path, run_id: str) -> Tuple[pd.DataFrame, Dict[str, Any]]:
    raw = read_any(path)
    n_raw = len(raw)

    df = normalize_columns(raw)
    df = tidy_strings(df, [
        "Authors", "Publication Year", "Title", "Abstract", "Author Keywords",
        "Source Title", "Volume", "Issue", "Begin Page", "End Page", "Article Number",
        "E-mail Address", "DOI", "DOI Link", "Pages"
    ])
    df = normalize_whitespace_all(df)
    df = ensure_email_column(df)
    df = normalize_doi_fields(df)
    df = derive_pages_from_pages_column(df)
    df = clean_page_columns(df)
    df = add_traceability_columns(df, run_id=run_id, file_path=path)
    df = select_and_order_full_schema(df, include_issue=INCLUDE_ISSUE)
    df = normalize_whitespace_all(df)

    df = cast_numeric_like_safe(df, [
        "Publication Year", "Volume", "Issue", "Begin Page", "End Page",
        "Article Number", "source_row_index", "source_priority"
    ])

    df = df.replace({"": None, pd.NA: None, "<NA>": None, "nan": None, "None": None})

    summary = {
        "file_name": path.name,
        "file_path": str(path.resolve()),
        "rows_raw": int(n_raw),
        "rows_cleaned": int(len(df)),
        "columns_raw": int(len(raw.columns)),
        "columns_final": int(len(df.columns)),
        "status": "ok",
    }
    return df, summary

# ============================================================
# DEDUPLICATION
# ============================================================

def build_record_keys(df: pd.DataFrame) -> pd.DataFrame:
    out = df.copy()
    out["_doi_norm"] = out["DOI"].astype("string").str.strip().str.lower().str.replace(
        r"^https?://(dx\.)?doi\.org/", "", regex=True
    )
    out["_doi_valid"] = valid_doi_mask(out["_doi_norm"].astype("string"))

    out["_title_norm"] = out["Title"].map(normalize_title_for_key)
    out["_year_norm"] = out["Publication Year"].astype("string").str.strip().replace(
        {"": pd.NA, "nan": pd.NA, "None": pd.NA, "<NA>": pd.NA}
    )
    out["_title_year_key"] = out["_title_norm"].fillna("") + "||" + out["_year_norm"].fillna("")

    score = pd.Series(0, index=out.index, dtype="int64")
    score += out["_doi_valid"].fillna(False).astype(int) * 100
    score += out["Abstract"].notna().astype(int) * 20
    score += out["Title"].notna().astype(int) * 10
    score += out["Publication Year"].notna().astype(int) * 5
    score += out["Author Keywords"].notna().astype(int) * 2
    out["_record_score"] = score
    return out

def sort_for_keep_policy(df: pd.DataFrame, keep_policy: str) -> pd.DataFrame:
    sort_cols, ascending = [], []
    if keep_policy == "best_record":
        sort_cols.append("_record_score")
        ascending.append(False)
    sort_cols += ["source_path", "source_row_index"]
    ascending += [True, True]
    return (
        df.sort_values(by=sort_cols, ascending=ascending, kind="mergesort")
        .reset_index(drop=False)
        .rename(columns={"index": "_orig_index"})
    )

def deduplicate_records(df: pd.DataFrame) -> Tuple[pd.DataFrame, pd.DataFrame, Dict[str, Any]]:
    work = build_record_keys(df)
    work = sort_for_keep_policy(work, KEEP_POLICY)

    keep_mask = pd.Series(True, index=work.index)
    dup_reason = pd.Series([None] * len(work), index=work.index, dtype=object)
    dup_group_key = pd.Series([None] * len(work), index=work.index, dtype=object)

    if DEDUPLICATION_MODE in {"strict_doi", "doi_then_title_year"}:
        valid_rows = work["_doi_valid"].fillna(False) & work["_doi_norm"].notna()
        seen = set()
        for i, row in work[valid_rows].iterrows():
            key = f"doi::{row['_doi_norm']}"
            if key in seen:
                keep_mask.loc[i] = False
                dup_reason.loc[i] = "duplicate_valid_doi"
                dup_group_key.loc[i] = key
            else:
                seen.add(key)

    if DEDUPLICATION_MODE == "doi_then_title_year":
        fallback_rows = (
            keep_mask
            & ~work["_doi_valid"].fillna(False)
            & work["_title_norm"].notna()
            & work["_year_norm"].notna()
        )
        seen = set()
        for i, row in work[fallback_rows].iterrows():
            key = f"title_year::{row['_title_year_key']}"
            if key in seen:
                keep_mask.loc[i] = False
                dup_reason.loc[i] = "duplicate_title_year_fallback"
                dup_group_key.loc[i] = key
            else:
                seen.add(key)

    kept = work[keep_mask].copy()
    removed = work[~keep_mask].copy()

    if not removed.empty:
        removed["_duplicate_reason"] = dup_reason.loc[removed.index].values
        removed["_duplicate_group_key"] = dup_group_key.loc[removed.index].values

    rows = []
    if not removed.empty:
        kept_lookup = kept.copy()
        kept_lookup["_k_doi"] = "doi::" + kept_lookup["_doi_norm"].astype("string")
        kept_lookup["_k_ty"] = "title_year::" + kept_lookup["_title_year_key"].astype("string")

        for _, r in removed.iterrows():
            gk = r.get("_duplicate_group_key")
            if not gk:
                continue
            if str(gk).startswith("doi::"):
                retained_match = kept_lookup[kept_lookup["_k_doi"] == gk]
            else:
                retained_match = kept_lookup[kept_lookup["_k_ty"] == gk]

            retained = retained_match.iloc[0] if not retained_match.empty else None

            item = {
                "duplicate_reason": r.get("_duplicate_reason"),
                "duplicate_group_key": gk,
                "removed_source_record_id": r.get("source_record_id"),
                "removed_source_row_index": r.get("source_row_index"),
                "removed_title": r.get("Title"),
                "removed_publication_year": r.get("Publication Year"),
                "removed_doi": r.get("DOI"),
                "removed_record_score": r.get("_record_score"),
            }
            if retained is not None:
                item.update({
                    "retained_source_record_id": retained.get("source_record_id"),
                    "retained_source_row_index": retained.get("source_row_index"),
                    "retained_title": retained.get("Title"),
                    "retained_publication_year": retained.get("Publication Year"),
                    "retained_doi": retained.get("DOI"),
                    "retained_record_score": retained.get("_record_score"),
                })
            rows.append(item)

    duplicates_log = pd.DataFrame(rows)

    helper_cols = [c for c in kept.columns if c.startswith("_")]
    deduped = kept.drop(columns=helper_cols, errors="ignore").copy()

    summary = {
        "deduplication_applied": True,
        "deduplication_mode": DEDUPLICATION_MODE,
        "keep_policy": KEEP_POLICY,
        "input_rows": int(len(df)),
        "output_rows": int(len(deduped)),
        "rows_removed": int(len(df) - len(deduped)),
        "removed_duplicate_valid_doi": int(
            (removed.get("_duplicate_reason", pd.Series(dtype="string")) == "duplicate_valid_doi").sum()
        ) if not removed.empty else 0,
        "removed_duplicate_title_year_fallback": int(
            (removed.get("_duplicate_reason", pd.Series(dtype="string")) == "duplicate_title_year_fallback").sum()
        ) if not removed.empty else 0,
    }

    return deduped, duplicates_log, summary

# ============================================================
# QUALITY REPORT (SIMPLE) + MISSINGNESS
# ============================================================

def compute_year_numeric(series: pd.Series) -> pd.Series:
    y = series.astype("string").str.strip()
    return pd.to_numeric(y.where(y.str.fullmatch(r"\d{4}").fillna(False)), errors="coerce")

def build_quality_report_simple(
    merged_df: pd.DataFrame,
    final_df: pd.DataFrame,
    dedupe_summary: Dict[str, Any]
) -> pd.DataFrame:
    doi_valid = valid_doi_mask(final_df["DOI"])
    y_num = compute_year_numeric(final_df["Publication Year"])

    metrics = {
        "records_after_merge": int(len(merged_df)),
        "duplicates_removed": int(dedupe_summary.get("rows_removed", 0)),
        "records_after_dedup": int(len(final_df)),
        "missing_title_count": int(final_df["Title"].isna().sum()),
        "missing_abstract_count": int(final_df["Abstract"].isna().sum()),
        "missing_doi_count": int(final_df["DOI"].isna().sum()),
        "invalid_doi_count": int((final_df["DOI"].notna() & ~doi_valid).sum()),
        "missing_year_count": int(final_df["Publication Year"].isna().sum()),
        "year_out_of_range_count": int(((y_num.notna()) & ((y_num < MIN_YEAR) | (y_num > MAX_YEAR))).sum()),
    }
    return pd.DataFrame({"metric": list(metrics.keys()), "value": list(metrics.values())})

def build_metric_definitions_en() -> pd.DataFrame:
    rows = [
        {"metric": "records_after_merge", "definition": "Records after importing and merging all input files (before deduplication).", "why_it_matters": "PRISMA: records identified."},
        {"metric": "duplicates_removed", "definition": "Records removed by deduplication.", "why_it_matters": "PRISMA: duplicates removed."},
        {"metric": "records_after_dedup", "definition": "Unique records kept after deduplication.", "why_it_matters": "Working set for Title/Abstract screening."},
        {"metric": "missing_title_count", "definition": "Records with missing Title.", "why_it_matters": "Impairs screening and traceability."},
        {"metric": "missing_abstract_count", "definition": "Records with missing Abstract.", "why_it_matters": "Impairs screening quality; may require early full-text retrieval."},
        {"metric": "missing_doi_count", "definition": "Records with missing DOI.", "why_it_matters": "DOI supports deduplication and retrieval."},
        {"metric": "invalid_doi_count", "definition": "Records with a non-standard DOI format.", "why_it_matters": "May break deduplication/linking; needs manual correction."},
        {"metric": "missing_year_count", "definition": "Records with missing Publication Year.", "why_it_matters": "Used for reporting and QC."},
        {"metric": "year_out_of_range_count", "definition": f"Records with Publication Year outside [{MIN_YEAR}, {MAX_YEAR}].", "why_it_matters": "Often indicates parsing/metadata issues."},
    ]
    return pd.DataFrame(rows)

def missing_sheet_base_cols(df: pd.DataFrame) -> List[str]:
    want = [
        "record_id",
        "DOI",
        "DOI Link",
        "Title",
        "Publication Year",
        "source_record_id",
        "source_row_index",
        "source_path",
    ]
    return [c for c in want if c in df.columns]

def build_missingness_sheets(final_df: pd.DataFrame) -> Dict[str, pd.DataFrame]:
    doi_valid = valid_doi_mask(final_df["DOI"])
    y_num = compute_year_numeric(final_df["Publication Year"])

    def sel(mask: pd.Series) -> pd.DataFrame:
        cols = missing_sheet_base_cols(final_df)
        return final_df.loc[mask, cols].copy()

    out: Dict[str, pd.DataFrame] = {
        "missing_title": sel(final_df["Title"].isna()),
        "missing_abstract": sel(final_df["Abstract"].isna()),
        "missing_doi": sel(final_df["DOI"].isna()),
        "invalid_doi": sel(final_df["DOI"].notna() & ~doi_valid),
        "missing_year": sel(final_df["Publication Year"].isna()),
        "year_out_of_range": sel((y_num.notna()) & ((y_num < MIN_YEAR) | (y_num > MAX_YEAR))),
    }

    for k, df_k in out.items():
        if df_k.empty:
            out[k] = pd.DataFrame(columns=missing_sheet_base_cols(final_df))

    return out

# ============================================================
# SEARCH PROTOCOL TEMPLATE
# ============================================================

def build_search_protocol_template_df() -> pd.DataFrame:
    return pd.DataFrame([["", "", "", ""] for _ in range(15)])

def format_search_protocol_sheet(ws, source_name: str = "WoS") -> None:
    from openpyxl.styles import Font

    title_text = f"Search protocol in {source_name}"

    for r in range(1, 20):
        for c in range(1, 5):
            ws.cell(row=r, column=c).value = None

    ws["B2"] = "Question:"
    ws["B4"] = title_text
    ws["B5"] = "Database:"
    ws["B6"] = "Editions:"
    ws["B7"] = "Search string:"
    ws["B8"] = "Field:"
    ws["B9"] = "Time limit:"
    ws["B10"] = "Search date:"
    ws["B11"] = "Search location:"
    ws["B12"] = "Number of records:"

    bold = Font(bold=True)
    ws["B2"].font = bold
    ws["B4"].font = bold
    ws.freeze_panes = "B2"

# ============================================================
# SCREENING VIEW + LOOKUP FORMULAS
# ============================================================

def build_screening_view_df(full_df: pd.DataFrame) -> pd.DataFrame:
    cols = [
        "source_record_id",  # hidden technical key
        "record_id",
        "DOI",
        "DOI Link",
        "Title",
        "Abstract",
        "Label",
        "Reason to Exclude",
        "Comment",
    ]
    sv = full_df[cols].copy()

    # IMPORTANT: empty strings => full remains blank until user edits screening_view
    sv["Label"] = ""
    sv["Reason to Exclude"] = ""
    sv["Comment"] = ""

    return sv

def excel_col_letter(n: int) -> str:
    result = ""
    while n > 0:
        n, rem = divmod(n - 1, 26)
        result = chr(65 + rem) + result
    return result

def add_lookup_formulas_to_full_sheet(ws_full, full_columns: List[str], screening_columns: List[str]) -> None:
    """
    In wos_full, fill Label / Reason to Exclude / Comment with formulas that lookup
    values from wos_screening_view by source_record_id.

    Uses XLOOKUP with INDEX/MATCH fallback and preserves blanks ("").
    """
    full_col_idx = {c: i + 1 for i, c in enumerate(full_columns)}
    sv_col_idx = {c: i + 1 for i, c in enumerate(screening_columns)}

    key_col_full = excel_col_letter(full_col_idx["source_record_id"])
    key_col_sv = excel_col_letter(sv_col_idx["source_record_id"])

    key_range = f"'{SCREENING_SHEET_NAME}'!${key_col_sv}:${key_col_sv}"

    for target_col in ["Label", "Reason to Exclude", "Comment"]:
        ret_col_sv = excel_col_letter(sv_col_idx[target_col])
        ret_range = f"'{SCREENING_SHEET_NAME}'!${ret_col_sv}:${ret_col_sv}"

        col_letter_full = excel_col_letter(full_col_idx[target_col])

        for row in range(2, ws_full.max_row + 1):
            key_ref = f"${key_col_full}{row}"

            # Blank-safe formula:
            # - If lookup returns "", keep "" (prevents Excel showing 0)
            # - If not found, keep ""
            formula = (
                f'=IFERROR('
                f'IF(XLOOKUP({key_ref},{key_range},{ret_range},"")="","",XLOOKUP({key_ref},{key_range},{ret_range},"")),'
                f'IFERROR('
                f'IF(INDEX({ret_range},MATCH({key_ref},{key_range},0))="","",INDEX({ret_range},MATCH({key_ref},{key_range},0))),'
                f'""'
                f')'
                f')'
            )
            ws_full[f"{col_letter_full}{row}"] = formula

# ============================================================
# CONDITIONAL FORMATTING (LABEL-BASED)
# ============================================================

def apply_label_row_conditional_formatting(ws, label_column_name: str = "Label") -> None:
    from openpyxl.formatting.rule import FormulaRule
    from openpyxl.styles import PatternFill

    if ws.max_row < 2 or ws.max_column < 1:
        return

    headers = {ws.cell(row=1, column=i).value: i for i in range(1, ws.max_column + 1)}
    if label_column_name not in headers:
        return

    label_col_letter = excel_col_letter(headers[label_column_name])
    start_col = excel_col_letter(1)
    end_col = excel_col_letter(ws.max_column)

    target_range = f"${start_col}$2:${end_col}${ws.max_row}"

    include_fill = PatternFill(fill_type="solid", fgColor="C6EFCE")
    exclude_fill = PatternFill(fill_type="solid", fgColor="FFC7CE")

    include_rule = FormulaRule(formula=[f"${label_col_letter}2=1"], fill=include_fill, stopIfTrue=False)
    exclude_rule = FormulaRule(formula=[f"${label_col_letter}2=0"], fill=exclude_fill, stopIfTrue=False)

    ws.conditional_formatting.add(target_range, include_rule)
    ws.conditional_formatting.add(target_range, exclude_rule)

# ============================================================
# EXCEL WRITING / MINIMAL FORMATTING
# ============================================================

def bold_header_row(ws) -> None:
    from openpyxl.styles import Font
    bold = Font(bold=True)
    for cell in ws[1]:
        cell.font = bold

def set_fixed_row_heights(ws, start_row: int, end_row: int, height: float) -> None:
    for r in range(start_row, end_row + 1):
        ws.row_dimensions[r].height = height

def apply_wrap_on_columns(ws, column_names: List[str]) -> None:
    from openpyxl.styles import Alignment
    headers = {ws.cell(row=1, column=i).value: i for i in range(1, ws.max_column + 1)}
    for name in column_names:
        if name not in headers:
            continue
        cidx = headers[name]
        for r in range(2, ws.max_row + 1):
            ws.cell(row=r, column=cidx).alignment = Alignment(
                horizontal="left", vertical="top", wrap_text=True
            )

def write_main_workbook(
    path_xlsx: Path,
    full_df: pd.DataFrame,
    screening_df: pd.DataFrame,
    file_summary_df: pd.DataFrame,
    input_manifest_df: pd.DataFrame,
    search_protocol_df: pd.DataFrame
) -> None:
    from openpyxl.worksheet.datavalidation import DataValidation

    path_xlsx.parent.mkdir(parents=True, exist_ok=True)

    with pd.ExcelWriter(path_xlsx, engine="openpyxl") as writer:
        search_protocol_df.to_excel(writer, index=False, header=False, sheet_name=SEARCH_PROTOCOL_SHEET_NAME)
        full_df.to_excel(writer, index=False, sheet_name=FULL_SHEET_NAME, na_rep="")
        screening_df.to_excel(writer, index=False, sheet_name=SCREENING_SHEET_NAME, na_rep="")
        file_summary_df.to_excel(writer, index=False, sheet_name="file_processing_summary", na_rep="")
        input_manifest_df.to_excel(writer, index=False, sheet_name="input_manifest", na_rep="")

        ws_sp = writer.sheets[SEARCH_PROTOCOL_SHEET_NAME]
        ws_full = writer.sheets[FULL_SHEET_NAME]
        ws_sv = writer.sheets[SCREENING_SHEET_NAME]
        ws_fs = writer.sheets["file_processing_summary"]
        ws_im = writer.sheets["input_manifest"]

        format_search_protocol_sheet(ws_sp, source_name=SOURCE_NAME)

        # Headers bold only
        bold_header_row(ws_full)
        bold_header_row(ws_sv)
        bold_header_row(ws_fs)
        bold_header_row(ws_im)

        # Fill formulas in full (after both sheets exist)
        add_lookup_formulas_to_full_sheet(
            ws_full=ws_full,
            full_columns=list(full_df.columns),
            screening_columns=list(screening_df.columns),
        )

        # Freeze panes so record_id is fixed (column B). With hidden key in col A => freeze at C2
        ws_sv.freeze_panes = "C2"

        headers_sv = {ws_sv.cell(row=1, column=i).value: i for i in range(1, ws_sv.max_column + 1)}

        # Hide technical key in screening view (column A)
        if "source_record_id" in headers_sv:
            ws_sv.column_dimensions[excel_col_letter(headers_sv["source_record_id"])].hidden = True

        # Prevent spillover: wrap Title/Abstract but fixed row height
        apply_wrap_on_columns(ws_sv, ["Title", "Abstract"])
        if ws_sv.max_row >= 2:
            set_fixed_row_heights(ws_sv, 2, ws_sv.max_row, SCREENING_VIEW_ROW_HEIGHT)

        # Data validation for Label: only 0/1 or blank
        if "Label" in headers_sv and ws_sv.max_row >= 2:
            label_col_letter = excel_col_letter(headers_sv["Label"])
            dv = DataValidation(type="list", formula1='"0,1"', allow_blank=True)
            dv.promptTitle = "Label"
            dv.prompt = "Use only: 1 = Include, 0 = Exclude (or leave blank)."
            dv.errorTitle = "Invalid Label"
            dv.error = "Only 0 or 1 is allowed in Label (or leave blank)."
            ws_sv.add_data_validation(dv)
            dv.add(f"{label_col_letter}2:{label_col_letter}{ws_sv.max_row}")

        # Hide technical columns in full by default
        headers_full = {ws_full.cell(row=1, column=i).value: i for i in range(1, ws_full.max_column + 1)}
        for c in FULL_HIDDEN_COLUMNS_DEFAULT:
            if c in headers_full:
                ws_full.column_dimensions[excel_col_letter(headers_full[c])].hidden = True

        # Conditional formatting (full + screening)
        apply_label_row_conditional_formatting(ws_full, label_column_name="Label")
        apply_label_row_conditional_formatting(ws_sv, label_column_name="Label")

        # Hide audit sheets by default
        ws_fs.sheet_state = "hidden"
        ws_im.sheet_state = "hidden"

def write_quality_workbook(
    path_xlsx: Path,
    quality_report_df: pd.DataFrame,
    metric_definitions_df: pd.DataFrame,
    missing_sheets: Dict[str, pd.DataFrame]
) -> None:
    path_xlsx.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(path_xlsx, engine="openpyxl") as writer:
        quality_report_df.to_excel(writer, index=False, sheet_name="quality_report", na_rep="")
        metric_definitions_df.to_excel(writer, index=False, sheet_name="metric_definitions", na_rep="")
        for name, df in missing_sheets.items():
            df.to_excel(writer, index=False, sheet_name=name, na_rep="")

        bold_header_row(writer.sheets["quality_report"])
        bold_header_row(writer.sheets["metric_definitions"])
        for name in missing_sheets.keys():
            bold_header_row(writer.sheets[name])

def write_duplicates_workbook(path_xlsx: Path, df: pd.DataFrame) -> None:
    path_xlsx.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(path_xlsx, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name="duplicates_log", na_rep="")
        bold_header_row(writer.sheets["duplicates_log"])

# ============================================================
# MAIN
# ============================================================

def main() -> None:
    validate_user_inputs()

    run_id = now_run_id()
    run_dir = OUTPUT_RUNS_DIR / run_id
    run_dir.mkdir(parents=True, exist_ok=True)

    run_log: Dict[str, Any] = {
        "run_id": run_id,
        "script_name": "build_wos_workbook.py",
        "script_version": SCRIPT_VERSION,
        "source": SOURCE_NAME,
        "started_at": datetime.now().isoformat(timespec="seconds"),
        "status": "running",
        "configuration": {
            "INPUT_DIR": str(INPUT_DIR.resolve()),
            "OUTPUT_RUNS_DIR": str(OUTPUT_RUNS_DIR.resolve()),
            "FILE_PATTERNS": FILE_PATTERNS,
            "EXPORT_CSV": EXPORT_CSV,
            "INCLUDE_ISSUE": INCLUDE_ISSUE,
            "ENABLE_DEDUPLICATION": ENABLE_DEDUPLICATION,
            "DEDUPLICATION_MODE": DEDUPLICATION_MODE,
            "KEEP_POLICY": KEEP_POLICY,
            "CONTINUE_ON_FILE_ERROR": CONTINUE_ON_FILE_ERROR,
        },
        "environment": {
            "python_version": sys.version,
            "platform": platform.platform(),
            "pandas_version": pd.__version__,
        },
        "files_discovered": [],
        "files_processed_ok": [],
        "files_failed": [],
        "summary": {},
        "outputs": {},
    }

    friendly_print("============================================================")
    friendly_print("BUILD WOS WORKBOOK")
    friendly_print("============================================================")
    friendly_print(f"Run ID: {run_id}")
    friendly_print(f"Run folder: {run_dir}")
    friendly_print("")

    friendly_print("Searching for input files...")
    files = list_input_files(INPUT_DIR, FILE_PATTERNS)
    if not files:
        err = f"No input files found in {INPUT_DIR} matching patterns {FILE_PATTERNS}"
        run_log["status"] = "failed"
        run_log["ended_at"] = datetime.now().isoformat(timespec="seconds")
        run_log["summary"]["error"] = err
        write_json(run_dir / "run_log.json", run_log)
        raise FileNotFoundError(err)

    run_log["files_discovered"] = [str(f.resolve()) for f in files]
    friendly_print(f"Files found: {len(files)}")

    friendly_print("Creating input manifest...")
    input_manifest_df = build_input_manifest(files)

    frames: List[pd.DataFrame] = []
    file_rows: List[Dict[str, Any]] = []

    friendly_print("Processing files...")
    for i, f in enumerate(files, start=1):
        friendly_print(f"  [{i}/{len(files)}] {f.name}")
        try:
            df_part, summ = process_file(f, run_id=run_id)
            frames.append(df_part)
            file_rows.append(summ)
            run_log["files_processed_ok"].append(summ)
        except Exception as e:
            err = {
                "file_name": f.name,
                "file_path": str(f.resolve()),
                "error_type": type(e).__name__,
                "error_message": str(e),
                "traceback": traceback.format_exc(limit=5),
            }
            file_rows.append({
                "file_name": f.name,
                "file_path": str(f.resolve()),
                "rows_raw": None,
                "rows_cleaned": None,
                "columns_raw": None,
                "columns_final": None,
                "status": "failed",
            })
            run_log["files_failed"].append(err)
            friendly_print(f"      ERROR: {type(e).__name__}: {e}")
            if not CONTINUE_ON_FILE_ERROR:
                run_log["status"] = "failed"
                run_log["ended_at"] = datetime.now().isoformat(timespec="seconds")
                write_json(run_dir / "run_log.json", run_log)
                raise

    if not frames:
        err = "No files were processed successfully."
        run_log["status"] = "failed"
        run_log["ended_at"] = datetime.now().isoformat(timespec="seconds")
        run_log["summary"]["error"] = err
        write_json(run_dir / "run_log.json", run_log)
        raise RuntimeError(err)

    file_summary_df = pd.DataFrame(file_rows)

    friendly_print("Merging processed records...")
    merged_df = pd.concat(frames, ignore_index=True)
    merged_df = select_and_order_full_schema(merged_df, include_issue=INCLUDE_ISSUE)

    duplicates_log_df = pd.DataFrame()
    dedupe_summary = {
        "deduplication_applied": False,
        "deduplication_mode": None,
        "keep_policy": None,
        "input_rows": int(len(merged_df)),
        "output_rows": int(len(merged_df)),
        "rows_removed": 0,
    }

    final_df = merged_df.copy()
    if ENABLE_DEDUPLICATION:
        friendly_print("Applying deduplication...")
        final_df, duplicates_log_df, dedupe_summary = deduplicate_records(merged_df)
    else:
        friendly_print("Deduplication disabled.")

    final_df = select_and_order_full_schema(final_df, include_issue=INCLUDE_ISSUE)

    # record_id AFTER dedupe
    final_df.insert(0, "record_id", range(1, len(final_df) + 1))

    screening_df = build_screening_view_df(final_df)
    search_protocol_df = build_search_protocol_template_df()

    quality_report_df = build_quality_report_simple(merged_df=merged_df, final_df=final_df, dedupe_summary=dedupe_summary)
    metric_definitions_df = build_metric_definitions_en()
    missing_sheets = build_missingness_sheets(final_df)

    if duplicates_log_df.empty:
        duplicates_log_df = pd.DataFrame(columns=[
            "duplicate_reason", "duplicate_group_key",
            "removed_source_record_id", "removed_source_row_index",
            "removed_title", "removed_publication_year", "removed_doi", "removed_record_score",
            "retained_source_record_id", "retained_source_row_index",
            "retained_title", "retained_publication_year", "retained_doi", "retained_record_score"
        ])

    friendly_print("Writing output files...")
    path_01 = run_dir / "01_wos_workbook.xlsx"
    path_02 = run_dir / "02_quality_report.xlsx"
    path_03 = run_dir / "03_duplicates_log.xlsx"
    path_log = run_dir / "run_log.json"

    write_main_workbook(
        path_xlsx=path_01,
        full_df=final_df,
        screening_df=screening_df,
        file_summary_df=file_summary_df,
        input_manifest_df=input_manifest_df,
        search_protocol_df=search_protocol_df,
    )
    write_quality_workbook(
        path_xlsx=path_02,
        quality_report_df=quality_report_df,
        metric_definitions_df=metric_definitions_df,
        missing_sheets=missing_sheets,
    )
    write_duplicates_workbook(path_xlsx=path_03, df=duplicates_log_df)

    run_log["status"] = "completed"
    run_log["ended_at"] = datetime.now().isoformat(timespec="seconds")
    run_log["summary"] = {
        "files_found": int(len(files)),
        "files_processed_ok": int((file_summary_df["status"] == "ok").sum()) if not file_summary_df.empty else 0,
        "files_failed": int((file_summary_df["status"] == "failed").sum()) if not file_summary_df.empty else 0,
        "rows_after_merge_before_deduplication": int(len(merged_df)),
        "rows_after_final_output": int(len(final_df)),
        "deduplication": dedupe_summary,
    }
    run_log["outputs"] = {
        "run_folder": str(run_dir.resolve()),
        "01_merged_cleaned_xlsx": str(path_01.resolve()),
        "02_quality_report_xlsx": str(path_02.resolve()),
        "03_duplicates_log_xlsx": str(path_03.resolve()),
        "run_log_json": str(path_log.resolve()),
    }
    write_json(path_log, run_log)

    friendly_print("")
    friendly_print("============================================================")
    friendly_print("DONE")
    friendly_print("============================================================")
    friendly_print(f"Main workbook: {path_01}")
    friendly_print(f"Rows in final wos_full: {len(final_df)}")
    friendly_print(f"Rows in wos_screening_view: {len(screening_df)}")
    if ENABLE_DEDUPLICATION:
        friendly_print(f"Rows removed by deduplication: {dedupe_summary.get('rows_removed', 0)}")
    friendly_print(f"Failed files: {len(run_log['files_failed'])}")
    friendly_print("")

if __name__ == "__main__":
    main()