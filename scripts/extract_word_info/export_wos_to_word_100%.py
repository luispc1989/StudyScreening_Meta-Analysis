from __future__ import annotations

import re
from pathlib import Path
from typing import Optional, Tuple, List, Dict

from openpyxl import load_workbook
from docx import Document

# =======================
# CONFIG
# =======================

EXCEL_PATH = Path(
    r"C:\Users\Luís Pinto Coelho\Desktop\Dissertação\metanalysis development\2 - literature search & study screening\data\wos\excel_files\WoS\wos_workbook.xlsx"
).resolve()

OUTPUT_DOCX = Path(
    r"C:\Users\Luís Pinto Coelho\Desktop\Dissertação\metanalysis development\2 - literature search & study screening\data\wos\excel_files\WoS\wos_titles_abstracts_full_v2.docx"
).resolve()

SHEET_NAME = "wos_full"

ROW_START = 2
ROW_END: Optional[int] = None  # If None, reads until the last available row

TITLE_COLUMN_NAME: Optional[str] = None      # Example: "Title"
ABSTRACT_COLUMN_NAME: Optional[str] = None   # Example: "Abstract"
STATUS_COLUMN_NAME: Optional[str] = None     # Example: "Label"

GROUP_SIZE = 10


def _norm_header(value) -> str:
    """Normalize header text for robust matching."""
    if value is None:
        return ""
    s = str(value).strip().lower()
    s = re.sub(r"\s+", " ", s)
    return s


def _find_columns_from_headers(
    header_values: List[object],
    title_name: Optional[str],
    abstract_name: Optional[str],
    status_name: Optional[str],
) -> Tuple[int, int, Optional[int]]:
    """
    Detect the Title, Abstract and optional Status columns from the header row.
    Returns 1-based column indexes: (title_col, abstract_col, status_col)
    """
    headers = {
        idx: _norm_header(value)
        for idx, value in enumerate(header_values, start=1)
    }

    # Title column
    if title_name:
        t_norm = _norm_header(title_name)
        title_col = next((c for c, h in headers.items() if h == t_norm), None)
    else:
        title_col = next(
            (
                c for c, h in headers.items()
                if h == "title"
                or "title" in h
                or h == "article title"
                or h == "document title"
                or h == "ti"
            ),
            None,
        )

    # Abstract column
    if abstract_name:
        a_norm = _norm_header(abstract_name)
        abstract_col = next((c for c, h in headers.items() if h == a_norm), None)
    else:
        abstract_col = next(
            (
                c for c, h in headers.items()
                if h == "abstract"
                or "abstract" in h
                or h == "ab"
            ),
            None,
        )

    # Status column
    status_col = None
    if status_name:
        s_norm = _norm_header(status_name)
        status_col = next((c for c, h in headers.items() if h == s_norm), None)
    else:
        status_col = next(
            (
                c for c, h in headers.items()
                if h in {"label", "include/exclude", "include exclude"}
                or "label" in h
                or "include/exclude" in h
                or "include exclude" in h
            ),
            None,
        )

    if not title_col or not abstract_col:
        raise ValueError(
            "Could not detect the Title and/or Abstract columns.\n"
            "Set TITLE_COLUMN_NAME and ABSTRACT_COLUMN_NAME with the exact header names from row 1.\n"
            f"Detected headers (column -> normalized header): {headers}"
        )

    return title_col, abstract_col, status_col


def read_articles(excel_path: Path) -> List[Dict]:
    """
    Read articles from the Excel workbook using sequential iteration.
    This avoids the performance issues caused by ws.cell(...) in read_only mode.
    """
    wb = load_workbook(excel_path, data_only=True, read_only=True)

    try:
        if SHEET_NAME not in wb.sheetnames:
            raise ValueError(
                f"Sheet '{SHEET_NAME}' not found. Available sheets: {wb.sheetnames}"
            )

        ws = wb[SHEET_NAME]

        rows_iter = ws.iter_rows(values_only=True)

        # Read header row
        try:
            header_row = next(rows_iter)
        except StopIteration:
            raise ValueError(f"Sheet '{SHEET_NAME}' is empty.")

        title_col, abstract_col, status_col = _find_columns_from_headers(
            list(header_row),
            TITLE_COLUMN_NAME,
            ABSTRACT_COLUMN_NAME,
            STATUS_COLUMN_NAME,
        )

        articles: List[Dict] = []

        # Excel row number starts at 2 because row 1 is the header
        for excel_row_num, row_values in enumerate(rows_iter, start=2):
            if excel_row_num < ROW_START:
                continue

            if ROW_END is not None and excel_row_num > ROW_END:
                break

            # Convert row tuple to list for safe indexed access
            row_list = list(row_values)

            def get_value(col_index_1_based: Optional[int]) -> object:
                if col_index_1_based is None:
                    return None
                idx0 = col_index_1_based - 1
                if 0 <= idx0 < len(row_list):
                    return row_list[idx0]
                return None

            title_val = get_value(title_col)
            abstract_val = get_value(abstract_col)
            status_val = get_value(status_col)

            title = "" if title_val is None else str(title_val).strip()
            abstract = "" if abstract_val is None else str(abstract_val).strip()

            # Skip fully empty rows
            if not title and not abstract:
                continue

            evaluated = status_val is not None and str(status_val).strip() != ""

            articles.append(
                {
                    "row": excel_row_num,
                    "title": title,
                    "abstract": abstract,
                    "evaluated": evaluated,
                }
            )

        return articles

    finally:
        wb.close()


def add_blank_page(doc: Document) -> None:
    """Insert a blank page."""
    doc.add_page_break()
    doc.add_paragraph("")
    doc.add_page_break()


def write_word(articles: List[Dict], output_docx: Path, group_size: int = 10) -> None:
    """Write the articles to a Word document."""
    doc = Document()

    articles_sorted = sorted(articles, key=lambda x: x["row"])

    for idx, article in enumerate(articles_sorted, start=1):
        row_num = article["row"]
        title = article["title"]
        abstract = article["abstract"]
        evaluated = article["evaluated"]

        title_text = f"{row_num}. {title}"
        if evaluated:
            title_text += " (ALREADY EVALUATED)"

        p_title = doc.add_paragraph()
        run = p_title.add_run(title_text)
        run.bold = True

        doc.add_paragraph(abstract)

        if (idx % group_size == 0) and (idx != len(articles_sorted)):
            add_blank_page(doc)

    output_docx.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output_docx)


def main() -> None:
    print(f"EXCEL_PATH  = {EXCEL_PATH}")
    print(f"OUTPUT_DOCX = {OUTPUT_DOCX}")
    print(f"EXCEL EXISTS? {EXCEL_PATH.exists()}")

    if not EXCEL_PATH.exists():
        raise FileNotFoundError(
            f"Excel file not found:\n{EXCEL_PATH}\n\n"
            "Check that the file exists at this exact location and that the filename is correct."
        )

    articles = read_articles(EXCEL_PATH)

    if not articles:
        raise ValueError(
            "No articles were found to export. "
            "Check the sheet name, row range, and Title/Abstract column detection."
        )

    write_word(articles, OUTPUT_DOCX, group_size=GROUP_SIZE)

    print(f"Word file created successfully: {OUTPUT_DOCX}")
    print(f"Articles exported: {len(articles)}")


if __name__ == "__main__":
    main()