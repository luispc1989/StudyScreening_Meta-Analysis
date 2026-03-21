from __future__ import annotations

import re
from pathlib import Path
from typing import Optional, Tuple, List, Dict

from openpyxl import load_workbook
from docx import Document


# =======================
# CONFIG (READY TO RUN)
# =======================
BASE_DIR = Path(r"C:\Users\Luís Pinto Coelho\Desktop\Dissertação\Dados\excel_files\WoS").resolve()

EXCEL_FILENAME = "matriz_wos_20%_10.01.2026.xlsx"
OUTPUT_DOCX_NAME = "wos_titles_abstracts_20pct.docx"

SHEET_NAME = "WoS_Merged_20%"
ROW_START = 2
ROW_END = 484

# If auto-detection fails, set these to the exact header names in row 1
TITLE_COLUMN_NAME: Optional[str] = None
ABSTRACT_COLUMN_NAME: Optional[str] = None

GROUP_SIZE = 10  # 10 em 10 artigos


EXCEL_PATH = BASE_DIR / EXCEL_FILENAME
OUTPUT_DOCX = BASE_DIR / OUTPUT_DOCX_NAME


def _norm_header(value) -> str:
    if value is None:
        return ""
    s = str(value).strip().lower()
    s = re.sub(r"\s+", " ", s)
    return s


def _find_columns(ws, title_name: Optional[str], abstract_name: Optional[str]) -> Tuple[int, int]:
    headers = {col: _norm_header(ws.cell(row=1, column=col).value) for col in range(1, ws.max_column + 1)}

    if title_name:
        t_norm = _norm_header(title_name)
        title_col = next((c for c, h in headers.items() if h == t_norm), None)
    else:
        title_col = next((c for c, h in headers.items() if ("title" in h) or (h == "ti")), None)

    if abstract_name:
        a_norm = _norm_header(abstract_name)
        abstract_col = next((c for c, h in headers.items() if h == a_norm), None)
    else:
        abstract_col = next((c for c, h in headers.items() if ("abstract" in h) or (h == "ab")), None)

    if not title_col or not abstract_col:
        raise ValueError(
            "Não foi possível detetar as colunas de Título/Abstract.\n"
            "Defina TITLE_COLUMN_NAME e ABSTRACT_COLUMN_NAME com o nome exato das colunas (linha 1).\n"
            f"Cabeçalhos detetados (col -> header normalizado): {headers}"
        )

    return title_col, abstract_col


def read_articles(excel_path: Path) -> List[Dict]:
    wb = load_workbook(excel_path, data_only=True, read_only=True)
    if SHEET_NAME not in wb.sheetnames:
        raise ValueError(f"Aba '{SHEET_NAME}' não existe. Abas disponíveis: {wb.sheetnames}")

    ws = wb[SHEET_NAME]
    title_col, abstract_col = _find_columns(ws, TITLE_COLUMN_NAME, ABSTRACT_COLUMN_NAME)

    articles: List[Dict] = []
    for r in range(ROW_START, ROW_END + 1):
        title = ws.cell(row=r, column=title_col).value
        abstract = ws.cell(row=r, column=abstract_col).value

        title_s = "" if title is None else str(title).strip()
        abstract_s = "" if abstract is None else str(abstract).strip()

        if not title_s and not abstract_s:
            continue

        articles.append({"row": r, "title": title_s, "abstract": abstract_s})

    wb.close()
    return articles


def add_blank_page(doc: Document) -> None:
    """
    Ensures a full blank page between groups:
    - Page break to move to new page
    - Add an empty paragraph (keeps the page "occupied" as blank)
    - Another page break to start next content after the blank page
    """
    doc.add_page_break()
    doc.add_paragraph("")
    doc.add_page_break()


def write_word(articles: List[Dict], output_docx: Path, group_size: int = 10) -> None:
    doc = Document()

    # Ordem: 484 -> 2 (primeiros 10: 484..475)
    articles_sorted = sorted(articles, key=lambda x: x["row"], reverse=True)

    for idx, a in enumerate(articles_sorted, start=1):
        p_title = doc.add_paragraph()
        run = p_title.add_run(f'{a["row"]}. {a["title"]}'.strip())
        run.bold = True

        doc.add_paragraph(a["abstract"])

        # After each group of 10, insert a FULL BLANK PAGE (except after last article)
        if (idx % group_size == 0) and (idx != len(articles_sorted)):
            add_blank_page(doc)

    doc.save(output_docx)


def main() -> None:
    if not EXCEL_PATH.exists():
        raise FileNotFoundError(
            f"Excel não encontrado: {EXCEL_PATH}\n"
            "Confirme se o nome do ficheiro tem mesmo extensão .xlsx e se está nessa pasta."
        )

    articles = read_articles(EXCEL_PATH)
    write_word(articles, OUTPUT_DOCX, group_size=GROUP_SIZE)

    print(f"Word criado em: {OUTPUT_DOCX}")
    print(f"Artigos exportados: {len(articles)}")


if __name__ == "__main__":
    main()
