from __future__ import annotations

import re
import time
from pathlib import Path
from typing import Any, Dict, List, Optional, Sequence

from openpyxl import load_workbook

from app.config import (
    BENCHMARK_RECORD_IDS,
    ENABLE_BENCHMARK_MODE,
    MAX_ROWS_TO_PROCESS,
    PDF_BASE_DIR,
    REQUIRED_COLUMNS,
    SHEET_NAME,
    build_final_workbook_path,
    build_temp_workbook_path,
)
from app.models import DownloadResult, Record, SessionState
from app.utils import build_pdf_filename


def load_workbook_and_sheet(workbook_path: Path, sheet_name: str):
    """
    Open the Excel workbook and return (workbook, worksheet).
    """
    if not workbook_path.exists():
        raise FileNotFoundError(f"Workbook not found: {workbook_path}")

    wb = load_workbook(workbook_path)

    if sheet_name not in wb.sheetnames:
        raise ValueError(f"Sheet not found: {sheet_name}")

    ws = wb[sheet_name]
    return wb, ws


def build_col_map(ws) -> Dict[str, int]:
    """
    Build a mapping from header name to Excel column index.
    """
    col_map: Dict[str, int] = {}

    for col_idx in range(1, ws.max_column + 1):
        header = ws.cell(row=1, column=col_idx).value
        if header is not None and str(header).strip():
            col_map[str(header).strip()] = col_idx

    return col_map


def validate_required_columns(
    col_map: Dict[str, int],
    sheet_name: str,
    required_columns: Optional[Sequence[str]] = None,
) -> None:
    """
    Ensure all required columns exist in the target sheet.
    """
    required = list(required_columns or REQUIRED_COLUMNS)
    missing = [c for c in required if c not in col_map]
    if missing:
        raise ValueError(f"Missing required columns in '{sheet_name}': {missing}")


def get_cell(ws, row_idx: int, col_map: Dict[str, int], header: str):
    return ws.cell(row=row_idx, column=col_map[header]).value


def get_optional_cell(ws, row_idx: int, col_map: Dict[str, int], header: str, default: Any = ""):
    """
    Return the cell value if the column exists, otherwise return default.
    """
    if header not in col_map:
        return default
    return ws.cell(row=row_idx, column=col_map[header]).value


def set_cell(ws, row_idx: int, col_map: Dict[str, int], header: str, value) -> None:
    ws.cell(row=row_idx, column=col_map[header], value=value)


def extract_row_values(
    ws,
    row_idx: int,
    col_map: Dict[str, int],
    headers: Sequence[str],
) -> Dict[str, Any]:
    """
    Extract a dictionary of values for the requested headers.
    Missing columns are returned as empty strings.
    """
    return {header: get_optional_cell(ws, row_idx, col_map, header, "") for header in headers}


def is_blank(value) -> bool:
    return value is None or str(value).strip() == ""


def normalize_doi_like(value: str) -> Optional[str]:
    """
    Lightweight DOI normalizer for workbook eligibility checks.
    """
    if is_blank(value):
        return None

    doi = str(value).strip()
    doi = re.sub(r"^https?://(dx\.)?doi\.org/", "", doi, flags=re.IGNORECASE).strip()

    if re.match(r"^10\.\d{4,9}/\S+$", doi):
        return doi

    return None


def row_has_usable_doi(ws, row_idx: int, col_map: Dict[str, int]) -> bool:
    """
    Return True when the row already has a usable DOI or DOI link.
    """
    doi_raw = str(get_optional_cell(ws, row_idx, col_map, "DOI", "") or "").strip()
    doi_link_raw = str(get_optional_cell(ws, row_idx, col_map, "DOI Link", "") or "").strip()

    if doi_link_raw:
        return True

    return normalize_doi_like(doi_raw) is not None


def row_is_already_downloaded(ws, row_idx: int, col_map: Dict[str, int]) -> bool:
    """
    Return True when the row already has a downloaded PDF according to workbook state.
    """
    pdf_downloaded = get_optional_cell(ws, row_idx, col_map, "pdf_downloaded", "")
    pdf_status = str(get_optional_cell(ws, row_idx, col_map, "pdf_download_status", "") or "").strip().lower()

    if str(pdf_downloaded).strip() == "1":
        return True

    return pdf_status == "downloaded"


def build_record(ws, row_idx: int, col_map: Dict[str, int]) -> Optional[Record]:
    """
    Build a Record object from a given worksheet row.
    Return None when the row is not usable.
    """
    record_id = str(get_cell(ws, row_idx, col_map, "record_id") or "").strip()
    if not record_id:
        return None

    if ENABLE_BENCHMARK_MODE and BENCHMARK_RECORD_IDS and record_id not in BENCHMARK_RECORD_IDS:
        return None

    title = str(get_cell(ws, row_idx, col_map, "Title") or "").strip()
    doi_raw = str(get_cell(ws, row_idx, col_map, "DOI") or "").strip()
    doi_link_raw = str(get_cell(ws, row_idx, col_map, "DOI Link") or "").strip()

    file_name = build_pdf_filename(record_id, title)
    pdf_path = PDF_BASE_DIR / file_name
    relative_path = f"pdf_files/{file_name}"

    return Record(
        row_idx=row_idx,
        record_id=record_id,
        title=title,
        doi_raw=doi_raw,
        doi_link_raw=doi_link_raw,
        file_name=file_name,
        pdf_path=pdf_path,
        relative_path=relative_path,
    )


def collect_records(ws, col_map: Dict[str, int], max_rows_to_process: Optional[int]) -> List[Record]:
    """
    Read all usable rows from the worksheet and return Record objects.
    """
    records: List[Record] = []
    logical_seen = 0

    for row_idx in range(2, ws.max_row + 1):
        record = build_record(ws, row_idx, col_map)
        if record is None:
            continue

        records.append(record)
        logical_seen += 1

        if max_rows_to_process is not None and logical_seen >= max_rows_to_process:
            break

    return records


def collect_phase0_candidate_rows(
    ws,
    col_map: Dict[str, int],
    max_rows_to_process: Optional[int],
) -> List[int]:
    """
    Return worksheet row indices eligible for DOI enrichment.

    Rules:
    - must have record_id
    - must have Title
    - must NOT already have a usable DOI
    - must NOT already have a downloaded PDF
    """
    candidate_rows: List[int] = []
    logical_seen = 0

    for row_idx in range(2, ws.max_row + 1):
        record_id = str(get_optional_cell(ws, row_idx, col_map, "record_id", "") or "").strip()
        if not record_id:
            continue

        if ENABLE_BENCHMARK_MODE and BENCHMARK_RECORD_IDS and record_id not in BENCHMARK_RECORD_IDS:
            continue

        title = str(get_optional_cell(ws, row_idx, col_map, "Title", "") or "").strip()
        if not title:
            continue

        if row_has_usable_doi(ws, row_idx, col_map):
            continue

        if row_is_already_downloaded(ws, row_idx, col_map):
            continue

        candidate_rows.append(row_idx)
        logical_seen += 1

        if max_rows_to_process is not None and logical_seen >= max_rows_to_process:
            break

    return candidate_rows


def write_result(ws, row_idx: int, col_map: Dict[str, int], result: DownloadResult) -> None:
    """
    Write a DownloadResult back into the worksheet.
    """
    set_cell(ws, row_idx, col_map, "pdf_downloaded", result.pdf_downloaded)
    set_cell(ws, row_idx, col_map, "pdf_download_status", result.pdf_download_status)
    set_cell(ws, row_idx, col_map, "pdf_file_name", result.pdf_file_name)
    set_cell(ws, row_idx, col_map, "pdf_source_url", result.pdf_source_url)
    set_cell(ws, row_idx, col_map, "pdf_local_path", result.pdf_local_path)
    set_cell(
        ws,
        row_idx,
        col_map,
        "pdf_http_status",
        result.pdf_http_status if result.pdf_http_status is not None else "",
    )
    set_cell(ws, row_idx, col_map, "pdf_checked_at", result.pdf_checked_at)


def write_doi_enrichment_result(
    ws,
    row_idx: int,
    col_map: Dict[str, int],
    doi: str,
    doi_link: str,
) -> None:
    """
    Write DOI enrichment results back into the main worksheet.

    Important:
    - only DOI and DOI Link are written
    - no auxiliary DOI columns are created in the main workbook
    """
    set_cell(ws, row_idx, col_map, "DOI", doi)
    set_cell(ws, row_idx, col_map, "DOI Link", doi_link)


def refresh_record_from_worksheet(ws, row_idx: int, col_map: Dict[str, int]) -> Optional[Record]:
    """
    Rebuild a single Record object from the current worksheet state.
    Useful after DOI enrichment changes.
    """
    return build_record(ws, row_idx, col_map)


def refresh_records_from_worksheet(
    ws,
    col_map: Dict[str, int],
    max_rows_to_process: Optional[int] = None,
) -> List[Record]:
    """
    Rebuild the full Record list from the current worksheet state.
    """
    return collect_records(ws, col_map, max_rows_to_process)


def initialize_session_state(
    workbook_path: Path,
    sheet_name: str = SHEET_NAME,
    max_rows_to_process: Optional[int] = None,
) -> SessionState:
    """
    Load workbook + worksheet into a SessionState and build the initial records list.
    """
    wb, ws = load_workbook_and_sheet(workbook_path, sheet_name)
    col_map = build_col_map(ws)
    validate_required_columns(col_map, sheet_name)

    effective_max_rows = max_rows_to_process if max_rows_to_process is not None else MAX_ROWS_TO_PROCESS
    records = collect_records(ws, col_map, effective_max_rows)

    return SessionState(
        workbook_path=workbook_path,
        sheet_name=sheet_name,
        workbook=wb,
        worksheet=ws,
        col_map=col_map,
        records=records,
    )


def refresh_session_records(
    session_state: SessionState,
    max_rows_to_process: Optional[int] = None,
) -> None:
    """
    Rebuild session_state.records from the current worksheet state.
    """
    if not session_state.is_workbook_loaded:
        raise ValueError("SessionState workbook is not loaded.")

    effective_max_rows = max_rows_to_process if max_rows_to_process is not None else MAX_ROWS_TO_PROCESS
    session_state.records = refresh_records_from_worksheet(
        session_state.worksheet,
        session_state.col_map,
        effective_max_rows,
    )


def rebuild_session_column_map(session_state: SessionState) -> None:
    """
    Refresh the header map from the current worksheet.
    """
    if not session_state.is_workbook_loaded:
        raise ValueError("SessionState workbook is not loaded.")

    session_state.col_map = build_col_map(session_state.worksheet)


def safe_save_workbook(wb, output_path: Path, max_attempts: int = 2) -> bool:
    """
    Save workbook safely through a temporary file and replacement.
    """
    output_path.parent.mkdir(parents=True, exist_ok=True)

    for attempt in range(1, max_attempts + 1):
        try:
            temp_path = output_path.with_name(output_path.stem + ".__tmp__.xlsx")
            wb.save(temp_path)

            if output_path.exists():
                try:
                    output_path.unlink()
                except PermissionError:
                    pass

            temp_path.replace(output_path)
            return True

        except KeyboardInterrupt:
            raise
        except Exception:
            if attempt == max_attempts:
                return False
            time.sleep(0.5)

    return False


def save_session_checkpoint(
    session_state: SessionState,
    label: str,
    timestamp: Optional[str] = None,
) -> Optional[Path]:
    """
    Save a temporary checkpoint workbook for recovery/debugging purposes.
    """
    if not session_state.is_workbook_loaded:
        raise ValueError("SessionState workbook is not loaded.")

    checkpoint_path = build_temp_workbook_path(label=label, timestamp=timestamp)
    ok = safe_save_workbook(session_state.workbook, checkpoint_path)
    if not ok:
        return None

    session_state.temp_workbook_path = checkpoint_path
    return checkpoint_path


def save_final_workbook(
    session_state: SessionState,
    output_path: Optional[Path] = None,
    timestamp: Optional[str] = None,
) -> Optional[Path]:
    """
    Save the final workbook for the session.
    """
    if not session_state.is_workbook_loaded:
        raise ValueError("SessionState workbook is not loaded.")

    final_path = output_path or build_final_workbook_path(timestamp=timestamp)
    ok = safe_save_workbook(session_state.workbook, final_path)
    if not ok:
        return None

    session_state.final_workbook_path = final_path
    return final_path