from __future__ import annotations

import os
import re
import time
from concurrent.futures import FIRST_COMPLETED, ThreadPoolExecutor, wait
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Dict, List, Optional, Tuple
from urllib.parse import urljoin

import requests
from openpyxl import load_workbook
from requests.adapters import HTTPAdapter
from urllib3.util.retry import Retry


# =========================
# CONFIGURATION
# =========================

EXTERNAL_WOS_DIR = Path(
    r"C:\Users\Luís Pinto Coelho\Desktop\Dissertação\metanalysis development\2 - literature search & study screening\data\wos"
)

EXCEL_INPUT_DIR = EXTERNAL_WOS_DIR / "excel_files"
PDF_BASE_DIR = EXTERNAL_WOS_DIR / "pdf_files"

PREFERRED_WORKBOOK_PATH = EXCEL_INPUT_DIR / "wos_workbook_03.03.2026_v2.xlsx"

if PREFERRED_WORKBOOK_PATH.exists():
    WORKBOOK_PATH = PREFERRED_WORKBOOK_PATH
else:
    candidates = sorted(
        [p for p in EXCEL_INPUT_DIR.glob("*.xlsx") if not p.name.startswith("~$")],
        key=lambda p: p.stat().st_mtime,
        reverse=True,
    )

    if not candidates:
        raise FileNotFoundError(f"No .xlsx workbook found in: {EXCEL_INPUT_DIR}")

    WORKBOOK_PATH = candidates[0]

OUTPUT_WORKBOOK_PATH = EXCEL_INPUT_DIR / f"{WORKBOOK_PATH.stem}_pdf_downloaded.xlsx"

SHEET_NAME = "wos_full_text"

# Parallelism
MAX_WORKERS = 4

# Performance-oriented but conservative settings
CONNECT_TIMEOUT = 8
READ_TIMEOUT = 15
SLEEP_BETWEEN_REQUESTS = 0.05
MAX_ROWS_TO_PROCESS: Optional[int] = None

# Benchmark mode
ENABLE_BENCHMARK_MODE = False
BENCHMARK_RECORD_IDS: set[str] = set()

# Save safety checkpoint every N processed rows
SAVE_EVERY_N_ROWS = 100

# Dashboard refresh cadence
REFRESH_DASHBOARD_EVERY_N_COMPLETIONS = 5

# Optional future extension hook
ENABLE_SPECIALIZED_RESOLVERS = False

HEADERS = {
    "User-Agent": (
        "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
        "AppleWebKit/537.36 (KHTML, like Gecko) "
        "Chrome/122.0 Safari/537.36"
    ),
    "Accept": (
        "text/html,application/xhtml+xml,application/xml;q=0.9,"
        "application/pdf,*/*;q=0.8"
    ),
}

REQUIRED_INPUT_COLUMNS = [
    "record_id",
    "Title",
    "DOI",
    "DOI Link",
]

REQUIRED_OUTPUT_COLUMNS = [
    "pdf_downloaded",
    "pdf_download_status",
    "pdf_file_name",
    "pdf_source_url",
    "pdf_local_path",
    "pdf_http_status",
    "pdf_checked_at",
]

REQUIRED_COLUMNS = REQUIRED_INPUT_COLUMNS + REQUIRED_OUTPUT_COLUMNS


# =========================
# DATA MODELS
# =========================

@dataclass
class Record:
    row_idx: int
    record_id: str
    title: str
    doi_raw: str
    doi_link_raw: str
    file_name: str
    pdf_path: Path
    relative_path: str


@dataclass
class DownloadResult:
    pdf_downloaded: int
    pdf_download_status: str
    pdf_file_name: str
    pdf_source_url: str
    pdf_local_path: str
    pdf_http_status: Optional[int]
    pdf_checked_at: str


# =========================
# GENERIC HELPERS
# =========================

def now_str() -> str:
    return datetime.now().strftime("%Y-%m-%d %H:%M:%S")


def is_blank(value) -> bool:
    return value is None or str(value).strip() == ""


def sanitize_filename(text: str, max_len: int = 140) -> str:
    text = str(text or "").strip()

    text = re.sub(r'[\\/:*?"<>|]', " ", text)
    text = re.sub(r"[\r\n\t]+", " ", text)
    text = re.sub(r"\s+", " ", text).strip()
    text = text.replace(" ", "_")
    text = re.sub(r"_+", "_", text).strip("_")

    if not text:
        text = "untitled"

    return text[:max_len].rstrip("_")


def build_pdf_filename(record_id: str, title: str) -> str:
    safe_title = sanitize_filename(title)
    return f"{record_id}__{safe_title}.pdf"


def normalize_doi(doi: str) -> Optional[str]:
    if is_blank(doi):
        return None

    doi = str(doi).strip()
    doi = re.sub(r"^https?://(dx\.)?doi\.org/", "", doi, flags=re.IGNORECASE)
    doi = doi.strip()

    if re.match(r"^10\.\d{4,9}/\S+$", doi):
        return doi

    return None


def make_doi_url(doi: str) -> str:
    return f"https://doi.org/{doi}"


def response_looks_like_pdf(resp: requests.Response) -> bool:
    content_type = (resp.headers.get("Content-Type") or "").lower()
    if "application/pdf" in content_type:
        return True

    try:
        return resp.content[:5] == b"%PDF-"
    except Exception:
        return False


def save_pdf_bytes(pdf_path: Path, content: bytes) -> bool:
    if not content:
        return False

    pdf_path.parent.mkdir(parents=True, exist_ok=True)
    pdf_path.write_bytes(content)

    return pdf_path.exists() and pdf_path.stat().st_size > 0


def extract_pdf_url_from_html(base_url: str, html: str) -> Optional[str]:
    if not html:
        return None

    patterns = [
        r'<meta[^>]+name=["\']citation_pdf_url["\'][^>]+content=["\']([^"\']+)["\']',
        r'<meta[^>]+content=["\']([^"\']+)["\'][^>]+name=["\']citation_pdf_url["\']',
        r'href=["\']([^"\']+\.pdf(?:\?[^"\']*)?)["\']',
        r'src=["\']([^"\']+\.pdf(?:\?[^"\']*)?)["\']',
    ]

    for pattern in patterns:
        match = re.search(pattern, html, flags=re.IGNORECASE)
        if match:
            return urljoin(base_url, match.group(1))

    return None


def classify_html_failure(status_code: Optional[int], html: str) -> str:
    text = (html or "").lower()

    if status_code == 403:
        return "paywalled"

    paywall_markers = [
        "purchase pdf",
        "buy article",
        "access through your institution",
        "institutional access",
        "subscribe to journal",
        "get access",
        "login via institution",
        "rent this article",
    ]
    if any(marker in text for marker in paywall_markers):
        return "paywalled"

    metadata_markers = [
        "abstract",
        "references",
        "metrics",
        "article info",
    ]
    if any(marker in text for marker in metadata_markers):
        return "metadata_only"

    if status_code == 404:
        return "not_found"

    if status_code is not None and status_code >= 500:
        return "broken_link"

    return "manual_check"


def format_seconds(seconds: float) -> str:
    seconds = max(0, int(seconds))
    hours, rem = divmod(seconds, 3600)
    minutes, secs = divmod(rem, 60)
    return f"{hours:02d}:{minutes:02d}:{secs:02d}"


def deduplicate_urls(urls: List[str]) -> List[str]:
    seen = set()
    unique_urls: List[str] = []

    for url in urls:
        cleaned = str(url or "").strip()
        if not cleaned:
            continue
        if cleaned not in seen:
            seen.add(cleaned)
            unique_urls.append(cleaned)

    return unique_urls


def build_session() -> requests.Session:
    session = requests.Session()

    retry = Retry(
        total=2,
        read=2,
        connect=2,
        backoff_factor=0.3,
        status_forcelist=[429, 500, 502, 503, 504],
        allowed_methods=["GET", "HEAD", "OPTIONS"],
        raise_on_status=False,
    )

    adapter = HTTPAdapter(
        max_retries=retry,
        pool_connections=20,
        pool_maxsize=20,
    )

    session.mount("http://", adapter)
    session.mount("https://", adapter)

    return session


# =========================
# EXCEL HELPERS
# =========================

def build_col_map(ws) -> Dict[str, int]:
    col_map: Dict[str, int] = {}
    for col_idx in range(1, ws.max_column + 1):
        header = ws.cell(row=1, column=col_idx).value
        if header is not None and str(header).strip():
            col_map[str(header).strip()] = col_idx
    return col_map


def validate_required_columns(col_map: Dict[str, int], sheet_name: str) -> None:
    missing = [c for c in REQUIRED_COLUMNS if c not in col_map]
    if missing:
        raise ValueError(f"Missing required columns in '{sheet_name}': {missing}")


def get_cell(ws, row_idx: int, col_map: Dict[str, int], header: str):
    return ws.cell(row=row_idx, column=col_map[header]).value


def set_cell(ws, row_idx: int, col_map: Dict[str, int], header: str, value) -> None:
    ws.cell(row=row_idx, column=col_map[header], value=value)


def build_record(ws, row_idx: int, col_map: Dict[str, int]) -> Optional[Record]:
    record_id = str(get_cell(ws, row_idx, col_map, "record_id") or "").strip()
    if not record_id:
        return None

    if ENABLE_BENCHMARK_MODE and BENCHMARK_RECORD_IDS and record_id not in BENCHMARK_RECORD_IDS:
        return None

    title = str(get_cell(ws, row_idx, col_map, "Title") or "").strip()
    doi_raw = str(get_cell(ws, row_idx, col_map, "DOI") or "").strip()
    doi_link_raw = str(get_cell(ws, row_idx, col_map, "DOI Link") or "").strip()

    file_name = build_pdf_filename(record_id, title)
    pdf_path = PDF_BASE_DIR / file_name
    relative_path = f"pdf_files/{file_name}"

    return Record(
        row_idx=row_idx,
        record_id=record_id,
        title=title,
        doi_raw=doi_raw,
        doi_link_raw=doi_link_raw,
        file_name=file_name,
        pdf_path=pdf_path,
        relative_path=relative_path,
    )


def write_result(ws, row_idx: int, col_map: Dict[str, int], result: DownloadResult) -> None:
    set_cell(ws, row_idx, col_map, "pdf_downloaded", result.pdf_downloaded)
    set_cell(ws, row_idx, col_map, "pdf_download_status", result.pdf_download_status)
    set_cell(ws, row_idx, col_map, "pdf_file_name", result.pdf_file_name)
    set_cell(ws, row_idx, col_map, "pdf_source_url", result.pdf_source_url)
    set_cell(ws, row_idx, col_map, "pdf_local_path", result.pdf_local_path)
    set_cell(
        ws,
        row_idx,
        col_map,
        "pdf_http_status",
        result.pdf_http_status if result.pdf_http_status is not None else "",
    )
    set_cell(ws, row_idx, col_map, "pdf_checked_at", result.pdf_checked_at)


def collect_records(ws, col_map: Dict[str, int], max_rows_to_process: Optional[int]) -> List[Record]:
    records: List[Record] = []
    logical_seen = 0

    for row_idx in range(2, ws.max_row + 1):
        record = build_record(ws, row_idx, col_map)
        if record is None:
            continue

        records.append(record)
        logical_seen += 1

        if max_rows_to_process is not None and logical_seen >= max_rows_to_process:
            break

    return records


# =========================
# DOWNLOAD CORE
# =========================

def try_download_from_url(
    url: str,
    session: requests.Session,
) -> Tuple[bool, str, Optional[int], Optional[bytes], str]:
    """
    Returns:
        (success, status_label, http_status, pdf_bytes, final_url_used)
    """
    try:
        resp = session.get(
            url,
            headers=HEADERS,
            timeout=(CONNECT_TIMEOUT, READ_TIMEOUT),
            allow_redirects=True,
        )
        http_status = resp.status_code
        final_url = resp.url

        if http_status >= 400:
            if http_status == 403:
                return False, "paywalled", http_status, None, final_url
            if http_status == 404:
                return False, "not_found", http_status, None, final_url
            return False, "broken_link", http_status, None, final_url

        if response_looks_like_pdf(resp):
            return True, "downloaded", http_status, resp.content, final_url

        html = resp.text or ""
        pdf_url = extract_pdf_url_from_html(final_url, html)

        if pdf_url:
            pdf_resp = session.get(
                pdf_url,
                headers=HEADERS,
                timeout=(CONNECT_TIMEOUT, READ_TIMEOUT),
                allow_redirects=True,
            )
            pdf_http_status = pdf_resp.status_code
            pdf_final_url = pdf_resp.url

            if pdf_http_status >= 400:
                if pdf_http_status == 403:
                    return False, "paywalled", pdf_http_status, None, pdf_final_url
                if pdf_http_status == 404:
                    return False, "not_found", pdf_http_status, None, pdf_final_url
                return False, "broken_link", pdf_http_status, None, pdf_final_url

            if response_looks_like_pdf(pdf_resp):
                return True, "downloaded", pdf_http_status, pdf_resp.content, pdf_final_url

            return False, "manual_check", pdf_http_status, None, pdf_final_url

        return False, classify_html_failure(http_status, html), http_status, None, final_url

    except requests.Timeout:
        return False, "broken_link", None, None, url
    except requests.RequestException:
        return False, "broken_link", None, None, url


def build_candidate_urls(record: Record) -> List[str]:
    urls: List[str] = []

    doi_link = str(record.doi_link_raw or "").strip()
    doi = normalize_doi(record.doi_raw)

    if doi_link:
        urls.append(doi_link)

    if doi:
        urls.append(make_doi_url(doi))

    return deduplicate_urls(urls)


def make_result(
    record: Record,
    downloaded: int,
    status: str,
    source_url: str = "",
    local_path: str = "",
    http_status: Optional[int] = None,
) -> DownloadResult:
    return DownloadResult(
        pdf_downloaded=downloaded,
        pdf_download_status=status,
        pdf_file_name=record.file_name,
        pdf_source_url=source_url,
        pdf_local_path=local_path,
        pdf_http_status=http_status,
        pdf_checked_at=now_str(),
    )


def try_base_download(record: Record, session: requests.Session) -> DownloadResult:
    if record.pdf_path.exists() and record.pdf_path.stat().st_size > 0:
        return make_result(
            record=record,
            downloaded=1,
            status="duplicate_pdf",
            local_path=record.relative_path,
        )

    candidate_urls = build_candidate_urls(record)

    if not candidate_urls:
        return make_result(
            record=record,
            downloaded=0,
            status="invalid_doi",
        )

    final_status = "not_found"
    final_http_status: Optional[int] = None
    final_source_url = ""

    for url in candidate_urls:
        success, status_label, http_status, pdf_bytes, final_url = try_download_from_url(url, session)

        final_status = status_label
        final_http_status = http_status
        final_source_url = final_url or url

        if success and pdf_bytes:
            saved = save_pdf_bytes(record.pdf_path, pdf_bytes)
            if saved:
                return make_result(
                    record=record,
                    downloaded=1,
                    status="downloaded",
                    source_url=final_source_url,
                    local_path=record.relative_path,
                    http_status=http_status,
                )

            final_status = "manual_check"
            final_http_status = http_status
            final_source_url = final_url or url

    return make_result(
        record=record,
        downloaded=0,
        status=final_status,
        source_url=final_source_url,
        http_status=final_http_status,
    )


def try_specialized_resolvers(
    record: Record,
    session: requests.Session,
    previous_result: DownloadResult,
) -> DownloadResult:
    return previous_result


def process_record(record: Record) -> Tuple[Record, DownloadResult]:
    session = build_session()
    try:
        result = try_base_download(record, session)

        if result.pdf_downloaded == 1:
            return record, result

        if ENABLE_SPECIALIZED_RESOLVERS:
            result = try_specialized_resolvers(record, session, result)

        return record, result
    finally:
        session.close()


# =========================
# TERMINAL DASHBOARD
# =========================

def clear_terminal() -> None:
    os.system("cls" if os.name == "nt" else "clear")


def render_terminal_dashboard(
    processed: int,
    total: int,
    workers: int,
    start_time: float,
    last_record_id: str,
    last_status: str,
    success_count: int,
    downloaded_now_count: int,
) -> None:
    remaining_count = max(0, total - processed)
    failed_count = max(0, processed - success_count)
    elapsed = time.time() - start_time

    if processed > 0:
        estimated_total = (elapsed / processed) * total
        estimated_remaining = max(0.0, estimated_total - elapsed)
        success_rate = (success_count / processed) * 100
    else:
        estimated_total = 0.0
        estimated_remaining = 0.0
        success_rate = 0.0

    clear_terminal()

    print("=" * 78)
    print("PDF DOWNLOAD STATUS")
    print("=" * 78)
    print(f"Artigos diagnosticados : {processed}")
    print(f"Artigos por diagnosticar: {remaining_count}")
    print(f"Total a diagnosticar    : {total}")
    print(f"Workers paralelos       : {workers}")
    print("-" * 78)
    print(f"PDFs disponíveis        : {success_count}")
    print(f"PDFs descarregados agora: {downloaded_now_count}")
    print(f"Falhas                  : {failed_count}")
    print(f"Taxa de sucesso         : {success_rate:6.2f}%")
    print("-" * 78)
    print(f"Tempo decorrido         : {format_seconds(elapsed)}")
    print(f"Tempo estimado total    : {format_seconds(estimated_total)}")
    print(f"Tempo restante estimado : {format_seconds(estimated_remaining)}")
    print("-" * 78)
    print(f"Último registo          : {last_record_id}")
    print(f"Último estado           : {last_status}")
    print("=" * 78)


# =========================
# MAIN PROCESS
# =========================

def main() -> None:
    if not WORKBOOK_PATH.exists():
        raise FileNotFoundError(f"Workbook not found: {WORKBOOK_PATH}")

    PDF_BASE_DIR.mkdir(parents=True, exist_ok=True)

    wb = load_workbook(WORKBOOK_PATH)
    if SHEET_NAME not in wb.sheetnames:
        raise ValueError(f"Sheet not found: {SHEET_NAME}")

    ws = wb[SHEET_NAME]
    col_map = build_col_map(ws)
    validate_required_columns(col_map, SHEET_NAME)

    records = collect_records(ws, col_map, MAX_ROWS_TO_PROCESS)
    total_to_process = len(records)

    processed = 0
    success_count = 0
    downloaded_now_count = 0
    start_time = time.time()
    last_record_id = ""
    last_status = ""

    render_terminal_dashboard(
        processed=0,
        total=total_to_process,
        workers=MAX_WORKERS,
        start_time=start_time,
        last_record_id="",
        last_status="starting",
        success_count=0,
        downloaded_now_count=0,
    )

    with ThreadPoolExecutor(max_workers=MAX_WORKERS) as executor:
        futures = {}
        next_idx = 0

        while next_idx < total_to_process and len(futures) < MAX_WORKERS:
            record = records[next_idx]
            futures[executor.submit(process_record, record)] = record
            next_idx += 1
            if SLEEP_BETWEEN_REQUESTS > 0:
                time.sleep(SLEEP_BETWEEN_REQUESTS)

        while futures:
            done, _ = wait(list(futures.keys()), return_when=FIRST_COMPLETED)

            for future in done:
                original_record = futures.pop(future)
                record, result = future.result()

                write_result(ws, record.row_idx, col_map, result)

                processed += 1
                last_record_id = record.record_id
                last_status = result.pdf_download_status

                if result.pdf_downloaded == 1:
                    success_count += 1

                if result.pdf_download_status == "downloaded":
                    downloaded_now_count += 1

                if (
                    processed == 1
                    or processed % REFRESH_DASHBOARD_EVERY_N_COMPLETIONS == 0
                    or processed == total_to_process
                ):
                    render_terminal_dashboard(
                        processed=processed,
                        total=total_to_process,
                        workers=MAX_WORKERS,
                        start_time=start_time,
                        last_record_id=last_record_id,
                        last_status=last_status,
                        success_count=success_count,
                        downloaded_now_count=downloaded_now_count,
                    )

                if processed % SAVE_EVERY_N_ROWS == 0:
                    wb.save(OUTPUT_WORKBOOK_PATH)

                while next_idx < total_to_process and len(futures) < MAX_WORKERS:
                    next_record = records[next_idx]
                    futures[executor.submit(process_record, next_record)] = next_record
                    next_idx += 1
                    if SLEEP_BETWEEN_REQUESTS > 0:
                        time.sleep(SLEEP_BETWEEN_REQUESTS)

    wb.save(OUTPUT_WORKBOOK_PATH)

    render_terminal_dashboard(
        processed=processed,
        total=total_to_process,
        workers=MAX_WORKERS,
        start_time=start_time,
        last_record_id=last_record_id,
        last_status=last_status if last_status else "finished",
        success_count=success_count,
        downloaded_now_count=downloaded_now_count,
    )

    print()
    print(f"PDFs disponíveis        : {success_count}")
    print(f"PDFs descarregados agora: {downloaded_now_count}")
    print(f"Saved workbook          : {OUTPUT_WORKBOOK_PATH}")
    print(f"PDF folder              : {PDF_BASE_DIR}")


if __name__ == "__main__":
    main()