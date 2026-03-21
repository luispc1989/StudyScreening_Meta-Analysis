from __future__ import annotations

import re
import time
from datetime import datetime
from pathlib import Path
from typing import Optional, Tuple
from urllib.parse import urljoin

import requests
from openpyxl import load_workbook


# =========================
# CONFIGURATION
# =========================

EXTERNAL_WOS_DIR = Path(
    r"C:\Users\Luís Pinto Coelho\Desktop\Dissertação\metanalysis development\2 - literature search & study screening\data\wos"
)

EXCEL_INPUT_DIR = EXTERNAL_WOS_DIR / "excel_files"
PDF_BASE_DIR = EXTERNAL_WOS_DIR / "pdf_files"

PREFERRED_WORKBOOK_PATH = EXCEL_INPUT_DIR / "wos_workbook_03.03.2026_v2.1_with_email_in_view.xlsx"

if PREFERRED_WORKBOOK_PATH.exists():
    WORKBOOK_PATH = PREFERRED_WORKBOOK_PATH
else:
    candidates = sorted(
        [p for p in EXCEL_INPUT_DIR.glob("*.xlsx") if not p.name.startswith("~$")],
        key=lambda p: p.stat().st_mtime,
        reverse=True,
    )

    if not candidates:
        raise FileNotFoundError(f"No .xlsx workbook found in: {EXCEL_INPUT_DIR}")

    WORKBOOK_PATH = candidates[0]

OUTPUT_WORKBOOK_PATH = EXCEL_INPUT_DIR / f"{WORKBOOK_PATH.stem}_pdf_downloaded.xlsx"

SHEET_NAME = "wos_full_text"

REQUEST_TIMEOUT = 25
SLEEP_BETWEEN_REQUESTS = 1.0

# For test runs, use a small number like 10 or 20.
# For full processing, set to None.
MAX_ROWS_TO_PROCESS: Optional[int] = None  # to define the limit of processing

HEADERS = {
    "User-Agent": (
        "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
        "AppleWebKit/537.36 (KHTML, like Gecko) "
        "Chrome/122.0 Safari/537.36"
    ),
    "Accept": (
        "text/html,application/xhtml+xml,application/xml;q=0.9,"
        "application/pdf,*/*;q=0.8"
    ),
}


# =========================
# HELPER FUNCTIONS
# =========================

def now_str() -> str:
    return datetime.now().strftime("%Y-%m-%d %H:%M:%S")


def is_blank(value) -> bool:
    return value is None or str(value).strip() == ""


def sanitize_filename(text: str, max_len: int = 140) -> str:
    text = str(text or "").strip()

    # Replace invalid Windows filename characters
    text = re.sub(r'[\\/:*?"<>|]', " ", text)

    # Remove control characters
    text = re.sub(r"[\r\n\t]+", " ", text)

    # Collapse whitespace
    text = re.sub(r"\s+", " ", text).strip()

    # Convert spaces to underscores
    text = text.replace(" ", "_")

    # Collapse multiple underscores
    text = re.sub(r"_+", "_", text).strip("_")

    if not text:
        text = "untitled"

    return text[:max_len].rstrip("_")


def build_pdf_filename(record_id: str, title: str) -> str:
    safe_title = sanitize_filename(title)
    return f"{record_id}__{safe_title}.pdf"


def normalize_doi(doi: str) -> Optional[str]:
    if is_blank(doi):
        return None

    doi = str(doi).strip()
    doi = re.sub(r"^https?://(dx\.)?doi\.org/", "", doi, flags=re.IGNORECASE)
    doi = doi.strip()

    # Minimal DOI format check
    if re.match(r"^10\.\d{4,9}/\S+$", doi):
        return doi

    return None


def make_doi_url(doi: str) -> str:
    return f"https://doi.org/{doi}"


def response_looks_like_pdf(resp: requests.Response) -> bool:
    content_type = (resp.headers.get("Content-Type") or "").lower()
    if "application/pdf" in content_type:
        return True

    # Fallback: check PDF magic bytes
    try:
        return resp.content[:5] == b"%PDF-"
    except Exception:
        return False


def save_pdf_bytes(pdf_path: Path, content: bytes) -> bool:
    if not content:
        return False

    pdf_path.parent.mkdir(parents=True, exist_ok=True)
    pdf_path.write_bytes(content)

    return pdf_path.exists() and pdf_path.stat().st_size > 0


def extract_pdf_url_from_html(base_url: str, html: str) -> Optional[str]:
    if not html:
        return None

    patterns = [
        r'<meta[^>]+name=["\']citation_pdf_url["\'][^>]+content=["\']([^"\']+)["\']',
        r'<meta[^>]+content=["\']([^"\']+)["\'][^>]+name=["\']citation_pdf_url["\']',
        r'href=["\']([^"\']+\.pdf(?:\?[^"\']*)?)["\']',
        r'src=["\']([^"\']+\.pdf(?:\?[^"\']*)?)["\']',
    ]

    for pattern in patterns:
        match = re.search(pattern, html, flags=re.IGNORECASE)
        if match:
            return urljoin(base_url, match.group(1))

    return None


def classify_html_failure(status_code: Optional[int], html: str) -> str:
    text = (html or "").lower()

    if status_code == 403:
        return "paywalled"

    paywall_markers = [
        "purchase pdf",
        "buy article",
        "access through your institution",
        "institutional access",
        "subscribe to journal",
        "get access",
        "login via institution",
        "rent this article",
    ]
    if any(marker in text for marker in paywall_markers):
        return "paywalled"

    metadata_markers = [
        "abstract",
        "references",
        "metrics",
        "article info",
    ]
    if any(marker in text for marker in metadata_markers):
        return "metadata_only"

    if status_code == 404:
        return "not_found"

    if status_code is not None and status_code >= 500:
        return "broken_link"

    return "manual_check"


def try_download_from_url(
    url: str,
    session: requests.Session,
) -> Tuple[bool, str, Optional[int], Optional[bytes], str]:
    """
    Returns:
        (success, status_label, http_status, pdf_bytes, final_url_used)
    """
    try:
        resp = session.get(
            url,
            headers=HEADERS,
            timeout=REQUEST_TIMEOUT,
            allow_redirects=True,
        )
        http_status = resp.status_code
        final_url = resp.url

        if http_status >= 400:
            if http_status == 403:
                return False, "paywalled", http_status, None, final_url
            if http_status == 404:
                return False, "not_found", http_status, None, final_url
            return False, "broken_link", http_status, None, final_url

        # Direct PDF
        if response_looks_like_pdf(resp):
            return True, "downloaded", http_status, resp.content, final_url

        # Not a direct PDF; inspect HTML
        html = resp.text or ""
        pdf_url = extract_pdf_url_from_html(final_url, html)

        if pdf_url:
            pdf_resp = session.get(
                pdf_url,
                headers=HEADERS,
                timeout=REQUEST_TIMEOUT,
                allow_redirects=True,
            )
            pdf_http_status = pdf_resp.status_code
            pdf_final_url = pdf_resp.url

            if pdf_http_status >= 400:
                if pdf_http_status == 403:
                    return False, "paywalled", pdf_http_status, None, pdf_final_url
                if pdf_http_status == 404:
                    return False, "not_found", pdf_http_status, None, pdf_final_url
                return False, "broken_link", pdf_http_status, None, pdf_final_url

            if response_looks_like_pdf(pdf_resp):
                return True, "downloaded", pdf_http_status, pdf_resp.content, pdf_final_url

            return False, "manual_check", pdf_http_status, None, pdf_final_url

        return False, classify_html_failure(http_status, html), http_status, None, final_url

    except requests.Timeout:
        return False, "broken_link", None, None, url
    except requests.RequestException:
        return False, "broken_link", None, None, url


def set_cell(ws, row_idx: int, col_map: dict[str, int], header: str, value) -> None:
    ws.cell(row=row_idx, column=col_map[header], value=value)


def get_cell(ws, row_idx: int, col_map: dict[str, int], header: str):
    return ws.cell(row=row_idx, column=col_map[header]).value


def clear_pdf_result_fields(
    ws,
    row_idx: int,
    col_map: dict[str, int],
    file_name: str,
) -> None:
    set_cell(ws, row_idx, col_map, "pdf_downloaded", "")
    set_cell(ws, row_idx, col_map, "pdf_download_status", "")
    set_cell(ws, row_idx, col_map, "pdf_file_name", file_name)
    set_cell(ws, row_idx, col_map, "pdf_local_path", "")
    set_cell(ws, row_idx, col_map, "pdf_source_url", "")
    set_cell(ws, row_idx, col_map, "pdf_http_status", "")
    set_cell(ws, row_idx, col_map, "pdf_checked_at", "")


def validate_required_columns(col_map: dict[str, int]) -> None:
    required_columns = [
        "record_id",
        "Title",
        "DOI",
        "DOI Link",
        "pdf_downloaded",
        "pdf_download_status",
        "pdf_file_name",
        "pdf_local_path",
        "pdf_source_url",
        "pdf_http_status",
        "pdf_checked_at",
    ]
    missing = [c for c in required_columns if c not in col_map]
    if missing:
        raise ValueError(f"Missing required columns in '{SHEET_NAME}': {missing}")


# =========================
# MAIN PROCESS
# =========================

def main() -> None:
    if not WORKBOOK_PATH.exists():
        raise FileNotFoundError(f"Workbook not found: {WORKBOOK_PATH}")

    PDF_BASE_DIR.mkdir(parents=True, exist_ok=True)

    wb = load_workbook(WORKBOOK_PATH)
    if SHEET_NAME not in wb.sheetnames:
        raise ValueError(f"Sheet not found: {SHEET_NAME}")

    ws = wb[SHEET_NAME]

    # Build column map from header row
    col_map: dict[str, int] = {}
    for col_idx in range(1, ws.max_column + 1):
        header = ws.cell(row=1, column=col_idx).value
        if header is not None and str(header).strip():
            col_map[str(header).strip()] = col_idx

    validate_required_columns(col_map)

    session = requests.Session()

    processed = 0
    attempted = 0
    success_count = 0

    print(f"Input workbook: {WORKBOOK_PATH}")
    print(f"Output workbook: {OUTPUT_WORKBOOK_PATH}")
    print(f"PDF folder: {PDF_BASE_DIR}")
    print(f"Sheet: {SHEET_NAME}")
    print(f"MAX_ROWS_TO_PROCESS: {MAX_ROWS_TO_PROCESS}")

    for row_idx in range(2, ws.max_row + 1):
        if MAX_ROWS_TO_PROCESS is not None and attempted >= MAX_ROWS_TO_PROCESS:
            break

        record_id = str(get_cell(ws, row_idx, col_map, "record_id") or "").strip()
        title = str(get_cell(ws, row_idx, col_map, "Title") or "").strip()
        doi_raw = get_cell(ws, row_idx, col_map, "DOI")
        doi_link_raw = get_cell(ws, row_idx, col_map, "DOI Link")

        if not record_id:
            continue

        attempted += 1

        file_name = build_pdf_filename(record_id, title)
        relative_path = f"pdf_files/{file_name}"
        pdf_path = PDF_BASE_DIR / file_name

        # If the PDF already exists physically, treat as logical success
        if pdf_path.exists() and pdf_path.stat().st_size > 0:
            set_cell(ws, row_idx, col_map, "pdf_downloaded", 1)
            set_cell(ws, row_idx, col_map, "pdf_download_status", "duplicate_pdf")
            set_cell(ws, row_idx, col_map, "pdf_file_name", file_name)
            set_cell(ws, row_idx, col_map, "pdf_local_path", relative_path)
            set_cell(ws, row_idx, col_map, "pdf_source_url", "")
            set_cell(ws, row_idx, col_map, "pdf_http_status", "")
            set_cell(ws, row_idx, col_map, "pdf_checked_at", now_str())

            processed += 1
            success_count += 1
            print(f"[{attempted}] {record_id}: duplicate_pdf")
            continue

        clear_pdf_result_fields(ws, row_idx, col_map, file_name)

        doi = normalize_doi(str(doi_raw or ""))
        doi_link = str(doi_link_raw or "").strip()

        candidate_urls: list[str] = []

        if doi_link:
            candidate_urls.append(doi_link)

        if doi:
            doi_url = make_doi_url(doi)
            if doi_url not in candidate_urls:
                candidate_urls.append(doi_url)

        if not candidate_urls:
            set_cell(ws, row_idx, col_map, "pdf_downloaded", 0)
            set_cell(ws, row_idx, col_map, "pdf_download_status", "invalid_doi")
            set_cell(ws, row_idx, col_map, "pdf_file_name", file_name)
            set_cell(ws, row_idx, col_map, "pdf_local_path", "")
            set_cell(ws, row_idx, col_map, "pdf_source_url", "")
            set_cell(ws, row_idx, col_map, "pdf_http_status", "")
            set_cell(ws, row_idx, col_map, "pdf_checked_at", now_str())

            processed += 1
            print(f"[{attempted}] {record_id}: invalid_doi")
            time.sleep(SLEEP_BETWEEN_REQUESTS)
            continue

        final_success = False
        final_status = "not_found"
        final_http_status: Optional[int] = None
        final_source_url = ""

        for url in candidate_urls:
            success, status_label, http_status, pdf_bytes, final_url = try_download_from_url(
                url,
                session,
            )

            final_status = status_label
            final_http_status = http_status
            final_source_url = final_url or url

            if success and pdf_bytes:
                saved = save_pdf_bytes(pdf_path, pdf_bytes)
                if saved:
                    set_cell(ws, row_idx, col_map, "pdf_downloaded", 1)
                    set_cell(ws, row_idx, col_map, "pdf_download_status", "downloaded")
                    set_cell(ws, row_idx, col_map, "pdf_file_name", file_name)
                    set_cell(ws, row_idx, col_map, "pdf_local_path", relative_path)
                    set_cell(ws, row_idx, col_map, "pdf_source_url", final_source_url)
                    set_cell(
                        ws,
                        row_idx,
                        col_map,
                        "pdf_http_status",
                        http_status if http_status is not None else "",
                    )
                    set_cell(ws, row_idx, col_map, "pdf_checked_at", now_str())

                    final_success = True
                    success_count += 1
                    print(f"[{attempted}] {record_id}: downloaded")
                    break
                else:
                    final_status = "manual_check"
                    final_success = False

        if not final_success:
            set_cell(ws, row_idx, col_map, "pdf_downloaded", 0)
            set_cell(ws, row_idx, col_map, "pdf_download_status", final_status)
            set_cell(ws, row_idx, col_map, "pdf_file_name", file_name)
            set_cell(ws, row_idx, col_map, "pdf_local_path", "")
            set_cell(ws, row_idx, col_map, "pdf_source_url", final_source_url)
            set_cell(
                ws,
                row_idx,
                col_map,
                "pdf_http_status",
                final_http_status if final_http_status is not None else "",
            )
            set_cell(ws, row_idx, col_map, "pdf_checked_at", now_str())
            print(f"[{attempted}] {record_id}: {final_status}")

        processed += 1
        time.sleep(SLEEP_BETWEEN_REQUESTS)

    wb.save(OUTPUT_WORKBOOK_PATH)

    print("\nDone.")
    print(f"Processed rows: {processed}")
    print(f"Attempted rows: {attempted}")
    print(f"Logical successes: {success_count}")
    print(f"Saved workbook: {OUTPUT_WORKBOOK_PATH}")
    print(f"PDF folder: {PDF_BASE_DIR}")


if __name__ == "__main__":
    main()