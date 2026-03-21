from pathlib import Path
from urllib.parse import urlparse
from collections import Counter, defaultdict
import re
import openpyxl
from openpyxl.styles import Font, PatternFill, Border, Alignment
from openpyxl.utils import get_column_letter


# =========================
# CONFIG
# =========================
INPUT_FILE = Path(r"C:\Users\Luís Pinto Coelho\Desktop\wos_workbook_03.03.2026_v2_pdf_attempted_updated_pdf_downloaded.xlsx")
OUTPUT_FILE = Path(r"C:\Users\Luís Pinto Coelho\Desktop\wos_workbook_03.03.2026_v2_pdf_attempted_updated_pdf_downloaded_error_tabs.xlsx")

SCREENING_SHEET = "wos_screening_view_full_text"
FULLTEXT_SHEET = "wos_full_text"
STATS_SHEET = "pdf_error_stats"

ERROR_SHEETS_ORDER = [
    "paywalled_institutional_access",
    "manual_check",
    "invalid_doi",
    "captcha_or_security_check",
    "paywalled_purchase_required",
    "not_found",
    "broken_link",
]


# =========================
# HELPERS
# =========================
def normalize_text(value):
    if value is None:
        return ""
    return str(value).strip()


def normalize_host(url):
    url = normalize_text(url)
    if not url:
        return "(missing)"

    try:
        parsed = urlparse(url)
        host = (parsed.netloc or "").strip().lower()

        if not host and url.startswith("www."):
            host = url.lower()

        if not host:
            return "(missing)"

        if ":" in host:
            host = host.split(":")[0]

        if host.startswith("www."):
            host = host[4:]

        return host or "(missing)"
    except Exception:
        return "(missing)"


def make_safe_sheet_name(name):
    name = re.sub(r"[:\\/?*\[\]]", "_", name)
    return name[:31]


def clear_sheet_content_and_basic_format(ws):
    for row in ws.iter_rows():
        for cell in row:
            cell.value = None
            cell._style = openpyxl.styles.Style()
            if cell.hyperlink:
                cell._hyperlink = None


def autosize_columns(ws, extra_padding=2, max_width=80):
    for col_cells in ws.columns:
        col_idx = col_cells[0].column
        max_len = 0
        for cell in col_cells:
            try:
                value = "" if cell.value is None else str(cell.value)
            except Exception:
                value = ""
            if len(value) > max_len:
                max_len = len(value)

        width = min(max_len + extra_padding, max_width)
        ws.column_dimensions[get_column_letter(col_idx)].width = max(10, width)


def copy_cell_style(src_cell, dst_cell):
    if src_cell.has_style:
        dst_cell._style = src_cell._style.copy()
    if src_cell.font:
        dst_cell.font = src_cell.font.copy()
    if src_cell.fill:
        dst_cell.fill = src_cell.fill.copy()
    if src_cell.border:
        dst_cell.border = src_cell.border.copy()
    if src_cell.alignment:
        dst_cell.alignment = src_cell.alignment.copy()
    if src_cell.number_format:
        dst_cell.number_format = src_cell.number_format
    if src_cell.protection:
        dst_cell.protection = src_cell.protection.copy()


def find_headers(ws):
    headers = {}
    for col in range(1, ws.max_column + 1):
        value = ws.cell(1, col).value
        if value is not None:
            headers[str(value).strip()] = col
    return headers


def write_row(ws, row_idx, values):
    for col_idx, value in enumerate(values, start=1):
        ws.cell(row=row_idx, column=col_idx, value=value)


# =========================
# LOAD WORKBOOK
# =========================
wb = openpyxl.load_workbook(INPUT_FILE)

ws_screen = wb[SCREENING_SHEET]
ws_full = wb[FULLTEXT_SHEET]

screen_headers = find_headers(ws_screen)
full_headers = find_headers(ws_full)

required_screen = [
    "record_id",
    "pdf_download_status",
    "pdf_source_url",
    "Title",
    "DOI",
    "DOI Link",
]
required_full = [
    "record_id",
    "Authors",
    "Title",
    "DOI",
    "DOI Link",
    "pdf_source_url",
    "pdf_download_status",
]

for col_name in required_screen:
    if col_name not in screen_headers:
        raise ValueError(f"Missing column in {SCREENING_SHEET}: {col_name}")

for col_name in required_full:
    if col_name not in full_headers:
        raise ValueError(f"Missing column in {FULLTEXT_SHEET}: {col_name}")


# =========================
# BUILD LOOKUP FROM wos_full_text
# =========================
full_lookup = {}

for row in range(2, ws_full.max_row + 1):
    record_id = normalize_text(ws_full.cell(row, full_headers["record_id"]).value)
    if not record_id:
        continue

    full_lookup[record_id] = {
        "record_id": record_id,
        "title": ws_full.cell(row, full_headers["Title"]).value,
        "authors": ws_full.cell(row, full_headers["Authors"]).value,
        "doi": ws_full.cell(row, full_headers["DOI"]).value,
        "doi_link": ws_full.cell(row, full_headers["DOI Link"]).value,
        "pdf_source_url": ws_full.cell(row, full_headers["pdf_source_url"]).value,
        "pdf_download_status": ws_full.cell(row, full_headers["pdf_download_status"]).value,
    }


# =========================
# EXTRACT ERRORS FROM wos_screening_view_full_text
# =========================
error_rows = defaultdict(list)
error_counter = Counter()
error_website_counter = defaultdict(Counter)
error_website_combo_counter = Counter()

for row in range(2, ws_screen.max_row + 1):
    pdf_status = normalize_text(ws_screen.cell(row, screen_headers["pdf_download_status"]).value)

    if not pdf_status:
        continue

    if pdf_status.lower() == "downloaded":
        continue

    record_id = normalize_text(ws_screen.cell(row, screen_headers["record_id"]).value)
    pdf_source_url = ws_screen.cell(row, screen_headers["pdf_source_url"]).value
    website = normalize_host(pdf_source_url)

    error_counter[pdf_status] += 1
    error_website_counter[pdf_status][website] += 1
    error_website_combo_counter[(pdf_status, website)] += 1

    full_info = full_lookup.get(record_id, {})

    row_data = {
        "record_id": record_id,
        "title": full_info.get("title") or ws_screen.cell(row, screen_headers["Title"]).value,
        "authors": full_info.get("authors"),
        "doi": full_info.get("doi") or ws_screen.cell(row, screen_headers["DOI"]).value,
        "doi_link": full_info.get("doi_link") or ws_screen.cell(row, screen_headers["DOI Link"]).value,
        "pdf_source_url": pdf_source_url if pdf_source_url else full_info.get("pdf_source_url"),
        "pdf_download_status": pdf_status,
    }

    error_rows[pdf_status].append(row_data)

total_errors = sum(error_counter.values())


# =========================
# REMOVE PREVIOUS OUTPUT SHEETS IF THEY EXIST
# =========================
sheets_to_remove = [STATS_SHEET] + [make_safe_sheet_name(x) for x in ERROR_SHEETS_ORDER]
for sheet_name in sheets_to_remove:
    if sheet_name in wb.sheetnames:
        del wb[sheet_name]


# =========================
# CREATE pdf_error_stats
# =========================
ws_stats = wb.create_sheet(title=STATS_SHEET)

# Title
ws_stats["A1"] = "PDF download status - ranking de erros e top websites"
ws_stats["A1"].font = Font(bold=True, size=14)
ws_stats.merge_cells("A1:F1")

ws_stats["A2"] = f"Base analisada: {SCREENING_SHEET}"
ws_stats["A2"].font = Font(italic=True)

ws_stats["A3"] = "Nota: o ranking de erros exclui o estado 'downloaded'. O campo 'website' corresponde ao host de pdf_source_url."
ws_stats["A3"].font = Font(italic=True)

# Main ranking table
main_headers = ["Rank", "Erro (pdf_download_status)", "N", "% erros", "% total registos"]
main_row = 5
write_row(ws_stats, main_row, main_headers)

header_fill_dark = PatternFill(fill_type="solid", fgColor="000000")
header_font_white = Font(color="FFFFFF", bold=True)

for c in range(1, len(main_headers) + 1):
    cell = ws_stats.cell(main_row, c)
    cell.fill = header_fill_dark
    cell.font = header_font_white
    cell.alignment = Alignment(horizontal="center")

sorted_errors = sorted(error_counter.items(), key=lambda x: (-x[1], x[0].lower()))

row_ptr = main_row + 1
for rank, (error_name, n) in enumerate(sorted_errors, start=1):
    pct_errors = n / total_errors if total_errors else 0
    pct_total = n / (ws_screen.max_row - 1) if ws_screen.max_row > 1 else 0

    ws_stats.cell(row_ptr, 1, rank)
    ws_stats.cell(row_ptr, 2, error_name)
    ws_stats.cell(row_ptr, 3, n)
    ws_stats.cell(row_ptr, 4, pct_errors)
    ws_stats.cell(row_ptr, 5, pct_total)

    ws_stats.cell(row_ptr, 4).number_format = "0.0%"
    ws_stats.cell(row_ptr, 5).number_format = "0.0%"
    row_ptr += 1

# Top combinations block
combo_start_col = 8
combo_headers = ["Rank", "Erro", "Website", "N"]

ws_stats.cell(5, combo_start_col, "Combinações erro + website mais frequentes")
ws_stats.cell(5, combo_start_col).font = Font(color="FFFFFF", bold=True)
ws_stats.cell(5, combo_start_col).fill = PatternFill(fill_type="solid", fgColor="1F4E78")
ws_stats.merge_cells(start_row=5, start_column=combo_start_col, end_row=5, end_column=combo_start_col + 3)

for i, h in enumerate(combo_headers, start=combo_start_col):
    cell = ws_stats.cell(6, i, h)
    cell.fill = PatternFill(fill_type="solid", fgColor="D9E2F3")
    cell.font = Font(bold=True)
    cell.alignment = Alignment(horizontal="center")

combo_sorted = sorted(
    error_website_combo_counter.items(),
    key=lambda x: (-x[1], x[0][0].lower(), x[0][1].lower())
)

combo_row = 7
for rank, ((error_name, website), n) in enumerate(combo_sorted[:15], start=1):
    ws_stats.cell(combo_row, combo_start_col, rank)
    ws_stats.cell(combo_row, combo_start_col + 1, error_name)
    ws_stats.cell(combo_row, combo_start_col + 2, website)
    ws_stats.cell(combo_row, combo_start_col + 3, n)
    combo_row += 1

# Top websites by error
section_row = row_ptr + 2
for error_name, _ in sorted_errors:
    ws_stats.cell(section_row, 1, f"Top websites — {error_name}")
    ws_stats.cell(section_row, 1).font = Font(color="FFFFFF", bold=True)
    ws_stats.cell(section_row, 1).fill = PatternFill(fill_type="solid", fgColor="1F4E78")
    ws_stats.merge_cells(start_row=section_row, start_column=1, end_row=section_row, end_column=3)

    section_row += 1
    write_row(ws_stats, section_row, ["Rank", "Website", "N"])
    for c in range(1, 4):
        cell = ws_stats.cell(section_row, c)
        cell.fill = PatternFill(fill_type="solid", fgColor="D9E2F3")
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center")

    section_row += 1
    websites_sorted = sorted(
        error_website_counter[error_name].items(),
        key=lambda x: (-x[1], x[0].lower())
    )[:10]

    for rank, (website, n) in enumerate(websites_sorted, start=1):
        ws_stats.cell(section_row, 1, rank)
        ws_stats.cell(section_row, 2, website)
        ws_stats.cell(section_row, 3, n)
        section_row += 1

    section_row += 2

autosize_columns(ws_stats)


# =========================
# CREATE ONE SHEET PER ERROR
# =========================
error_sheet_headers = [
    "record_id",
    "title",
    "authors",
    "doi",
    "doi_link",
    "pdf_source_url",
    "pdf_download_status",
]

# Capture a simple header style from wos_full_text row 1, if available
header_style_source = {}
for idx, name in enumerate(["record_id", "Title", "Authors", "DOI", "DOI Link", "pdf_source_url", "pdf_download_status"], start=1):
    if name in full_headers:
        src_cell = ws_full.cell(1, full_headers[name])
        header_style_source[name] = src_cell

for error_name in ERROR_SHEETS_ORDER:
    if error_name not in error_rows:
        continue

    ws_err = wb.create_sheet(title=make_safe_sheet_name(error_name))

    # Header row
    for col_idx, out_header in enumerate(error_sheet_headers, start=1):
        cell = ws_err.cell(1, col_idx, out_header)

        # normal style, close to wos_full_text
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="general", vertical="bottom")

    # Data rows
    current_row = 2
    for item in error_rows[error_name]:
        values = [
            item["record_id"],
            item["title"],
            item["authors"],
            item["doi"],
            item["doi_link"],
            item["pdf_source_url"],
            item["pdf_download_status"],
        ]
        write_row(ws_err, current_row, values)
        current_row += 1

    # Hyperlinks for doi_link and pdf_source_url
    for row_idx in range(2, current_row):
        doi_link_cell = ws_err.cell(row_idx, 5)
        pdf_url_cell = ws_err.cell(row_idx, 6)

        if doi_link_cell.value:
            doi_link_cell.hyperlink = str(doi_link_cell.value)
            doi_link_cell.style = "Hyperlink"

        if pdf_url_cell.value:
            pdf_url_cell.hyperlink = str(pdf_url_cell.value)
            pdf_url_cell.style = "Hyperlink"

    autosize_columns(ws_err)


# =========================
# SAVE
# =========================
wb.save(OUTPUT_FILE)
print(f"Created file: {OUTPUT_FILE}")