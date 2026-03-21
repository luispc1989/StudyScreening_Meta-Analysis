
from __future__ import annotations

import time
from pathlib import Path
from typing import Dict, List, Optional

from openpyxl import load_workbook

from app.config import (
    BENCHMARK_RECORD_IDS,
    ENABLE_BENCHMARK_MODE,
    PDF_BASE_DIR,
    REQUIRED_COLUMNS,
)
from app.models import DownloadResult, Record
from app.utils import build_pdf_filename


def load_workbook_and_sheet(workbook_path: Path, sheet_name: str):
    """
    Open the Excel workbook and return (workbook, worksheet).
    """
    if not workbook_path.exists():
        raise FileNotFoundError(f"Workbook not found: {workbook_path}")

    wb = load_workbook(workbook_path)

    if sheet_name not in wb.sheetnames:
        raise ValueError(f"Sheet not found: {sheet_name}")

    ws = wb[sheet_name]
    return wb, ws


def build_col_map(ws) -> Dict[str, int]:
    """
    Build a mapping from header name to Excel column index.
    """
    col_map: Dict[str, int] = {}

    for col_idx in range(1, ws.max_column + 1):
        header = ws.cell(row=1, column=col_idx).value
        if header is not None and str(header).strip():
            col_map[str(header).strip()] = col_idx

    return col_map


def validate_required_columns(col_map: Dict[str, int], sheet_name: str) -> None:
    """
    Ensure all required columns exist in the target sheet.
    """
    missing = [c for c in REQUIRED_COLUMNS if c not in col_map]
    if missing:
        raise ValueError(f"Missing required columns in '{sheet_name}': {missing}")


def get_cell(ws, row_idx: int, col_map: Dict[str, int], header: str):
    return ws.cell(row=row_idx, column=col_map[header]).value


def set_cell(ws, row_idx: int, col_map: Dict[str, int], header: str, value) -> None:
    ws.cell(row=row_idx, column=col_map[header], value=value)


def build_record(ws, row_idx: int, col_map: Dict[str, int]) -> Optional[Record]:
    """
    Build a Record object from a given worksheet row.
    Return None when the row is not usable.
    """
    record_id = str(get_cell(ws, row_idx, col_map, "record_id") or "").strip()
    if not record_id:
        return None

    if ENABLE_BENCHMARK_MODE and BENCHMARK_RECORD_IDS and record_id not in BENCHMARK_RECORD_IDS:
        return None

    title = str(get_cell(ws, row_idx, col_map, "Title") or "").strip()
    doi_raw = str(get_cell(ws, row_idx, col_map, "DOI") or "").strip()
    doi_link_raw = str(get_cell(ws, row_idx, col_map, "DOI Link") or "").strip()

    file_name = build_pdf_filename(record_id, title)
    pdf_path = PDF_BASE_DIR / file_name
    relative_path = f"pdf_files/{file_name}"

    return Record(
        row_idx=row_idx,
        record_id=record_id,
        title=title,
        doi_raw=doi_raw,
        doi_link_raw=doi_link_raw,
        file_name=file_name,
        pdf_path=pdf_path,
        relative_path=relative_path,
    )


def collect_records(ws, col_map: Dict[str, int], max_rows_to_process: Optional[int]) -> List[Record]:
    """
    Read all usable rows from the worksheet and return Record objects.
    """
    records: List[Record] = []
    logical_seen = 0

    for row_idx in range(2, ws.max_row + 1):
        record = build_record(ws, row_idx, col_map)
        if record is None:
            continue

        records.append(record)
        logical_seen += 1

        if max_rows_to_process is not None and logical_seen >= max_rows_to_process:
            break

    return records


def write_result(ws, row_idx: int, col_map: Dict[str, int], result: DownloadResult) -> None:
    """
    Write a DownloadResult back into the worksheet.
    """
    set_cell(ws, row_idx, col_map, "pdf_downloaded", result.pdf_downloaded)
    set_cell(ws, row_idx, col_map, "pdf_download_status", result.pdf_download_status)
    set_cell(ws, row_idx, col_map, "pdf_file_name", result.pdf_file_name)
    set_cell(ws, row_idx, col_map, "pdf_source_url", result.pdf_source_url)
    set_cell(ws, row_idx, col_map, "pdf_local_path", result.pdf_local_path)
    set_cell(
        ws,
        row_idx,
        col_map,
        "pdf_http_status",
        result.pdf_http_status if result.pdf_http_status is not None else "",
    )
    set_cell(ws, row_idx, col_map, "pdf_checked_at", result.pdf_checked_at)


def safe_save_workbook(wb, output_path: Path, max_attempts: int = 2) -> bool:
    """
    Save workbook safely through a temporary file and replacement.
    """
    for attempt in range(1, max_attempts + 1):
        try:
            temp_path = output_path.with_name(output_path.stem + ".__tmp__.xlsx")
            wb.save(temp_path)

            if output_path.exists():
                try:
                    output_path.unlink()
                except PermissionError:
                    pass

            temp_path.replace(output_path)
            return True

        except KeyboardInterrupt:
            raise
        except Exception:
            if attempt == max_attempts:
                return False
            time.sleep(0.5)

    return False