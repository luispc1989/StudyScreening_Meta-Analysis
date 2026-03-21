from __future__ import annotations

import re
import time
import unicodedata
from dataclasses import dataclass
from difflib import SequenceMatcher
from pathlib import Path
from typing import Dict, List, Optional, Tuple

import requests
from openpyxl import load_workbook


# =========================================================
# CONFIGURATION
# =========================================================

INPUT_XLSX = Path(r"C:\Users\Luís Pinto Coelho\Desktop\Dissertação\metanalysis development\2 - literature search & study screening\data\wos\excel_files\artigos_sem_doi_listagem.xlsx")
OUTPUT_XLSX = Path(r"C:\Users\Luís Pinto Coelho\Desktop\Dissertação\metanalysis development\2 - literature search & study screening\data\wos\excel_files\artigos_sem_doi_listagem_doi_preenchido.xlsx")
SHEET_NAME = "missing_doi"

CROSSREF_MAILTO = "your_email@example.com"   # Replace with your email
REQUEST_TIMEOUT = 25
SLEEP_BETWEEN_REQUESTS = 0.6
MAX_ROWS_TO_PROCESS = None   # None = process all rows
OVERWRITE_EXISTING_DOI = False

MIN_TITLE_SIMILARITY = 0.88
MIN_CONFIDENCE_SCORE = 75.0

USER_AGENT = (
    "DOI-Finder/2.0 "
    "(mailto:{mailto}; Python requests; workbook DOI enrichment)"
)


# =========================================================
# DATA STRUCTURES
# =========================================================

@dataclass
class Candidate:
    doi: str
    source: str
    title: str
    year: Optional[int]
    authors: List[str]
    journal: str = ""
    score: float = 0.0
    title_similarity: float = 0.0
    year_score: float = 0.0
    author_score: float = 0.0


# =========================================================
# TEXT HELPERS
# =========================================================

def strip_accents(text: str) -> str:
    text = unicodedata.normalize("NFKD", text)
    return "".join(ch for ch in text if not unicodedata.combining(ch))


def normalize_text(text: Optional[str]) -> str:
    if not text:
        return ""
    text = strip_accents(str(text)).lower()
    text = text.replace("&", " and ")
    text = re.sub(r"[\u2010-\u2015]", "-", text)
    text = re.sub(r"[^a-z0-9\s-]", " ", text)
    text = re.sub(r"\s+", " ", text).strip()
    return text


def compact_title(text: Optional[str]) -> str:
    text = normalize_text(text)
    text = re.sub(r"\b(a|an|the)\b", " ", text)
    text = re.sub(r"\s+", " ", text).strip()
    return text


def title_similarity(a: str, b: str) -> float:
    a_norm = compact_title(a)
    b_norm = compact_title(b)
    if not a_norm or not b_norm:
        return 0.0
    return SequenceMatcher(None, a_norm, b_norm).ratio()


def parse_year(value) -> Optional[int]:
    if value is None:
        return None
    text = str(value).strip()
    m = re.search(r"(19|20)\d{2}", text)
    return int(m.group(0)) if m else None


def split_authors(authors_text: Optional[str]) -> List[str]:
    if not authors_text:
        return []
    text = str(authors_text)
    text = text.replace(";", "|").replace(",", "|")
    parts = [p.strip() for p in text.split("|") if p.strip()]
    cleaned = []
    for part in parts:
        part_norm = normalize_text(part)
        if part_norm:
            cleaned.append(part_norm)
    return cleaned


def extract_family_names(authors: List[str]) -> List[str]:
    families = []
    for author in authors:
        tokens = author.split()
        if tokens:
            families.append(tokens[-1])
    return families


def score_year(expected: Optional[int], candidate: Optional[int]) -> float:
    if expected is None or candidate is None:
        return 0.5
    diff = abs(expected - candidate)
    if diff == 0:
        return 1.0
    if diff == 1:
        return 0.7
    if diff == 2:
        return 0.4
    return 0.0


def score_authors(expected_authors: List[str], candidate_authors: List[str]) -> float:
    if not expected_authors or not candidate_authors:
        return 0.5
    expected_families = set(extract_family_names(expected_authors))
    candidate_families = set(extract_family_names(candidate_authors))
    if not expected_families or not candidate_families:
        return 0.5
    overlap = expected_families.intersection(candidate_families)
    if not overlap:
        return 0.0
    return len(overlap) / max(1, min(len(expected_families), len(candidate_families)))


def build_confidence_score(
    expected_title: str,
    expected_year: Optional[int],
    expected_authors: List[str],
    candidate: Candidate,
) -> Candidate:
    t_score = title_similarity(expected_title, candidate.title)
    y_score = score_year(expected_year, candidate.year)
    a_score = score_authors(expected_authors, candidate.authors)
    total = (t_score * 70.0) + (y_score * 15.0) + (a_score * 15.0)

    candidate.title_similarity = t_score
    candidate.year_score = y_score
    candidate.author_score = a_score
    candidate.score = total
    return candidate


# =========================================================
# API HELPERS
# =========================================================

def get_session(mailto: str) -> requests.Session:
    session = requests.Session()
    session.headers.update(
        {
            "User-Agent": USER_AGENT.format(mailto=mailto),
            "Accept": "application/json",
        }
    )
    return session


def safe_get_json(session: requests.Session, url: str, params: Optional[Dict] = None) -> Optional[Dict]:
    try:
        response = session.get(url, params=params, timeout=REQUEST_TIMEOUT)
        response.raise_for_status()
        return response.json()
    except Exception:
        return None


def sleep_short() -> None:
    time.sleep(SLEEP_BETWEEN_REQUESTS)


def search_crossref(
    session: requests.Session,
    title: str,
    year: Optional[int],
    authors_text: Optional[str],
) -> List[Candidate]:
    bibliographic = title
    if year:
        bibliographic += f" {year}"
    if authors_text:
        bibliographic += f" {authors_text}"

    params = {
        "query.bibliographic": bibliographic,
        "rows": 8,
        "select": "DOI,title,author,published-print,published-online,issued,container-title",
        "mailto": CROSSREF_MAILTO,
    }
    data = safe_get_json(session, "https://api.crossref.org/works", params=params)
    if not data:
        return []

    items = data.get("message", {}).get("items", [])
    candidates: List[Candidate] = []

    for item in items:
        doi = item.get("DOI")
        titles = item.get("title") or []
        cand_title = titles[0] if titles else ""
        cand_year = None

        for field in ("published-print", "published-online", "issued"):
            date_parts = item.get(field, {}).get("date-parts")
            if date_parts and date_parts[0]:
                cand_year = date_parts[0][0]
                break

        authors = []
        for author in item.get("author", []) or []:
            full = " ".join(
                p for p in [author.get("given", ""), author.get("family", "")] if p
            ).strip()
            if full:
                authors.append(normalize_text(full))

        journal_list = item.get("container-title") or []
        journal = journal_list[0] if journal_list else ""

        if doi:
            candidates.append(
                Candidate(
                    doi=doi.strip(),
                    source="Crossref",
                    title=cand_title,
                    year=cand_year,
                    authors=authors,
                    journal=journal,
                )
            )

    return candidates


def search_openalex(session: requests.Session, title: str) -> List[Candidate]:
    params = {
        "search": title,
        "per_page": 8,
        "select": "doi,display_name,publication_year,authorships,primary_location,host_venue",
    }
    data = safe_get_json(session, "https://api.openalex.org/works", params=params)
    if not data:
        return []

    results = data.get("results", [])
    candidates: List[Candidate] = []

    for item in results:
        doi_url = item.get("doi")
        if not doi_url:
            continue
        doi = doi_url.replace("https://doi.org/", "").replace("http://doi.org/", "").strip()

        authors = []
        for auth in item.get("authorships", []) or []:
            name = ((auth.get("author") or {}).get("display_name")) or ""
            if name:
                authors.append(normalize_text(name))

        journal = ""
        host_venue = item.get("host_venue") or {}
        primary_location = item.get("primary_location") or {}
        source = (primary_location.get("source") or {})
        if source.get("display_name"):
            journal = source["display_name"]
        elif host_venue.get("display_name"):
            journal = host_venue["display_name"]

        candidates.append(
            Candidate(
                doi=doi,
                source="OpenAlex",
                title=item.get("display_name", "") or "",
                year=item.get("publication_year"),
                authors=authors,
                journal=journal,
            )
        )

    return candidates


def search_datacite(session: requests.Session, title: str) -> List[Candidate]:
    params = {
        "query": title,
        "page[size]": 8,
    }
    data = safe_get_json(session, "https://api.datacite.org/dois", params=params)
    if not data:
        return []

    candidates: List[Candidate] = []
    for item in data.get("data", []) or []:
        attrs = item.get("attributes", {}) or {}
        doi = attrs.get("doi") or ""
        if not doi:
            continue

        titles = attrs.get("titles") or []
        cand_title = ""
        if titles and isinstance(titles, list):
            first = titles[0]
            if isinstance(first, dict):
                cand_title = first.get("title", "")
            else:
                cand_title = str(first)

        year = attrs.get("publicationYear")
        try:
            year = int(year) if year is not None else None
        except Exception:
            year = None

        creators = attrs.get("creators") or []
        authors = []
        for creator in creators:
            if isinstance(creator, dict):
                name = creator.get("name") or ""
                if name:
                    authors.append(normalize_text(name))

        publisher = attrs.get("publisher", "") or ""
        candidates.append(
            Candidate(
                doi=doi.strip(),
                source="DataCite",
                title=cand_title,
                year=year,
                authors=authors,
                journal=publisher,
            )
        )
    return candidates


def search_europe_pmc(
    session: requests.Session,
    title: str,
    year: Optional[int],
) -> List[Candidate]:
    query = f'TITLE:"{title}"'
    if year:
        query += f" AND PUB_YEAR:{year}"

    params = {
        "query": query,
        "format": "json",
        "pageSize": 8,
        "resultType": "core",
    }
    data = safe_get_json(session, "https://www.ebi.ac.uk/europepmc/webservices/rest/search", params=params)
    if not data:
        return []

    candidates: List[Candidate] = []
    result_list = (data.get("resultList") or {}).get("result", []) or []
    for item in result_list:
        doi = item.get("doi") or ""
        if not doi:
            continue

        authors = []
        author_string = item.get("authorString") or ""
        if author_string:
            for part in re.split(r",|;", author_string):
                part = part.strip()
                if part:
                    authors.append(normalize_text(part))

        pub_year = item.get("pubYear")
        try:
            pub_year = int(pub_year) if pub_year else None
        except Exception:
            pub_year = None

        journal = item.get("journalTitle") or ""
        cand_title = item.get("title") or ""

        candidates.append(
            Candidate(
                doi=doi.strip(),
                source="EuropePMC",
                title=cand_title,
                year=pub_year,
                authors=authors,
                journal=journal,
            )
        )

    return candidates


def deduplicate_candidates(candidates: List[Candidate]) -> List[Candidate]:
    seen = set()
    unique = []
    for cand in candidates:
        key = (cand.doi.lower().strip(), compact_title(cand.title))
        if key in seen:
            continue
        seen.add(key)
        unique.append(cand)
    return unique


def pick_best_candidate(
    title: str,
    year: Optional[int],
    authors_text: Optional[str],
    candidates: List[Candidate],
) -> Optional[Candidate]:
    expected_authors = split_authors(authors_text)
    rescored = [
        build_confidence_score(title, year, expected_authors, cand)
        for cand in candidates
        if cand.doi
    ]
    if not rescored:
        return None

    rescored.sort(key=lambda c: (c.score, c.title_similarity, c.author_score), reverse=True)
    best = rescored[0]

    if best.title_similarity < MIN_TITLE_SIMILARITY:
        return None
    if best.score < MIN_CONFIDENCE_SCORE:
        return None

    return best


def find_best_doi(
    session: requests.Session,
    title: str,
    year: Optional[int],
    authors_text: Optional[str],
) -> Tuple[Optional[Candidate], List[Candidate]]:
    all_candidates: List[Candidate] = []

    all_candidates.extend(search_crossref(session, title, year, authors_text))
    sleep_short()
    all_candidates.extend(search_openalex(session, title))
    sleep_short()
    all_candidates.extend(search_datacite(session, title))
    sleep_short()
    all_candidates.extend(search_europe_pmc(session, title, year))
    sleep_short()

    all_candidates = deduplicate_candidates(all_candidates)

    expected_authors = split_authors(authors_text)
    rescored = [
        build_confidence_score(title, year, expected_authors, cand)
        for cand in all_candidates
    ]
    rescored.sort(key=lambda c: c.score, reverse=True)

    best = rescored[0] if rescored else None
    if best and best.title_similarity >= MIN_TITLE_SIMILARITY and best.score >= MIN_CONFIDENCE_SCORE:
        return best, rescored
    return None, rescored


# =========================================================
# EXCEL HELPERS
# =========================================================

def get_header_map(ws) -> Dict[str, int]:
    header_map: Dict[str, int] = {}
    for col in range(1, ws.max_column + 1):
        value = ws.cell(row=1, column=col).value
        if value:
            header_map[str(value).strip()] = col
    return header_map


def ensure_column(ws, header_map: Dict[str, int], column_name: str) -> int:
    if column_name in header_map:
        return header_map[column_name]
    new_col = ws.max_column + 1
    ws.cell(row=1, column=new_col).value = column_name
    header_map[column_name] = new_col
    return new_col


def doi_to_link(doi: str) -> str:
    doi = doi.strip()
    return f"https://doi.org/{doi}" if doi else ""


# =========================================================
# MAIN
# =========================================================

def main() -> int:
    if not INPUT_XLSX.exists():
        print(f"[ERROR] Input file not found: {INPUT_XLSX}")
        return 1

    wb = load_workbook(INPUT_XLSX)
    if SHEET_NAME not in wb.sheetnames:
        print(f"[ERROR] Sheet not found: {SHEET_NAME}")
        return 1

    ws = wb[SHEET_NAME]
    headers = get_header_map(ws)

    required = ["Title", "Authors", "Publication Year", "DOI", "DOI Link"]
    missing = [col for col in required if col not in headers]
    if missing:
        print(f"[ERROR] Missing required columns: {missing}")
        return 1

    doi_source_col = ensure_column(ws, headers, "doi_source")
    doi_score_col = ensure_column(ws, headers, "doi_confidence")
    doi_match_title_col = ensure_column(ws, headers, "doi_matched_title")
    doi_match_year_col = ensure_column(ws, headers, "doi_matched_year")
    doi_match_journal_col = ensure_column(ws, headers, "doi_matched_journal")
    doi_notes_col = ensure_column(ws, headers, "doi_notes")

    session = get_session(CROSSREF_MAILTO)

    processed = 0
    found = 0
    skipped_existing = 0
    unresolved = 0

    for row in range(2, ws.max_row + 1):
        if MAX_ROWS_TO_PROCESS is not None and processed >= MAX_ROWS_TO_PROCESS:
            break

        title = ws.cell(row=row, column=headers["Title"]).value
        authors_text = ws.cell(row=row, column=headers["Authors"]).value
        year = parse_year(ws.cell(row=row, column=headers["Publication Year"]).value)
        doi_value = ws.cell(row=row, column=headers["DOI"]).value

        if not title or not str(title).strip():
            continue

        if doi_value and str(doi_value).strip() and not OVERWRITE_EXISTING_DOI:
            skipped_existing += 1
            continue

        processed += 1
        print(f"[{processed}] Searching DOI for: {str(title)[:110]}")

        best, candidates = find_best_doi(session, str(title), year, str(authors_text or ""))

        if best:
            ws.cell(row=row, column=headers["DOI"]).value = best.doi
            ws.cell(row=row, column=headers["DOI Link"]).value = doi_to_link(best.doi)
            ws.cell(row=row, column=doi_source_col).value = best.source
            ws.cell(row=row, column=doi_score_col).value = round(best.score, 2)
            ws.cell(row=row, column=doi_match_title_col).value = best.title
            ws.cell(row=row, column=doi_match_year_col).value = best.year
            ws.cell(row=row, column=doi_match_journal_col).value = best.journal
            ws.cell(row=row, column=doi_notes_col).value = (
                f"title_sim={best.title_similarity:.3f}; "
                f"year_score={best.year_score:.3f}; "
                f"author_score={best.author_score:.3f}"
            )
            found += 1
        else:
            unresolved += 1
            top_preview = "; ".join(
                f"{c.source}:{c.doi} [{c.score:.1f}]"
                for c in candidates[:4]
            )
            ws.cell(row=row, column=doi_source_col).value = ""
            ws.cell(row=row, column=doi_score_col).value = ""
            ws.cell(row=row, column=doi_match_title_col).value = ""
            ws.cell(row=row, column=doi_match_year_col).value = ""
            ws.cell(row=row, column=doi_match_journal_col).value = ""
            ws.cell(row=row, column=doi_notes_col).value = (
                f"No confident match. Top candidates: {top_preview}" if top_preview else "No candidates found."
            )

        if processed % 20 == 0:
            wb.save(OUTPUT_XLSX)
            print(f"Progress saved to: {OUTPUT_XLSX}")

    wb.save(OUTPUT_XLSX)

    print("\\nDONE")
    print(f"Processed rows:   {processed}")
    print(f"Found DOI:        {found}")
    print(f"Unresolved:       {unresolved}")
    print(f"Skipped existing: {skipped_existing}")
    print(f"Output file:      {OUTPUT_XLSX.resolve()}")

    return 0


if __name__ == "__main__":
    raise SystemExit(main())
